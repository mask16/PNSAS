# -*- coding: utf-8 -*-
"""AS이력부품내역.xlsx 를 기존 접수의 report_details.parts 에 등록한다.

기본: 엑셀에 있는 접수번호의 부품 목록을 엑셀 기준으로 넣는다.
이미 부품이 있으면 건너뛴다. 덮어쓰려면: python import_as_parts_xlsx.py --replace
"""
from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

import app as m

SEED_MARKER = "__xlsx_as_parts__"
DEFAULT_XLSX = Path(r"c:\Users\AdministratorB\Desktop\AS이력부품내역.xlsx")


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text.lower() in {"none", "nan"}:
        return ""
    return text


def _qty(value):
    text = _cell(value).replace(",", "")
    if not text:
        return 1
    try:
        num = float(text)
    except ValueError:
        return 1
    if num.is_integer():
        return int(num)
    return num


def _price(value):
    text = _cell(value).replace(",", "").replace("원", "")
    if not text:
        return 0
    try:
        return int(round(float(text)))
    except ValueError:
        return 0


def _row_map(headers, values):
    out = {}
    for i, header in enumerate(headers):
        key = _cell(header)
        if key:
            out[key] = values[i] if i < len(values) else None
    return out


def _read_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = None
    rows = []
    for i, raw in enumerate(ws.iter_rows(values_only=True), 1):
        values = list(raw)
        if i == 1:
            headers = [_cell(v) for v in values]
            continue
        data = _row_map(headers, values)
        if not _cell(data.get("접수번호")):
            continue
        if not _cell(data.get("부품명")):
            continue
        rows.append(data)
    wb.close()
    return rows


def _clean_part(row):
    name = _cell(row.get("부품명"))
    qty = _qty(row.get("수량"))
    unit_price = _price(row.get("단가"))
    return {
        "name": name,
        "spec": _cell(row.get("규격")),
        "qty": qty,
        "unit_price": unit_price,
        "subtotal": int(round(float(qty) * unit_price)),
    }


def main():
    parser = argparse.ArgumentParser(description="AS이력부품내역.xlsx DB 등록")
    parser.add_argument("xlsx", nargs="?", default=str(DEFAULT_XLSX))
    parser.add_argument("--replace", action="store_true", help="이미 부품이 있어도 엑셀 값으로 덮어씀")
    args = parser.parse_args()
    path = Path(args.xlsx)
    if not path.exists():
        print(f"파일을 찾을 수 없습니다: {path}")
        return 1

    rows = _read_rows(path)
    if not rows:
        print("엑셀에 등록할 행이 없습니다.")
        return 1

    grouped = defaultdict(list)
    for row in rows:
        grouped[_cell(row.get("접수번호"))].append(_clean_part(row))

    existing_receipts = {r.get("receipt_no") for r in m.RECEIPTS}
    updated = 0
    skipped_has_parts = 0
    skipped_missing = 0
    missing_nos = []
    price_missing = sum(1 for row in rows if not _cell(row.get("단가")))

    for receipt_no, parts in grouped.items():
        if receipt_no not in existing_receipts:
            skipped_missing += 1
            missing_nos.append(receipt_no)
            continue
        detail = m.REPORT_DETAILS.get(receipt_no)
        if detail is None:
            skipped_missing += 1
            missing_nos.append(receipt_no)
            continue
        current = [p for p in (detail.get("parts") or []) if (p.get("name") or "").strip()]
        if current and not args.replace:
            skipped_has_parts += 1
            continue
        detail["parts"] = parts
        if not detail.get("seed_tag"):
            detail["seed_tag"] = SEED_MARKER
        receipt = next((r for r in m.RECEIPTS if r.get("receipt_no") == receipt_no), None)
        as_name = (receipt or {}).get("as_staff_name") or "엑셀부품등록"
        as_id = (receipt or {}).get("as_staff") or ""
        m._stamp_history_modified(detail, {"name": as_name, "id": as_id})
        if receipt is not None:
            receipt["parts_total"] = m.v32.calc_parts_total(parts)
        updated += 1

    m._save_service_data()

    print(f"파일: {path}")
    print(f"엑셀 부품 행: {len(rows)}건 / 접수번호 {len(grouped)}건")
    print(f"부품 등록: {updated}건")
    print(f"이미 부품 있어 건너뜀: {skipped_has_parts}건")
    print(f"접수 없음 건너뜀: {skipped_missing}건")
    if missing_nos:
        print("없는 접수번호:")
        for no in missing_nos:
            print(" -", no)
    if price_missing:
        print(f"단가 비어 0원 처리: {price_missing}행")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
