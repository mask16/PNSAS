"""웰딩 장비 AS 처리 시스템 - Flask 애플리케이션"""

import os
import csv
import io
import json
import re
import sqlite3
import time
import uuid
import unicodedata
from datetime import date, datetime, timedelta
from functools import wraps
from pathlib import Path
from threading import RLock, Thread

from flask import Flask, flash, render_template, session, redirect, url_for, request, send_file, Response, jsonify
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename

from pdf_report import build_report_pdf
from email_notify import (
    load_mail_config,
    mail_configured,
    mail_error_message,
    notify_escalation_changed,
    notify_receipt_completed,
    notify_receipt_registered,
    notify_staff_assigned,
    notify_support_requested,
    notify_tech_mgmt_reported,
    send_receipt_notice,
)
import v32_rules as v32
import product_line_ext as plex

app = Flask(__name__)
app.secret_key = os.environ.get("PNS_SECRET_KEY", "pns-as-system-secret-key-change-in-production")

BASE_DIR = Path(__file__).resolve().parent
load_mail_config(BASE_DIR)
USER_DB_PATH = BASE_DIR / "users.db"
LEGACY_SERVICE_DATA_PATH = BASE_DIR / "receipts_data.json"
UPLOAD_DIR = BASE_DIR / "static" / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
ALLOWED_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp"}
SERVICE_DATA_LOCK = RLock()


def _user_db():
    conn = sqlite3.connect(USER_DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def _init_user_db():
    """계정 DB 초기화 및 최초 관리자(admin/admin) 생성."""
    with _user_db() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL UNIQUE COLLATE NOCASE,
                password_hash TEXT NOT NULL,
                display_name TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'staff',
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS teams (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                code TEXT NOT NULL UNIQUE COLLATE NOCASE,
                name TEXT NOT NULL,
                category TEXT NOT NULL DEFAULT 'domestic',
                sort_order INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS receipt_staff (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                code TEXT NOT NULL UNIQUE COLLATE NOCASE,
                name TEXT NOT NULL,
                team_code TEXT NOT NULL,
                sort_order INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS as_staff (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                code TEXT NOT NULL UNIQUE COLLATE NOCASE,
                name TEXT NOT NULL,
                team_code TEXT NOT NULL,
                sort_order INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS org_people (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                code TEXT NOT NULL UNIQUE COLLATE NOCASE,
                name TEXT NOT NULL,
                title TEXT NOT NULL DEFAULT '',
                role_note TEXT NOT NULL DEFAULT '',
                team_code TEXT NOT NULL,
                can_receipt INTEGER NOT NULL DEFAULT 1,
                can_as INTEGER NOT NULL DEFAULT 1,
                sort_order INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        team_columns = {
            row["name"] for row in conn.execute("PRAGMA table_info(teams)").fetchall()
        }
        if "note" not in team_columns:
            conn.execute("ALTER TABLE teams ADD COLUMN note TEXT NOT NULL DEFAULT ''")
        people_columns = {
            row["name"] for row in conn.execute("PRAGMA table_info(org_people)").fetchall()
        }
        if "email" not in people_columns:
            conn.execute("ALTER TABLE org_people ADD COLUMN email TEXT NOT NULL DEFAULT ''")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS receipts (
                id INTEGER PRIMARY KEY,
                receipt_no TEXT NOT NULL UNIQUE,
                receipt_date TEXT NOT NULL,
                team_id TEXT NOT NULL,
                team TEXT NOT NULL,
                customer_name TEXT NOT NULL,
                customer_contact TEXT,
                equipment_name TEXT NOT NULL,
                symptoms TEXT NOT NULL,
                as_staff TEXT,
                as_staff_name TEXT,
                receipt_staff TEXT,
                receipt_staff_name TEXT,
                status TEXT NOT NULL DEFAULT 'received',
                charge_type TEXT NOT NULL DEFAULT 'free',
                billing_cost TEXT,
                completed_date TEXT,
                created_at TEXT NOT NULL,
                photos_json TEXT NOT NULL DEFAULT '[]',
                phone TEXT,
                ship_date TEXT,
                support_quality INTEGER NOT NULL DEFAULT 0,
                support_tech INTEGER NOT NULL DEFAULT 0,
                created_by_user_id INTEGER,
                created_by_username TEXT
            )
            """
        )
        receipt_columns = {
            row["name"] for row in conn.execute("PRAGMA table_info(receipts)").fetchall()
        }
        if "created_by_user_id" not in receipt_columns:
            conn.execute("ALTER TABLE receipts ADD COLUMN created_by_user_id INTEGER")
        if "created_by_username" not in receipt_columns:
            conn.execute("ALTER TABLE receipts ADD COLUMN created_by_username TEXT")
        # 최고기술경영실: 심각도·개입·보상 협의
        for col, ddl in (
            ("severity", "TEXT NOT NULL DEFAULT 'medium'"),
            ("intervention_status", "TEXT NOT NULL DEFAULT 'none'"),
            ("compensation_status", "TEXT NOT NULL DEFAULT 'none'"),
            ("intervention_note", "TEXT NOT NULL DEFAULT ''"),
            ("compensation_note", "TEXT NOT NULL DEFAULT ''"),
            ("intervention_by", "TEXT NOT NULL DEFAULT ''"),
            ("intervention_at", "TEXT"),
        ):
            if col not in receipt_columns:
                conn.execute(f"ALTER TABLE receipts ADD COLUMN {col} {ddl}")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS report_details (
                receipt_no TEXT PRIMARY KEY,
                data_json TEXT NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS app_meta (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL DEFAULT ''
            )
            """
        )

        def _team_marked_removed(code):
            row = conn.execute(
                "SELECT 1 FROM app_meta WHERE key = ?",
                (f"team_removed:{code}",),
            ).fetchone()
            return bool(row)

        default_teams = [
            ("domestic", "내수영업팀", "domestic", 10, ""),
            ("affiliate", "관계사영업팀", "affiliate", 20, ""),
            ("export", "수출영업팀", "domestic", 30, "접수 전담 · AS는 내수영업팀이 처리"),
            ("field_support", "현장 지원 부서", "service", 40, "지원요청 건만 조회 · 공용계정"),
            ("sales_head", "영업부 부서장", "service", 45, "접수+전권의뢰 · 이력관리"),
            ("executive", "경영진 · 임원", "service", 50, "전체 현황 조회 전용"),
            ("tech_mgmt", "최고기술경영실", "service", 55, "전 팀 AS 모니터링 · 고심각도 개입·보상 협의"),
            ("service", "AS팀", "service", 60, ""),
        ]
        team_count = conn.execute("SELECT COUNT(*) FROM teams").fetchone()[0]
        if team_count == 0:
            # 최초 설치: 기본 부서 일괄 생성
            conn.executemany(
                "INSERT OR IGNORE INTO teams (code, name, category, sort_order, note) VALUES (?, ?, ?, ?, ?)",
                default_teams,
            )
        else:
            # 업그레이드용: 삭제 기록이 없는 시스템 부서만 보강 (사용자 삭제분 복구 금지)
            ensure_codes = {
                "domestic",
                "affiliate",
                "export",
                "field_support",
                "executive",
                "tech_mgmt",
            }
            for row in default_teams:
                code = row[0]
                if code not in ensure_codes:
                    continue
                if _team_marked_removed(code):
                    continue
                conn.execute(
                    "INSERT OR IGNORE INTO teams (code, name, category, sort_order, note) VALUES (?, ?, ?, ?, ?)",
                    row,
                )
        # 영업부 부서장·AS팀(service)은 관리자가 삭제하면 기동 시 다시 만들지 않음
        # 기존 영업부 부서장 권한 마이그레이션 (부서가 실제로 존재할 때만)
        if (
            conn.execute(
                "SELECT 1 FROM teams WHERE code = 'sales_head' COLLATE NOCASE"
            ).fetchone()
            and not _team_marked_removed("sales_head")
        ):
            conn.execute(
                "UPDATE org_people SET team_code='sales_head', can_receipt=1, can_as=1 "
                "WHERE code='sales_head' OR name LIKE '%영업부 부서장%'"
            )
            conn.execute(
                "UPDATE org_people SET can_receipt=1, can_as=1 "
                "WHERE team_code='sales_head'"
            )
        # 대표이사·총괄: 영업부 부서장과 동등한 접수·이력 권한
        conn.execute(
            "UPDATE org_people SET can_receipt=1, can_as=1, role_note='전체 접수·의뢰 총괄' "
            "WHERE code IN ('exec_chief', 'exec_ceo') OR name IN ('대표이사', '총괄')"
        )
        conn.execute(
            "UPDATE org_people SET can_receipt=0, can_as=1 WHERE code='jang_hb'"
        )
        conn.execute(
            "UPDATE org_people SET can_receipt=1, can_as=0 WHERE code='lee_mj'"
        )
        receipt_columns = {
            row["name"] for row in conn.execute("PRAGMA table_info(receipts)").fetchall()
        }
        for col, ddl in (
            ("escalation_level", "TEXT NOT NULL DEFAULT 'normal'"),
            ("compensation_amount", "TEXT NOT NULL DEFAULT ''"),
            ("intervention_log", "TEXT NOT NULL DEFAULT '[]'"),
            ("billing_margin_rate", "TEXT"),
            ("billing_manual", "INTEGER NOT NULL DEFAULT 0"),
            ("billing_set_by", "TEXT NOT NULL DEFAULT ''"),
        ):
            if col not in receipt_columns:
                conn.execute(f"ALTER TABLE receipts ADD COLUMN {col} {ddl}")
        if "notice_sent" not in receipt_columns:
            conn.execute(
                "ALTER TABLE receipts ADD COLUMN notice_sent INTEGER NOT NULL DEFAULT 0"
            )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS notifications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_code TEXT NOT NULL DEFAULT '',
                team_code TEXT NOT NULL DEFAULT '',
                title TEXT NOT NULL,
                body TEXT NOT NULL DEFAULT '',
                link TEXT NOT NULL DEFAULT '',
                receipt_no TEXT NOT NULL DEFAULT '',
                is_read INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            "UPDATE teams SET note = ? WHERE code = ? AND IFNULL(note, '') = ''",
            ("접수 전담 · AS는 내수영업팀이 처리", "export"),
        )
        conn.execute(
            "UPDATE teams SET note = ? WHERE code = ? AND IFNULL(note, '') = ''",
            ("지원요청 건만 조회 · 공용계정", "field_support"),
        )
        conn.execute(
            "UPDATE teams SET note = ? WHERE code = ? AND IFNULL(note, '') = ''",
            ("전체 현황 조회 전용", "executive"),
        )
        conn.execute(
            "UPDATE teams SET note = ? WHERE code = ? AND (IFNULL(note, '') = '' OR note LIKE ?)",
            (
                "전 팀 AS 모니터링 · 고심각도 개입·보상 협의",
                "tech_mgmt",
                "%모니터링%",
            ),
        )
        # 기존에 만든 '최고기술경영' 부서(dept001 등)도 동일 권한으로 취급되도록 note 보강
        conn.execute(
            """
            UPDATE teams
            SET note = ?
            WHERE (name LIKE '%최고기술경영%' OR code = 'dept001')
              AND IFNULL(note, '') NOT LIKE '%개입%'
            """,
            ("전 팀 AS 모니터링 · 고심각도 개입·보상 협의",),
        )
        people_count = conn.execute("SELECT COUNT(*) FROM org_people").fetchone()[0]
        if people_count == 0:
            seed_people = [
                ("choi_yh", "최영호", "팀장", "접수·AS", "domestic", 1, 1, 10),
                ("park_ys", "박용수", "", "접수·AS", "domestic", 1, 1, 20),
                ("jang_hb", "장호빈", "", "AS 전담", "domestic", 0, 1, 30),
                ("kim_jh_aff", "김정훈", "팀장", "접수·AS", "affiliate", 1, 1, 10),
                ("park_kr", "박경률", "", "접수·AS", "affiliate", 1, 1, 20),
                ("jang_ys", "장유식", "", "접수·AS", "affiliate", 1, 1, 30),
                ("lee_jg", "이정건", "", "접수·AS", "affiliate", 1, 1, 40),
                ("kim_ys_ex", "김윤식", "팀장", "접수·AS 겸용", "export", 1, 1, 10),
                ("lee_mj", "이명지", "", "접수", "export", 1, 0, 20),
                ("quality_acc", "품질팀 계정", "", "", "field_support", 0, 0, 10),
                ("tech_acc", "기술부 계정", "", "", "field_support", 0, 0, 20),
                ("sales_head", "영업부 부서장", "", "전체 접수·의뢰 총괄", "sales_head", 1, 1, 10),
                ("tech_group", "기술부 그룹장", "", "", "executive", 0, 0, 20),
                ("tech_lab", "기술연구소장", "", "", "executive", 0, 0, 30),
                ("exec_chief", "총괄", "", "전체 접수·의뢰 총괄", "executive", 1, 1, 40),
                ("exec_ceo", "대표이사", "", "전체 접수·의뢰 총괄", "executive", 1, 1, 50),
                ("tech_mgmt_office", "최고기술경영실", "", "전 팀 모니터링·개입", "tech_mgmt", 0, 0, 10),
            ]
            # 기존 기준정보 담당자도 함께 이전
            for row in conn.execute(
                "SELECT code, name, team_code, sort_order FROM receipt_staff ORDER BY sort_order, id"
            ).fetchall():
                if any(p[0] == row["code"] for p in seed_people):
                    continue
                seed_people.append(
                    (row["code"], row["name"], "", "접수", row["team_code"], 1, 0, row["sort_order"] + 100)
                )
            for row in conn.execute(
                "SELECT code, name, team_code, sort_order FROM as_staff ORDER BY sort_order, id"
            ).fetchall():
                exists = next((i for i, p in enumerate(seed_people) if p[0] == row["code"]), None)
                if exists is not None:
                    person = list(seed_people[exists])
                    person[6] = 1
                    if not person[3]:
                        person[3] = "AS"
                    seed_people[exists] = tuple(person)
                    continue
                seed_people.append(
                    (row["code"], row["name"], "", "AS", row["team_code"], 0, 1, row["sort_order"] + 200)
                )
            conn.executemany(
                """
                INSERT OR IGNORE INTO org_people
                    (code, name, title, role_note, team_code, can_receipt, can_as, sort_order)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                seed_people,
            )
        conn.executemany(
            """
            INSERT OR IGNORE INTO receipt_staff (code, name, team_code, sort_order)
            VALUES (?, ?, ?, ?)
            """,
            [
                ("park_ys", "박용수", "domestic", 10),
                ("kim_jh", "김정호", "domestic", 20),
                ("lee_sm", "이승민", "export", 30),
            ],
        )
        conn.executemany(
            """
            INSERT OR IGNORE INTO as_staff (code, name, team_code, sort_order)
            VALUES (?, ?, ?, ?)
            """,
            [
                ("choi_as", "최기술", "service", 10),
                ("jung_as", "정수리", "service", 20),
                ("han_as", "한현장", "service", 30),
            ],
        )
        conn.executemany(
            """
            INSERT OR IGNORE INTO org_people
                (code, name, title, role_note, team_code, can_receipt, can_as, sort_order)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    "tech_mgmt_office",
                    "최고기술경영실",
                    "",
                    "전 팀 모니터링·개입",
                    "tech_mgmt",
                    0,
                    0,
                    10,
                ),
                ("exec_chief", "총괄", "", "전체 접수·의뢰 총괄", "executive", 1, 1, 40),
                ("exec_ceo", "대표이사", "", "전체 접수·의뢰 총괄", "executive", 1, 1, 50),
            ],
        )
        admin = conn.execute(
            "SELECT id FROM users WHERE username = ? COLLATE NOCASE", ("admin",)
        ).fetchone()
        if not admin:
            now = datetime.now().isoformat(timespec="seconds")
            conn.execute(
                """
                INSERT INTO users
                    (username, password_hash, display_name, role, active, created_at, updated_at)
                VALUES (?, ?, ?, ?, 1, ?, ?)
                """,
                ("admin", generate_password_hash("admin"), "관리자", "admin", now, now),
            )
        admin = conn.execute(
            "SELECT id, username FROM users WHERE username = ? COLLATE NOCASE", ("admin",)
        ).fetchone()
        conn.execute(
            """
            UPDATE receipts
            SET created_by_user_id = ?, created_by_username = ?
            WHERE created_by_user_id IS NULL
            """,
            (admin["id"], admin["username"]),
        )
        plex.ensure_schema_and_org(conn)


def _admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            flash("로그인이 필요합니다.", "error")
            return redirect(url_for("home", login=1, next=request.path))
        if session.get("user_role") != "admin":
            flash("관리자만 접근할 수 있습니다.", "error")
            return redirect(url_for("home"))
        return view(*args, **kwargs)

    return wrapped


def _product_home_endpoint():
    return "rods_home" if session.get("product_line") == "rods" else "home"


def _is_rods():
    return plex.is_rods()


def _enter_product_line(line):
    prev = session.get("product_line")
    if prev and prev != line and session.get("user_id"):
        session.clear()
    session["product_line"] = line


def _line_receipts(source=None):
    rows = source if source is not None else RECEIPTS
    return [r for r in rows if plex.receipt_matches_session(r)]


def _ui_labels():
    return plex.ui_labels()


SALES_RECEIPT_STAFF_TEAMS = ("domestic", "affiliate", "export")


def _teams_for_sales():
    """영업접수 AS 대응팀. 제품군 선택 시점의 세션 기준으로만 구성한다."""
    by_id = {t["id"]: t for t in _get_teams()}
    if _is_rods():
        team = by_id.get(plex.RODS_AS_TEAM)
        if not team:
            return []
        return [{**team, "label": "AS담당관"}]
    out = []
    for code in SALES_RECEIPT_STAFF_TEAMS:
        team = by_id.get(code)
        if team:
            out.append({**team, "label": team["name"]})
    return out


EXPORT_RESPONSE_CHOOSER_CODES = frozenset({"kim_ys_ex", "lee_mj"})
EXPORT_RESPONSE_CHOOSER_NAMES = ("김윤식", "이명지")


def _is_export_response_chooser(person=None):
    """수출영업팀 정규직(김윤식·이명지): AS 대응팀(수출/내수) 선택 가능."""
    if not person or _is_rods():
        return False
    try:
        if session.get("user_role") == "admin":
            return False
    except RuntimeError:
        pass
    try:
        if _is_sales_head(person):
            return False
    except NameError:
        if person.get("id") == "sales_head" or person.get("team_code") == "sales_head":
            return False
    code = (person.get("id") or person.get("code") or "").strip()
    if code in EXPORT_RESPONSE_CHOOSER_CODES:
        return True
    if (person.get("team_code") or "") != "export":
        return False
    name = (person.get("name") or "").strip()
    return any(name == n or name.startswith(n) for n in EXPORT_RESPONSE_CHOOSER_NAMES)


def _sales_form_teams(person):
    """접수 폼의 AS 대응팀 목록과 선택 가능 여부."""
    all_teams = _teams_for_sales()
    by_id = {t["id"]: t for t in all_teams}
    if _is_rods():
        return all_teams, False
    try:
        if session.get("user_role") == "admin":
            return all_teams, True
    except RuntimeError:
        pass
    try:
        if _is_sales_head(person):
            return all_teams, True
    except NameError:
        if person and (
            person.get("id") == "sales_head" or person.get("team_code") == "sales_head"
        ):
            return all_teams, True
    if _is_office_person(person):
        return all_teams, True
    if _is_export_response_chooser(person):
        out = []
        for code in ("export", "domestic"):
            team = by_id.get(code)
            if team:
                out.append(team)
        return out, True
    team_code = (person or {}).get("team_code") or ""
    if team_code in by_id:
        return [by_id[team_code]], False
    if all_teams:
        return [all_teams[0]], False
    return [], False


def _normalize_sales_team_id(person, requested, existing=None):
    """폼에서 온 팀을 계정별 허용 범위로 고정한다."""
    if _is_rods():
        return plex.RODS_AS_TEAM
    teams, _enabled = _sales_form_teams(person)
    allowed = {t["id"] for t in teams}
    req = (requested or "").strip()
    if req in allowed:
        return req
    if existing and (existing.get("team_id") or "") in allowed:
        return existing.get("team_id") or ""
    if teams:
        return teams[0]["id"]
    return req or "domestic"


def _apply_export_response_as_staff(person, team_id, as_staff_id=""):
    """수출 정규직 접수: 수출→김윤식, 내수→미지정 또는 장호빈만."""
    if not _is_export_response_chooser(person):
        return as_staff_id
    if (team_id or "") == "export":
        return EXPORT_AS_SELF_CODE
    if (team_id or "") == "domestic":
        as_id = (as_staff_id or "").strip()
        if as_id == EXPORT_AS_HANDLER_CODE:
            return EXPORT_AS_HANDLER_CODE
        return ""
    return as_staff_id


def _can_choose_receipt_staff(person=None):
    """영업부 부서장·관리자만 접수 담당자를 선택한다."""
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    try:
        return _is_sales_head(person)
    except NameError:
        return bool(
            person
            and (
                person.get("id") == "sales_head"
                or person.get("team_code") == "sales_head"
            )
        )


def _receipt_staff_choice_people():
    """접수 담당자 후보: AS 대응팀과 무관한 영업 3팀(또는 용접봉 AS담당관)."""
    teams = {plex.RODS_AS_TEAM} if _is_rods() else set(SALES_RECEIPT_STAFF_TEAMS)
    return [
        p
        for p in _get_org_people()
        if p.get("can_receipt")
        and plex.person_visible_on_line(p)
        and (p.get("team_code") or "") in teams
    ]


def _receipt_staff_choice_tree():
    allowed_ids = {p["id"] for p in _receipt_staff_choice_people()}
    tree = []
    for dept in _get_org_tree():
        members = [p for p in (dept.get("people") or []) if p.get("id") in allowed_ids]
        if members:
            tree.append({**dept, "people": members})
    return tree


def _locked_receipt_staff(person):
    if not person:
        return {"id": "", "name": session.get("user_name") or session.get("username") or ""}
    return {
        "id": person.get("id") or "",
        "name": person.get("display_name") or person.get("name") or "",
    }


def _resolve_receipt_staff(person, requested_id="", typed_name=""):
    """저장 시 접수 담당자. 일반 담당자는 로그인 본인 고정, 부서장만 3팀에서 지정."""
    if _is_office_person(person):
        return "", (typed_name or "").strip()
    if not _can_choose_receipt_staff(person):
        locked = _locked_receipt_staff(person)
        return locked["id"], locked["name"]
    allowed = {p["id"]: p for p in _receipt_staff_choice_people()}
    code = (requested_id or "").strip()
    picked = allowed.get(code)
    if not picked:
        return "", ""
    return picked["id"], picked.get("display_name") or picked.get("name") or ""


def _teams_for_history():
    teams = [
        t
        for t in _get_teams()
        if t["id"] not in ("executive", "sales_head", "tech_mgmt", "field_support", plex.SALES_OFFICE_TEAM)
        and not _is_tech_mgmt_team_code(t["id"])
    ]
    if _is_rods():
        teams = [t for t in teams if t["id"] == plex.RODS_AS_TEAM]
    else:
        teams = [t for t in teams if t["id"] != plex.RODS_AS_TEAM]
    return teams


def _default_affiliation(receipt=None):
    if receipt is not None:
        if plex.receipt_is_rod(receipt):
            return "AS담당관"
        return "내수영업팀"
    return "AS담당관" if _is_rods() else "내수영업팀"


def _affiliation_options(receipt=None):
    """이력관리 소속 드롭다운. AS 대응팀과 별개로 제품군에 맞는 표시값만 둔다."""
    if (receipt is not None and plex.receipt_is_rod(receipt)) or (
        receipt is None and _is_rods()
    ):
        return [
            {"value": "AS담당관", "label": "AS담당관"},
            {"value": "기술연구소", "label": "기술연구소"},
        ]
    by_id = {t["id"]: t for t in _get_teams()}
    out = []
    for code in ("domestic", "affiliate", "export"):
        team = by_id.get(code)
        if team:
            out.append({"value": team["name"], "label": team["name"]})
    if not out:
        out = [
            {"value": "내수영업팀", "label": "내수영업팀"},
            {"value": "관계사영업팀", "label": "관계사영업팀"},
            {"value": "수출영업팀", "label": "수출영업팀"},
        ]
    return out


def _normalize_affiliation(value, receipt=None):
    text = (value or "").strip()
    is_rod = plex.receipt_is_rod(receipt) if receipt is not None else _is_rods()
    options = {opt["value"] for opt in _affiliation_options(receipt)}
    if is_rod:
        compact = text.replace(" ", "")
        if compact in {"AS담당관실", "AS담당관", "rods_as", plex.RODS_AS_TEAM}:
            return "AS담당관"
        if "기술연구" in compact or compact in {"field_support", "rods_lab", "기술연구소계정"}:
            return "기술연구소"
        if text in options:
            return text
        return _default_affiliation(receipt)
    if text in options:
        return text
    if text in {"domestic", "내수"}:
        return "내수영업팀"
    if text in {"affiliate", "관계사"}:
        return "관계사영업팀"
    if text in {"export", "수출"}:
        return "수출영업팀"
    return text or _default_affiliation(receipt)


def _flash_once(message, category="message"):
    """같은 카테고리·문구의 flash가 세션에 이미 있으면 추가하지 않는다."""
    pending = session.get("_flashes") or []
    for item in pending:
        if (
            isinstance(item, (list, tuple))
            and len(item) >= 2
            and item[0] == category
            and item[1] == message
        ):
            return
    flash(message, category)


def _login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            flash("로그인 후 이용할 수 있습니다.", "error")
            return redirect(url_for(_product_home_endpoint(), login=1, next=request.path))
        return view(*args, **kwargs)

    return wrapped


def _executive_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            flash("로그인 후 이용할 수 있습니다.", "error")
            return redirect(url_for("home", login=1, next=request.path))
        person = _get_person_for_session()
        is_team_lead = False
        try:
            is_team_lead = bool(_is_sales_team_leader(person))
        except Exception:
            is_team_lead = False
        if (
            not _is_executive_person(person)
            and not _is_tech_mgmt_person(person)
            and not (person and (person.get("team_code") == "sales_head" or person.get("id") == "sales_head"))
            and not is_team_lead
            and session.get("user_role") != "admin"
        ):
            flash(
                "경영진·임원, 영업부 부서장·팀장 또는 최고기술경영실 계정만 이용할 수 있습니다.",
                "error",
            )
            return redirect(url_for("home"))
        return view(*args, **kwargs)

    return wrapped


def _tech_mgmt_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            flash("로그인 후 이용할 수 있습니다.", "error")
            return redirect(url_for("home", login=1, next=request.path))
        person = _get_person_for_session()
        if not _is_tech_mgmt_person(person) and session.get("user_role") != "admin":
            flash("최고기술경영실 계정만 이용할 수 있습니다.", "error")
            return redirect(url_for("home"))
        return view(*args, **kwargs)

    return wrapped


def _service_write_locked(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        with SERVICE_DATA_LOCK:
            return view(*args, **kwargs)

    return wrapped


def _valid_admin_csrf():
    return bool(
        request.form.get("csrf_token")
        and request.form.get("csrf_token") == session.get("csrf_token")
    )


def _valid_session_csrf():
    return _valid_admin_csrf()


def _safe_redirect_path(raw):
    path = (raw or "").strip()
    if path.startswith("/") and not path.startswith("//"):
        return path
    return url_for(_product_home_endpoint())


def _redirect_with_password_modal(next_path):
    base = _safe_redirect_path(next_path)
    joiner = "&" if "?" in base else "?"
    return redirect(f"{base}{joiner}password=1")


_init_user_db()


def _get_teams():
    with _user_db() as conn:
        rows = conn.execute(
            """
            SELECT id, code, name, category, IFNULL(note, '') AS note
            FROM teams
            ORDER BY sort_order, id
            """
        ).fetchall()
    return [
        {
            "db_id": row["id"],
            "id": row["code"],
            "name": row["name"],
            "category": row["category"],
            "note": row["note"] or "",
            "label": f"{row['name']} ({row['note']})" if row["note"] else row["name"],
        }
        for row in rows
    ]


def _parse_person_emails(raw):
    """세미콜론·쉼표·줄바꿈 구분 다중 이메일 파싱."""
    if not raw:
        return []
    if isinstance(raw, (list, tuple)):
        parts = raw
    else:
        parts = re.split(r"[;,\n]+", str(raw))
    emails = []
    seen = set()
    for part in parts:
        email = part.strip()
        if not email or "@" not in email:
            continue
        key = email.casefold()
        if key in seen:
            continue
        seen.add(key)
        emails.append(email)
    return emails


def _serialize_person_emails(raw):
    return ";".join(_parse_person_emails(raw))


def _collect_person_emails_from_form(form):
    items = form.getlist("emails")
    if not items:
        single = (form.get("email") or "").strip()
        if single:
            items = [single]
    return _serialize_person_emails(items)


def _get_org_people():
    with _user_db() as conn:
        rows = conn.execute(
            """
            SELECT p.id, p.code, p.name, p.title, p.role_note, p.team_code,
                   p.can_receipt, p.can_as, p.sort_order, IFNULL(p.email, '') AS email,
                   COALESCE(t.name, p.team_code) AS team_name,
                   IFNULL(t.note, '') AS team_note,
                   IFNULL(p.product_line, 'equipment') AS product_line
            FROM org_people p
            LEFT JOIN teams t ON t.code = p.team_code COLLATE NOCASE
            ORDER BY COALESCE(t.sort_order, 9999), p.sort_order, p.id
            """
        ).fetchall()
    people = []
    for row in rows:
        display_name = row["name"]
        if row["title"]:
            display_name = f"{row['name']} {row['title']}"
        if row["role_note"]:
            display_name = f"{display_name} ({row['role_note']})"
        people.append(
            {
                "db_id": row["id"],
                "id": row["code"],
                "name": row["name"],
                "title": row["title"] or "",
                "role_note": row["role_note"] or "",
                "team_code": row["team_code"],
                "team": row["team_name"],
                "team_note": row["team_note"] or "",
                "can_receipt": bool(row["can_receipt"]),
                "can_as": bool(row["can_as"]),
                "email": row["email"] or "",
                "emails": _parse_person_emails(row["email"]),
                "product_line": (row["product_line"] if "product_line" in row.keys() else "equipment")
                or "equipment",
                "display_name": display_name,
            }
        )
    return people


def _get_staff_emails(staff_code, email_map=None):
    if not staff_code:
        return []
    if email_map is not None:
        raw = email_map.get((staff_code or "").casefold()) or ""
    else:
        with _user_db() as conn:
            row = conn.execute(
                "SELECT email FROM org_people WHERE code = ? COLLATE NOCASE",
                (staff_code,),
            ).fetchone()
        raw = (row["email"] or "").strip() if row else ""
    return _parse_person_emails(raw)


def _get_staff_email(staff_code, email_map=None):
    emails = _get_staff_emails(staff_code, email_map=email_map)
    if not emails:
        return ""
    return _serialize_person_emails(emails)


def _build_staff_email_map():
    """코드→이메일(세미콜론 구분) 일괄 로드 (홈/목록 N+1 방지)."""
    with _user_db() as conn:
        rows = conn.execute(
            "SELECT code, IFNULL(email, '') AS email FROM org_people"
        ).fetchall()
    out = {}
    for row in rows:
        code = (row["code"] or "").strip()
        emails = _parse_person_emails(row["email"])
        if code and emails:
            out[code.casefold()] = _serialize_person_emails(emails)
    return out


def _receipt_staff_emails(receipt, email_map=None):
    """계정관리에 등록된 접수·AS 담당자 이메일만 반환(중복 제거)."""
    if email_map is None:
        email_map = _build_staff_email_map()
    emails = []
    seen = set()
    for code in (receipt.get("as_staff"), receipt.get("receipt_staff")):
        for email in _get_staff_emails(code, email_map=email_map):
            key = email.casefold()
            if key in seen:
                continue
            seen.add(key)
            emails.append(email)
    return emails


def _receipt_notice_sent(receipt):
    try:
        return int((receipt or {}).get("notice_sent") or 0) == 1
    except (TypeError, ValueError):
        return bool((receipt or {}).get("notice_sent"))


def _mark_receipt_notice_sent(receipt):
    if not receipt or _receipt_notice_sent(receipt):
        return
    receipt["notice_sent"] = 1
    _save_service_data()


def _mail_base_url():
    base = (os.environ.get("MAIL_BASE_URL") or "").strip()
    if base:
        return base.rstrip("/")
    try:
        return request.host_url.rstrip("/")
    except RuntimeError:
        return ""


def _get_org_tree(product_filtered=True):
    teams = _get_teams()
    people = _get_org_people()
    if product_filtered:
        people = [p for p in people if plex.person_visible_on_line(p)]
        if _is_rods():
            allow = {
                plex.RODS_AS_TEAM,
                plex.SALES_OFFICE_TEAM,
                "field_support",
                "tech_mgmt",
            }
            teams = [t for t in teams if t["id"] in allow]
        else:
            teams = [t for t in teams if t["id"] != plex.RODS_AS_TEAM]
    team_codes = {team["id"] for team in teams}
    tree = []
    for team in teams:
        members = [p for p in people if p["team_code"] == team["id"]]
        tree.append({**team, "people": members, "is_unassigned": False})
    # 부서가 삭제된 인원 등 소속 없는 인원
    unassigned = [
        p for p in people if (p.get("team_code") or "") not in team_codes
    ]
    if unassigned:
        tree.append(
            {
                "db_id": None,
                "id": "",
                "name": "미배정",
                "category": "",
                "note": "부서 삭제 등으로 소속이 없는 인원",
                "label": "미배정",
                "people": unassigned,
                "is_unassigned": True,
            }
        )
    return tree


def _get_all_users_index():
    """users 테이블 전체를 username 기준으로 조회."""
    with _user_db() as conn:
        rows = conn.execute(
            """
            SELECT id, username, display_name, role, active, created_at, updated_at
            FROM users
            ORDER BY role = 'admin' DESC, username COLLATE NOCASE
            """
        ).fetchall()
    by_username = {}
    indexed = []
    for row in rows:
        item = {
            "user_id": row["id"],
            "login_username": row["username"],
            "display_name": (row["display_name"] or row["username"] or "").strip(),
            "account_role": row["role"],
            "account_active": bool(row["active"]),
            "account_created_at": row["created_at"] or "",
            "account_updated_at": row["updated_at"] or "",
        }
        indexed.append(item)
        by_username[(row["username"] or "").casefold()] = item
    return indexed, by_username


def _person_with_account(person, by_username):
    uname = (person.get("name") or "").casefold()
    acc = by_username.get(uname)
    base = {
        **person,
        "user_id": None,
        "login_username": person.get("name") or "",
        "account_role": None,
        "account_active": False,
        "account_created_at": "",
        "account_updated_at": "",
        "has_login": False,
        "is_system_account": False,
    }
    if not acc:
        return base
    return {
        **base,
        **acc,
        "has_login": True,
    }


def _pseudo_account_person(user_info):
    display = user_info.get("display_name") or user_info.get("login_username") or ""
    login = user_info.get("login_username") or display
    return {
        "db_id": None,
        "id": f"user_{user_info['user_id']}",
        "name": display,
        "title": "",
        "role_note": "관리자" if user_info.get("account_role") == "admin" else "미연결 계정",
        "team_code": "",
        "team": "",
        "team_note": "",
        "can_receipt": False,
        "can_as": False,
        "email": "",
        "emails": [],
        "display_name": display,
        "product_line": "equipment",
        "user_id": user_info["user_id"],
        "login_username": login,
        "account_role": user_info.get("account_role"),
        "account_active": user_info.get("account_active"),
        "account_created_at": user_info.get("account_created_at") or "",
        "account_updated_at": user_info.get("account_updated_at") or "",
        "has_login": True,
        "is_system_account": True,
    }


def _get_org_tree_for_admin():
    """관리자 화면: 부서·인원 목록에 로그인 계정 정보를 통합."""
    tree = _get_org_tree(product_filtered=False)
    all_users, by_username = _get_all_users_index()
    matched_user_ids = set()
    enriched_tree = []

    for dept in tree:
        members = []
        for person in dept.get("people") or []:
            enriched = _person_with_account(person, by_username)
            if enriched.get("user_id"):
                matched_user_ids.add(enriched["user_id"])
            members.append(enriched)
        enriched_tree.append({**dept, "people": members})

    extras = [
        _pseudo_account_person(u) for u in all_users if u["user_id"] not in matched_user_ids
    ]
    admins = [p for p in extras if p.get("account_role") == "admin"]
    orphans = [p for p in extras if p.get("account_role") != "admin"]

    prefix = []
    if admins:
        prefix.append(
            {
                "db_id": None,
                "id": "_system_admin",
                "name": "시스템 관리자",
                "category": "",
                "note": "조직 인원 외 관리자 로그인",
                "label": "시스템 관리자",
                "people": admins,
                "is_unassigned": False,
                "is_system": True,
            }
        )

    if orphans:
        unassigned = next((d for d in enriched_tree if d.get("is_unassigned")), None)
        if unassigned:
            unassigned["people"] = list(unassigned.get("people") or []) + orphans
            if orphans:
                unassigned["note"] = "부서 미배정 · 조직에 없는 로그인 계정"
        else:
            enriched_tree.append(
                {
                    "db_id": None,
                    "id": "",
                    "name": "미배정",
                    "category": "",
                    "note": "부서 미배정 · 조직에 없는 로그인 계정",
                    "label": "미배정",
                    "people": orphans,
                    "is_unassigned": True,
                }
            )

    return prefix + enriched_tree, len(all_users)


def _get_login_dropdown():
    """로그인 모달용 부서·담당자 선택 목록."""
    groups = []
    for dept in _get_org_tree():
        members = []
        for person in dept.get("people") or []:
            username = (person.get("name") or "").strip()
            if not username:
                continue
            title = (person.get("title") or "").strip()
            label = f"{username} {title}".strip() if title else username
            members.append({"username": username, "label": label})
        members.sort(key=lambda item: item["label"])
        if members:
            groups.append(
                {
                    "team_code": dept.get("id") or "",
                    "team_name": dept.get("name") or dept.get("id") or "미배정",
                    "members": members,
                }
            )

    with _user_db() as conn:
        admins = conn.execute(
            """
            SELECT username, display_name
            FROM users
            WHERE role = 'admin' AND active = 1
            ORDER BY username COLLATE NOCASE
            """
        ).fetchall()
    if admins:
        groups.append(
            {
                "team_code": "admin",
                "team_name": "시스템 관리",
                "members": [
                    {
                        "username": row["username"],
                        "label": (row["display_name"] or row["username"]).strip(),
                    }
                    for row in admins
                ],
            }
        )
    return groups


def _get_staff_master(table_name):
    if table_name not in ("receipt_staff", "as_staff"):
        return []
    flag = "can_receipt" if table_name == "receipt_staff" else "can_as"
    people = _get_org_people()
    return [
        {
            "db_id": person["db_id"],
            "id": person["id"],
            "name": person["display_name"],
            "team_code": person["team_code"],
            "team": person["team"],
        }
        for person in people
        if person[flag] and plex.person_visible_on_line(person)
    ]


def _team_name(team_code):
    with _user_db() as conn:
        row = conn.execute(
            "SELECT name FROM teams WHERE code = ? COLLATE NOCASE", (team_code,)
        ).fetchone()
    return row["name"] if row else team_code


def _staff_name(table_name, staff_code):
    if not staff_code:
        return ""
    people = {p["id"]: p["display_name"] for p in _get_org_people()}
    if staff_code in people:
        return people[staff_code]
    if table_name not in ("receipt_staff", "as_staff"):
        return ""
    with _user_db() as conn:
        row = conn.execute(
            f"SELECT name FROM {table_name} WHERE code = ? COLLATE NOCASE", (staff_code,)
        ).fetchone()
    return row["name"] if row else ""


def _next_master_code(conn, table_name, prefix):
    """화면 입력 없이 사용할 내부 기준정보 코드를 자동 생성."""
    if table_name not in ("teams", "receipt_staff", "as_staff", "org_people"):
        raise ValueError("잘못된 기준정보 테이블입니다.")
    seq = 1
    while True:
        code = f"{prefix}{seq:03d}"
        exists = conn.execute(
            f"SELECT 1 FROM {table_name} WHERE code = ? COLLATE NOCASE", (code,)
        ).fetchone()
        if not exists:
            return code
        seq += 1


def _person_login_credentials(name):
    username = (name or "").strip()
    return username, "1234"


def _ensure_person_login(conn, name, old_name=None, password=None):
    """소속 인원 로그인 계정 생성/동기화. ID=이름, 기본 PW=1234.

    password가 주어지면 해당 값으로 설정한다.
    기존 계정은 비밀번호를 유지하고, 신규 생성·이름 변경 시에만 기본 비밀번호를 쓴다.
    """
    username, default_password = _person_login_credentials(name)
    if not username:
        return False

    now = datetime.now().isoformat(timespec="seconds")
    old_username = (old_name or "").strip()
    custom_password = (password or "").strip() or None

    def _apply_password_update(user_id, new_password):
        conn.execute(
            """
            UPDATE users
            SET username = ?, display_name = ?, password_hash = ?,
                role = 'staff', active = 1, updated_at = ?
            WHERE id = ?
            """,
            (username, name, generate_password_hash(new_password), now, user_id),
        )

    def _apply_profile_update(user_id):
        conn.execute(
            """
            UPDATE users
            SET username = ?, display_name = ?, role = 'staff',
                active = 1, updated_at = ?
            WHERE id = ?
            """,
            (username, name, now, user_id),
        )

    if old_username and old_username.casefold() != username.casefold():
        old_user = conn.execute(
            "SELECT id, role FROM users WHERE username = ? COLLATE NOCASE",
            (old_username,),
        ).fetchone()
        if old_user and old_user["role"] != "admin":
            try:
                # 이름(로그인 ID) 변경 시: 지정 비밀번호 또는 새 기본 비밀번호
                _apply_password_update(
                    old_user["id"], custom_password or default_password
                )
                return True
            except sqlite3.IntegrityError:
                pass

    existing = conn.execute(
        "SELECT id, role FROM users WHERE username = ? COLLATE NOCASE",
        (username,),
    ).fetchone()
    if existing:
        if existing["role"] == "admin":
            return False
        if custom_password:
            _apply_password_update(existing["id"], custom_password)
        else:
            _apply_profile_update(existing["id"])
        return True

    conn.execute(
        """
        INSERT INTO users
            (username, password_hash, display_name, role, active, created_at, updated_at)
        VALUES (?, ?, ?, 'staff', 1, ?, ?)
        """,
        (
            username,
            generate_password_hash(custom_password or default_password),
            name,
            now,
            now,
        ),
    )
    return True


def _sync_org_people_logins(conn=None):
    """기존 소속 인원에 로그인 계정이 없으면 자동 생성. 기존 비밀번호는 유지."""
    owns_conn = conn is None
    if owns_conn:
        conn = _user_db()
    try:
        rows = conn.execute("SELECT name FROM org_people ORDER BY id").fetchall()
        for row in rows:
            username = (row["name"] or "").strip()
            if not username:
                continue
            exists = conn.execute(
                "SELECT id FROM users WHERE username = ? COLLATE NOCASE",
                (username,),
            ).fetchone()
            if not exists:
                _ensure_person_login(conn, username)
        if owns_conn:
            conn.commit()
    finally:
        if owns_conn:
            conn.close()


def _reset_org_people_passwords(password="1234", conn=None):
    """소속 인원 로그인 계정의 비밀번호를 일괄 변경. admin 계정은 제외."""
    owns_conn = conn is None
    if owns_conn:
        conn = _user_db()
    updated = 0
    try:
        rows = conn.execute("SELECT name FROM org_people ORDER BY id").fetchall()
        now = datetime.now().isoformat(timespec="seconds")
        password_hash = generate_password_hash(password)
        for row in rows:
            username = (row["name"] or "").strip()
            if not username:
                continue
            user = conn.execute(
                "SELECT id, role FROM users WHERE username = ? COLLATE NOCASE",
                (username,),
            ).fetchone()
            if not user or user["role"] == "admin":
                # 계정 없으면 기본 비밀번호로 생성
                if not user:
                    _ensure_person_login(conn, username, password=password)
                    updated += 1
                continue
            conn.execute(
                """
                UPDATE users
                SET password_hash = ?, active = 1, updated_at = ?
                WHERE id = ?
                """,
                (password_hash, now, user["id"]),
            )
            updated += 1
        if owns_conn:
            conn.commit()
        return updated
    finally:
        if owns_conn:
            conn.close()


def _normalize_login_text(value):
    return unicodedata.normalize("NFC", (value or "").strip())


def _normalize_login_password(password):
    return _normalize_login_text(password).replace("＠", "@")


def _find_user_for_login(conn, username):
    username = _normalize_login_text(username)
    if not username:
        return None

    user = conn.execute(
        "SELECT * FROM users WHERE username = ? COLLATE NOCASE", (username,)
    ).fetchone()
    if user:
        return user

    person = conn.execute(
        "SELECT name FROM org_people WHERE name = ? COLLATE NOCASE", (username,)
    ).fetchone()
    if not person:
        return None

    _ensure_person_login(conn, person["name"])
    return conn.execute(
        "SELECT * FROM users WHERE username = ? COLLATE NOCASE", (username,)
    ).fetchone()


def _delete_person_login(conn, name):
    username = (name or "").strip()
    if not username:
        return
    user = conn.execute(
        "SELECT id, role FROM users WHERE username = ? COLLATE NOCASE",
        (username,),
    ).fetchone()
    if user and user["role"] != "admin":
        conn.execute("DELETE FROM users WHERE id = ?", (user["id"],))


_sync_org_people_logins()


# 접수번호 형식: HWDS-YYMMDD-NNN (같은 날짜 내 3자리 순번)

# 접수 목록 (추후 DB 연동)

RECEIPTS = [

    {

        "id": 1,

        "receipt_no": "HWDS-260102-001",

        "receipt_date": "2026-01-02",

        "team_id": "domestic",

        "team": "내수영업팀",

        "customer_name": "알앤티",

        "customer_contact": "대구사무소",

        "equipment_name": "H600MR",

        "symptoms": "아크 불량",

        "as_staff": "",

        "as_staff_name": "",

        "status": "completed",

        "charge_type": "free",

        "billing_cost": None,

        "completed_date": "2026-01-13",

        "created_at": "2026-01-02T09:00:00",

    },

    {

        "id": 2,

        "receipt_no": "HWDS-260102-002",

        "receipt_date": "2026-01-02",

        "team_id": "affiliate",

        "team": "관계사영업팀",

        "customer_name": "대한조선",

        "customer_contact": "본사 구매팀",

        "equipment_name": "Super M",

        "symptoms": "냉각수 누수",

        "as_staff": "",

        "as_staff_name": "",

        "status": "completed",

        "charge_type": "free",

        "billing_cost": None,

        "completed_date": "2026-01-10",

        "created_at": "2026-01-02T10:30:00",

    },

    {

        "id": 3,

        "receipt_no": "HWDS-260105-001",

        "receipt_date": "2026-01-05",

        "team_id": "domestic",

        "team": "내수영업팀",

        "customer_name": "삼성중공업",

        "customer_contact": "거제공장",

        "equipment_name": "H800MR",

        "symptoms": "와이어 공급 불량",

        "as_staff": "",

        "as_staff_name": "",

        "status": "completed",

        "charge_type": "paid",

        "billing_cost": "450000",

        "completed_date": "2026-01-12",

        "created_at": "2026-01-05T14:00:00",

    },

    {

        "id": 4,

        "receipt_no": "HWDS-260113-001",

        "receipt_date": "2026-01-13",

        "team_id": "export",

        "team": "수출영업팀",

        "customer_name": "Global Weld",

        "customer_contact": "해외영업팀",

        "equipment_name": "H500MR",

        "symptoms": "전원 불량",

        "as_staff": "",

        "as_staff_name": "",

        "status": "received",

        "charge_type": "paid",

        "billing_cost": None,

        "completed_date": None,

        "created_at": "2026-01-13T08:00:00",

        "photos": [],

        "support_quality": False,

        "support_tech": False,

    },

    {

        "id": 5,

        "receipt_no": "HWDS-260113-002",

        "receipt_date": "2026-01-13",

        "team_id": "affiliate",

        "team": "관계사영업팀",

        "customer_name": "현대건설기계",

        "customer_contact": "울산공장",

        "equipment_name": "H600MR",

        "symptoms": "터치패널 오류",

        "as_staff": "",

        "as_staff_name": "",

        "status": "in_progress",

        "charge_type": "free",

        "billing_cost": None,

        "completed_date": None,

        "created_at": "2026-01-13T09:30:00",

        "photos": [],

        "support_quality": True,

        "support_tech": False,

    },

    {

        "id": 6,

        "receipt_no": "HWDS-260113-003",

        "receipt_date": "2026-01-13",

        "team_id": "affiliate",

        "team": "관계사영업팀",

        "customer_name": "두산에너빌리티",

        "customer_contact": "창원공장",

        "equipment_name": "Super T",

        "symptoms": "가스 누출",

        "as_staff": "",

        "as_staff_name": "",

        "status": "received",

        "charge_type": "free",

        "billing_cost": None,

        "completed_date": None,

        "created_at": "2026-01-13T11:00:00",

        "photos": [],

        "support_quality": False,

        "support_tech": True,

    },

]


# 시스템은 빈 접수 목록으로 시작합니다.
RECEIPTS.clear()
_next_id = 1


# AS 보고서 상세 (이력관리 연동 데이터, 추후 DB)
REPORT_DETAILS = {
    "HWDS-260102-001": {
        "report_date": "2026-01-05",
        "phone": "053-123-4567",
        "serial_no": "SN-2024-00123",
        "re_repair": "해당없음",
        "receipt_staff_name": "박용수",
        "as_technician": "최기술",
        "ship_date": "2024-03-15",
        "warranty": "2026-03-14",
        "inspection_date": "2026-01-04",
        "affiliation": "내수영업팀",
        "process_location": "대구 현장",
        "ship_out_date": "2026-01-05",
        "inspection_result": "토치 접지 불량 확인",
        "repair_content": "접지선 재체결 및 토치 교체",
        "detailed_work": "—",
        "parts": [],
        "support_history": "지원 요청 이력 없음",
        "photos": [],
    },
    "HWDS-260102-002": {
        "report_date": "2026-01-08",
        "phone": "02-555-1234",
        "serial_no": "SN-2023-00456",
        "re_repair": "해당없음",
        "receipt_staff_name": "박용수",
        "as_technician": "정수리",
        "ship_date": "2023-06-01",
        "warranty": "2025-05-31",
        "inspection_date": "2026-01-07",
        "affiliation": "관계사영업팀",
        "process_location": "울산 공장",
        "ship_out_date": "2026-01-08",
        "inspection_result": "냉각수 호스 노후 확인",
        "repair_content": "냉각수 호스 교체 및 누수 점검",
        "detailed_work": "냉각수 순환 테스트 완료",
        "parts": [{"name": "냉각수 호스", "qty": 1, "unit_price": 85000, "subtotal": 85000}],
        "support_history": "지원 요청 이력 없음",
        "photos": [],
    },
    "HWDS-260105-001": {
        "report_date": "2026-01-10",
        "phone": "055-777-8899",
        "serial_no": "SN-2022-00789",
        "re_repair": "1회",
        "receipt_staff_name": "김정호",
        "as_technician": "한현장",
        "ship_date": "2022-11-20",
        "warranty": "만료",
        "inspection_date": "2026-01-09",
        "affiliation": "내수영업팀",
        "process_location": "거제 현장",
        "ship_out_date": "2026-01-10",
        "inspection_result": "와이어 드라이브 모터 이상",
        "repair_content": "와이어 드라이브 모터 교체",
        "detailed_work": "와이어 공급 테스트 및 아크 시운전 완료",
        "parts": [{"name": "와이어 드라이브 모터", "qty": 1, "unit_price": 320000, "subtotal": 320000}],
        "support_history": "기술부 지원 요청 (2026-01-08)",
        "photos": [],
    },
}

# 기존 샘플 이력/보고서 데이터도 모두 제거합니다.
REPORT_DETAILS.clear()


def _row_get(row, key, default=None):
    try:
        value = row[key]
    except (KeyError, IndexError):
        return default
    return default if value is None else value


def _parse_json_list(raw):
    if isinstance(raw, list):
        return raw
    try:
        data = json.loads(raw or "[]")
        return data if isinstance(data, list) else []
    except (TypeError, json.JSONDecodeError):
        return []


def _save_service_data():
    """접수·이력·보고서 데이터를 SQLite DB에 저장."""
    with SERVICE_DATA_LOCK, _user_db() as conn:
        conn.execute("DELETE FROM report_details")
        conn.execute("DELETE FROM receipts")
        conn.executemany(
            """
            INSERT INTO receipts (
                id, receipt_no, receipt_date, team_id, team,
                customer_name, customer_contact, equipment_name, symptoms,
                as_staff, as_staff_name, receipt_staff, receipt_staff_name,
                status, charge_type, billing_cost, completed_date, created_at,
                photos_json, phone, ship_date, support_quality, support_tech,
                created_by_user_id, created_by_username,
                severity, intervention_status, compensation_status,
                intervention_note, compensation_note, intervention_by, intervention_at,
                escalation_level, compensation_amount, intervention_log,
                billing_margin_rate, billing_manual, billing_set_by,
                product_line, spec, lot, office, notice_sent
            )
            VALUES (
                ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?, ?,
                ?, ?, ?, ?, ?
            )
            """,
            [
                (
                    item.get("id"),
                    item.get("receipt_no", ""),
                    item.get("receipt_date", ""),
                    item.get("team_id", ""),
                    item.get("team", ""),
                    item.get("customer_name", ""),
                    item.get("customer_contact", ""),
                    item.get("equipment_name", ""),
                    item.get("symptoms", ""),
                    item.get("as_staff", ""),
                    item.get("as_staff_name", ""),
                    item.get("receipt_staff", ""),
                    item.get("receipt_staff_name", ""),
                    item.get("status", "received"),
                    item.get("charge_type", "free"),
                    item.get("billing_cost"),
                    item.get("completed_date"),
                    item.get("created_at", datetime.now().isoformat()),
                    json.dumps(item.get("photos") or [], ensure_ascii=False),
                    item.get("phone", ""),
                    item.get("ship_date", ""),
                    1 if item.get("support_quality") else 0,
                    1 if item.get("support_tech") else 0,
                    item.get("created_by_user_id"),
                    item.get("created_by_username", ""),
                    item.get("severity") or "medium",
                    item.get("intervention_status") or "none",
                    item.get("compensation_status") or "none",
                    item.get("intervention_note") or "",
                    item.get("compensation_note") or "",
                    item.get("intervention_by") or "",
                    item.get("intervention_at"),
                    item.get("escalation_level") or v32.normalize_escalation(item),
                    item.get("compensation_amount") or "",
                    json.dumps(item.get("intervention_log") or [], ensure_ascii=False),
                    item.get("billing_margin_rate"),
                    1 if item.get("billing_manual") else 0,
                    item.get("billing_set_by") or "",
                    plex.normalize_receipt_line(item.get("product_line")),
                    item.get("spec") or "",
                    item.get("lot") or "",
                    item.get("office") or "",
                    1 if item.get("notice_sent") else 0,
                )
                for item in RECEIPTS
            ],
        )
        now = datetime.now().isoformat(timespec="seconds")
        conn.executemany(
            """
            INSERT INTO report_details (receipt_no, data_json, updated_at)
            VALUES (?, ?, ?)
            """,
            [
                (
                    receipt_no,
                    json.dumps(detail, ensure_ascii=False),
                    detail.get("history_updated_at") or now,
                )
                for receipt_no, detail in REPORT_DETAILS.items()
            ],
        )


def _load_service_data():
    """SQLite DB에서 접수·이력 데이터를 불러오고 기존 JSON은 최초 1회 이전."""
    global _next_id

    with _user_db() as conn:
        rows = conn.execute("SELECT * FROM receipts ORDER BY id").fetchall()
        detail_rows = conn.execute(
            "SELECT receipt_no, data_json FROM report_details"
        ).fetchall()

    def archive_legacy_file():
        if not LEGACY_SERVICE_DATA_PATH.exists():
            return
        archive_path = LEGACY_SERVICE_DATA_PATH.with_suffix(".migrated.json")
        try:
            if archive_path.exists():
                archive_path.unlink()
            LEGACY_SERVICE_DATA_PATH.replace(archive_path)
        except OSError:
            pass

    if rows:
        RECEIPTS.clear()
        for row in rows:
            try:
                photos = json.loads(row["photos_json"] or "[]")
            except json.JSONDecodeError:
                photos = []
            RECEIPTS.append(
                {
                    "id": row["id"],
                    "receipt_no": row["receipt_no"],
                    "receipt_date": row["receipt_date"],
                    "team_id": row["team_id"],
                    "team": row["team"],
                    "customer_name": row["customer_name"],
                    "customer_contact": row["customer_contact"] or "",
                    "equipment_name": row["equipment_name"],
                    "symptoms": row["symptoms"],
                    "as_staff": row["as_staff"] or "",
                    "as_staff_name": row["as_staff_name"] or "",
                    "receipt_staff": row["receipt_staff"] or "",
                    "receipt_staff_name": row["receipt_staff_name"] or "",
                    "status": row["status"],
                    "charge_type": row["charge_type"],
                    "billing_cost": row["billing_cost"],
                    "completed_date": row["completed_date"],
                    "created_at": row["created_at"],
                    "photos": photos,
                    "phone": row["phone"] or "",
                    "ship_date": row["ship_date"] or "",
                    "support_quality": bool(row["support_quality"]),
                    "support_tech": bool(row["support_tech"]),
                    "created_by_user_id": row["created_by_user_id"],
                    "created_by_username": row["created_by_username"] or "",
                    "severity": _row_get(row, "severity", "medium") or "medium",
                    "intervention_status": _row_get(row, "intervention_status", "none")
                    or "none",
                    "compensation_status": _row_get(row, "compensation_status", "none")
                    or "none",
                    "intervention_note": _row_get(row, "intervention_note", "") or "",
                    "compensation_note": _row_get(row, "compensation_note", "") or "",
                    "intervention_by": _row_get(row, "intervention_by", "") or "",
                    "intervention_at": _row_get(row, "intervention_at", None),
                    "escalation_level": _row_get(row, "escalation_level", "")
                    or v32.SEVERITY_TO_ESCALATION.get(
                        _row_get(row, "severity", "medium") or "medium", "normal"
                    ),
                    "compensation_amount": _row_get(row, "compensation_amount", "") or "",
                    "intervention_log": _parse_json_list(
                        _row_get(row, "intervention_log", "[]")
                    ),
                    "billing_margin_rate": _row_get(row, "billing_margin_rate", None),
                    "billing_manual": int(_row_get(row, "billing_manual", 0) or 0),
                    "billing_set_by": _row_get(row, "billing_set_by", "") or "",
                    "product_line": plex.normalize_receipt_line(
                        _row_get(row, "product_line", "equipment")
                    ),
                    "spec": _row_get(row, "spec", "") or "",
                    "lot": _row_get(row, "lot", "") or "",
                    "office": _row_get(row, "office", "") or "",
                    "notice_sent": int(_row_get(row, "notice_sent", 0) or 0),
                }
            )

        REPORT_DETAILS.clear()
        for row in detail_rows:
            try:
                detail = json.loads(row["data_json"])
            except json.JSONDecodeError:
                detail = {}
            REPORT_DETAILS[row["receipt_no"]] = detail
        archive_legacy_file()

    elif LEGACY_SERVICE_DATA_PATH.exists():
        try:
            payload = json.loads(LEGACY_SERVICE_DATA_PATH.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return {}
        receipts = payload.get("receipts")
        details = payload.get("report_details")
        if isinstance(receipts, list):
            RECEIPTS.clear()
            RECEIPTS.extend(receipts)
        if isinstance(details, dict):
            REPORT_DETAILS.clear()
            REPORT_DETAILS.update(details)
        if RECEIPTS or REPORT_DETAILS:
            _save_service_data()
            archive_legacy_file()

    _next_id = max((int(item.get("id", 0)) for item in RECEIPTS), default=0) + 1


_load_service_data()


def _normalize_photo(item, default_tag="customer"):
    """사진 항목을 {url, tag} 형태로 정규화"""
    allowed = {"customer", "before", "after", "parts"}
    if isinstance(item, dict):
        tag = item.get("tag") or default_tag
        if tag not in allowed:
            tag = default_tag
        return {"url": item.get("url", ""), "tag": tag}
    if isinstance(item, str) and item:
        tag = default_tag if default_tag in allowed else "customer"
        return {"url": item, "tag": tag}
    return None


def _photo_url_set(photos, default_tag="customer"):
    urls = set()
    for item in photos or []:
        norm = _normalize_photo(item, default_tag)
        if norm and norm.get("url"):
            urls.add(norm["url"])
    return urls


def _disk_photo_entries(receipt_no, subdir="", default_tag="customer"):
    """uploads/{접수번호}/[subdir] 실제 파일을 {url, tag} 목록으로 반환."""
    if not receipt_no:
        return []
    folder = UPLOAD_DIR / receipt_no
    if subdir:
        folder = folder / subdir
    if not folder.is_dir():
        return []
    entries = []
    for path in sorted(folder.iterdir()):
        if not path.is_file():
            continue
        if path.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue
        rel = (
            f"uploads/{receipt_no}/{subdir}/{path.name}"
            if subdir
            else f"uploads/{receipt_no}/{path.name}"
        )
        rel = rel.replace("//", "/")
        # 요청 컨텍스트 없이도 동작하도록 static 절대경로 사용
        url = f"/static/{rel}"
        entries.append({"url": url, "tag": default_tag})
    return entries


def _recover_orphan_upload_photos(receipt, persist=False):
    """DB photos가 비었는데 디스크에 파일이 있으면 연결 복구.

    - static/uploads/{접수번호}/* → 접수(고객전달) 사진
    - static/uploads/{접수번호}/field/* → 이력(현장) 사진
    같은 URL을 양쪽에 넣지 않는다.
    """
    if not receipt:
        return False
    receipt_no = (receipt.get("receipt_no") or "").strip()
    if not receipt_no:
        return False

    changed = False
    detail = REPORT_DETAILS.setdefault(receipt_no, {})

    disk_customer = _disk_photo_entries(receipt_no, "", "customer")
    disk_field = _disk_photo_entries(receipt_no, "field", "before")

    current_customer = list(receipt.get("photos") or [])
    current_field = list(detail.get("photos") or [])
    known = _photo_url_set(current_customer, "customer") | _photo_url_set(current_field, "before")

    if not current_customer and disk_customer:
        receipt["photos"] = disk_customer
        changed = True
        known |= {p["url"] for p in disk_customer}
    else:
        add_customer = [p for p in disk_customer if p["url"] not in known]
        if add_customer:
            receipt["photos"] = current_customer + add_customer
            changed = True
            known |= {p["url"] for p in add_customer}

    if not current_field and disk_field:
        detail["photos"] = disk_field
        changed = True
    else:
        add_field = [p for p in disk_field if p["url"] not in known]
        if add_field:
            detail["photos"] = list(detail.get("photos") or []) + add_field
            changed = True

    if _dedupe_receipt_and_detail_photos(receipt):
        changed = True

    if changed and persist:
        try:
            _save_service_data()
        except Exception:
            pass
    return changed


def _photo_url_key(url):
    """비교용 URL 키(쿼리·중복 슬래시 정규화)."""
    text = (url or "").strip()
    if not text:
        return ""
    text = text.split("?", 1)[0]
    return text.replace("\\", "/").rstrip("/")


def _unique_photos(photos, default_tag="customer"):
    out = []
    seen = set()
    for item in photos or []:
        norm = _normalize_photo(item, default_tag)
        if not norm or not norm.get("url"):
            continue
        key = _photo_url_key(norm["url"])
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(norm)
    return out


def _dedupe_receipt_and_detail_photos(receipt):
    """접수·이력 photos 내부/상호 중복 제거. 변경 여부 반환."""
    if not receipt:
        return False
    receipt_no = (receipt.get("receipt_no") or "").strip()
    detail = REPORT_DETAILS.setdefault(receipt_no, {}) if receipt_no else {}

    before_c = list(receipt.get("photos") or [])
    before_f = list(detail.get("photos") or [])

    customer = _unique_photos(before_c, "customer")
    field = _unique_photos(before_f, "before")

    customer_keys = {_photo_url_key(p["url"]) for p in customer}
    # 접수(루트)에 이미 있는 URL은 이력에서 제거 → 보고서 12장 중복 방지
    field = [p for p in field if _photo_url_key(p["url"]) not in customer_keys]

    # field 경로 사진은 이력 전용으로 유지, 접수 쪽에 잘못 들어간 경우 제거
    cleaned_customer = []
    for p in customer:
        key = _photo_url_key(p["url"])
        if "/field/" in key:
            if all(_photo_url_key(x["url"]) != key for x in field):
                field.append({**p, "tag": p.get("tag") if p.get("tag") in {"before", "after", "parts"} else "before"})
            continue
        cleaned_customer.append(p)

    receipt["photos"] = cleaned_customer
    detail["photos"] = field
    return before_c != cleaned_customer or before_f != field


def _flatten_photo_urls(photos):
    """보고서/미리보기용 URL 목록"""
    urls = []
    for p in photos or []:
        norm = _normalize_photo(p)
        if norm and norm["url"]:
            urls.append(norm["url"])
    return urls


def _build_report(receipt):
    """접수 + 이력 데이터로 AS 보고서 구성"""
    _recover_orphan_upload_photos(receipt, persist=True)
    detail = REPORT_DETAILS.get(receipt["receipt_no"], {})
    status_map = {
        "received": "접수",
        "draft": "임시저장",
        "in_progress": "진행중",
        "completed": "완료",
    }
    charge_label = "무상" if receipt["charge_type"] == "free" else "유상"

    inspection = detail.get("inspection_result", "")
    symptoms = receipt["symptoms"]
    if inspection:
        symptom_text = f"{symptoms} → 점검결과: {inspection}"
    else:
        symptom_text = symptoms

    parts = detail.get("parts", [])
    parts_total = sum(p.get("subtotal", 0) for p in parts)

    if receipt["charge_type"] == "free":
        billing_note = "이 건은 무상 처리 건으로 고객 청구비용이 없습니다."
        billing_amount = "—"
    elif receipt.get("billing_cost"):
        billing_amount = f"{int(receipt['billing_cost']):,}원"
        billing_note = f"고객 청구비용: {billing_amount}"
    else:
        billing_amount = "미책정"
        billing_note = "유상 처리 건이나 청구비용이 아직 책정되지 않았습니다."

    ship_warranty = "—"
    ship_date = detail.get("ship_date") or receipt.get("ship_date")
    if ship_date:
        ship_warranty = ship_date
        if detail.get("warranty"):
            ship_warranty += f" / {detail['warranty']}"

    loc = detail.get("process_location") or "—"
    if loc in ("field", "현장"):
        loc_label = "현장"
    elif loc in ("factory", "공장"):
        loc_label = "공장"
    else:
        loc_label = loc

    re_repair = detail.get("re_repair", "N")
    if re_repair in ("Y", "1회", "예"):
        re_repair_label = "Y"
    elif re_repair in ("N", "해당없음", "아니오"):
        re_repair_label = "N"
    else:
        re_repair_label = re_repair

    customer_photos = _unique_photos(receipt.get("photos") or [], "customer")
    field_photos = _unique_photos(detail.get("photos") or [], "before")
    # 보고서에서도 URL 기준 1회만 표시
    seen = set()
    all_photos = []
    for p in customer_photos + field_photos:
        if not p or not p.get("url"):
            continue
        key = _photo_url_key(p["url"])
        if key in seen:
            continue
        seen.add(key)
        all_photos.append(p)

    return {
        "receipt_no": receipt["receipt_no"],
        "report_date": detail.get("report_date", receipt.get("completed_date") or receipt["receipt_date"]),
        "charge_type": receipt["charge_type"],
        "charge_label": charge_label,
        "status": receipt["status"],
        "status_label": status_map.get(receipt["status"], receipt["status"]),
        "status_class": {
            "received": "received",
            "draft": "progress",
            "in_progress": "progress",
            "parts_waiting": "progress",
            "on_hold": "progress",
            "completed": "completed",
        }.get(
            receipt["status"], receipt["status"]
        ),
        "customer_name": receipt["customer_name"],
        "customer_contact": receipt.get("customer_contact") or "—",
        "office": receipt.get("office") or "—",
        "phone": receipt.get("phone") or detail.get("phone") or "—",
        "equipment_name": receipt["equipment_name"],
        "spec": receipt.get("spec") or "",
        "lot": receipt.get("lot") or "",
        "product_line": plex.normalize_receipt_line(receipt.get("product_line")),
        "product_identity": plex.product_identity_text(receipt),
        "serial_no": (
            plex.product_identity_text(receipt)
            if plex.receipt_is_rod(receipt)
            else detail.get("serial_no", "—")
        )
        or "—",
        "re_repair": re_repair_label,
        "receipt_staff_name": (
            _staff_name("receipt_staff", receipt.get("receipt_staff"))
            or receipt.get("receipt_staff_name")
            or detail.get("receipt_staff_name")
            or "—"
        ),
        "as_technician": (
            _staff_name("as_staff", receipt.get("as_staff"))
            or receipt.get("as_staff_name")
            or detail.get("as_technician")
            or "미지정"
        ),
        "ship_warranty": ship_warranty,
        "inspection_date": detail.get("inspection_date", "—"),
        "affiliation": _normalize_affiliation(
            detail.get("affiliation") or receipt.get("team") or "",
            receipt,
        )
        or "—",
        "process_location": loc_label,
        "ship_out_date": detail.get("ship_out_date", receipt.get("completed_date") or "—"),
        "symptom_text": symptom_text,
        "repair_content": detail.get("repair_content", "—"),
        "detailed_work": detail.get("detailed_work", "—"),
        "parts": parts,
        "parts_total": parts_total,
        "billing_amount": billing_amount,
        "billing_note": billing_note,
        "support_history": detail.get("support_history", "지원 요청 이력 없음"),
        "support_feedback": detail.get("support_feedback") or {},
        "support_quality": bool(
            receipt.get("support_quality") or detail.get("support_quality")
        ),
        "support_tech": bool(
            receipt.get("support_tech") or detail.get("support_tech")
        ),
        "support_status_labels": {
            "waiting": "대기중",
            "confirmed": "확인함",
            "replied": "회신완료",
            "rejected": "반려됨",
        },
        "photos": _flatten_photo_urls(all_photos),
        "photo_items": all_photos,
        "ui_labels": plex.ui_labels(
            "rods" if plex.receipt_is_rod(receipt) else "welding"
        ),
        # 6.11 최고기술경영실 개입 현황 (전 직원 조회)
        **_tech_mgmt_report_fields(receipt),
    }


def _tech_mgmt_report_fields(receipt):
    """AS보고서용 개입·보상 조회 필드."""
    esc = v32.normalize_escalation(receipt)
    logs = receipt.get("intervention_log") or []
    if isinstance(logs, str):
        logs = _parse_json_list(logs)
    return {
        "escalation_level": esc,
        "escalation_label": dict(v32.ESCALATION_OPTIONS).get(esc, esc),
        "intervention_status": receipt.get("intervention_status") or "none",
        "intervention_label": _intervention_label(receipt.get("intervention_status")),
        "intervention_note": receipt.get("intervention_note") or "",
        "intervention_by": receipt.get("intervention_by") or "",
        "intervention_at": receipt.get("intervention_at") or "",
        "intervention_log": logs,
        "compensation_status": receipt.get("compensation_status") or "none",
        "compensation_label": _compensation_label(receipt.get("compensation_status")),
        "compensation_amount": receipt.get("compensation_amount") or "",
        "compensation_note": receipt.get("compensation_note") or "",
    }


def _allowed_image(filename):
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def _save_uploaded_photos(receipt_no, files, subdir="", tags=None):
    """사진을 static/uploads/{접수번호}/[subdir] 에 저장.
    tags가 있으면 [{url, tag}] 목록, 없으면 URL 문자열 목록 반환."""
    saved = []
    if not files:
        return saved

    folder = UPLOAD_DIR / receipt_no
    if subdir:
        folder = folder / subdir
    folder.mkdir(parents=True, exist_ok=True)

    for idx, f in enumerate(files):
        if not f or not f.filename:
            continue
        if not _allowed_image(f.filename):
            continue
        original = secure_filename(f.filename)
        ext = Path(original).suffix.lower() or ".jpg"
        unique_name = f"{uuid.uuid4().hex}{ext}"
        f.save(folder / unique_name)
        rel = f"uploads/{receipt_no}/{subdir}/{unique_name}" if subdir else f"uploads/{receipt_no}/{unique_name}"
        url = url_for("static", filename=rel.replace("//", "/"))
        if tags is not None:
            tag = tags[idx] if idx < len(tags) else "before"
            if tag not in {"customer", "before", "after", "parts"}:
                tag = "before"
            saved.append({"url": url, "tag": tag or "before"})
        else:
            saved.append(url)
    return saved


def _get_receipt_by_no(receipt_no):
    return next((r for r in RECEIPTS if r["receipt_no"] == receipt_no), None)


def _format_mobile_phone(value):
    value = (value or "").strip()
    digits = "".join(char for char in value if char.isdigit())
    if len(digits) == 11 and digits.startswith("010"):
        return f"{digits[:3]}-{digits[3:7]}-{digits[7:]}"
    return value


def _is_receipt_owner(receipt):
    return bool(
        receipt
        and session.get("user_id")
        and receipt.get("created_by_user_id") == session.get("user_id")
    )


def _is_receipt_registrant(receipt, person=None, user_id=None):
    """영업접수 등록자 여부 (수정·삭제·지정·변경 권한 기준)."""
    if not receipt:
        return False
    if person is None:
        person = _get_person_for_session()
    if not person:
        return False
    uid = user_id
    if uid is None:
        try:
            uid = session.get("user_id")
        except RuntimeError:
            uid = None
    if uid and receipt.get("created_by_user_id") == uid:
        return True
    if _is_office_person(person):
        return _is_my_receipt(receipt, person, uid)
    # 레거시: created_by_user_id 없을 때 접수 담당자=등록자
    if not receipt.get("created_by_user_id"):
        if receipt.get("receipt_staff") == person.get("id"):
            return True
        staff_name = (receipt.get("receipt_staff_name") or "").strip()
        person_name = (person.get("name") or "").strip()
        if staff_name and person_name and staff_name.startswith(person_name):
            return True
    return False


def _is_office_person(person=None):
    """지역 영업사무소 공용계정 여부."""
    if person is None:
        try:
            person = _get_person_for_session()
        except Exception:
            person = None
    if not person:
        return False
    if (person.get("team_code") or "") == plex.SALES_OFFICE_TEAM:
        return True
    code = (person.get("id") or "").strip()
    return code.startswith("office_")


def _office_label(person=None):
    if person is None:
        person = _get_person_for_session()
    if not person:
        return ""
    return (person.get("name") or "").strip()


def _sales_office_people():
    """계정관리 '지역 영업사무소'에 등록된 사무소 (표시·저장값은 이름)."""
    offices = []
    seen = set()
    for person in _get_org_people():
        if (person.get("team_code") or "") != plex.SALES_OFFICE_TEAM:
            continue
        if not plex.person_visible_on_line(person):
            continue
        name = (person.get("name") or "").strip()
        code = (person.get("id") or "").strip()
        if not name:
            continue
        key = name.casefold()
        if key in seen:
            continue
        seen.add(key)
        offices.append({"code": code, "name": name})
    return offices


def _sales_office_choices(current=None):
    """담당 사무소 선택지. 지역 영업사무소 등록값을 사용한다."""
    choices = list(_sales_office_people())
    current_name = (current or "").strip()
    if current_name and not any(
        office["name"] == current_name or office["code"] == current_name
        for office in choices
    ):
        choices.append({"code": "", "name": current_name})
    return choices


def _can_pick_assigned_office(person=None):
    """정직원(내수·관계사·수출)·부서장·관리자만 담당 사무소를 지정한다."""
    if _is_office_person(person):
        return False
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    try:
        if _is_sales_head(person):
            return True
    except NameError:
        if person and (
            person.get("id") == "sales_head" or person.get("team_code") == "sales_head"
        ):
            return True
    if not person:
        return False
    return (person.get("team_code") or "") in SALES_RECEIPT_STAFF_TEAMS


def _normalize_assigned_office(raw):
    value = (raw or "").strip()
    if not value:
        return ""
    for office in _sales_office_people():
        if value == office["name"] or value == office["code"]:
            return office["name"]
    return ""


def _receipt_office_matches(receipt, person):
    office = (receipt.get("office") or "").strip()
    if not office or not person:
        return False
    name = (person.get("name") or "").strip()
    code = (person.get("id") or "").strip()
    office_cf = office.casefold()
    return office_cf in {name.casefold(), code.casefold()}


def _office_notify_team_code(receipt):
    """완료 알림용 사무소 코드(팀 단위)."""
    office = (receipt.get("office") or "").strip()
    if not office:
        return ""
    office_cf = office.casefold()
    for person in _get_org_people():
        if (person.get("team_code") or "") != plex.SALES_OFFICE_TEAM:
            continue
        if (person.get("name") or "").strip().casefold() == office_cf:
            return person.get("id") or ""
        if (person.get("id") or "").strip().casefold() == office_cf:
            return person.get("id") or ""
    return office


def _is_my_receipt(receipt, person, user_id):
    """내가 접수한 건. 사무소 공용계정은 office 기준, 그 외는 접수담당자 기준."""
    if not person:
        return False
    if _is_office_person(person):
        if _receipt_office_matches(receipt, person):
            return True
        if (receipt.get("office") or "").strip():
            return False
        if user_id and receipt.get("created_by_user_id") == user_id:
            return True
        created_name = (receipt.get("created_by_username") or "").strip()
        person_name = (person.get("name") or "").strip()
        return bool(
            created_name
            and person_name
            and created_name.casefold() == person_name.casefold()
        )
    if receipt.get("receipt_staff") == person["id"]:
        return True
    staff_name = (receipt.get("receipt_staff_name") or "").strip()
    return bool(staff_name and staff_name.startswith(person["name"]))


def _can_view_all_sales_receipts(person=None):
    """영업접수 목록 전체 조회: 부서장·경영진·최고기술경영실·admin."""
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if person is None:
        try:
            person = _get_person_for_session()
        except Exception:
            person = None
    if not person:
        return False
    try:
        if _is_sales_head(person):
            return True
    except NameError:
        if person.get("team_code") == "sales_head" or person.get("id") == "sales_head":
            return True
    if _is_executive_person(person) or _is_tech_mgmt_person(person):
        return True
    return False


def _can_view_sales_receipt(receipt, person=None):
    """영업접수 목록/편집 대상 단건 조회 가능 여부 (AS보고서는 사용하지 않음)."""
    if not receipt:
        return False
    if person is None:
        try:
            person = _get_person_for_session()
        except Exception:
            person = None
    if _can_view_all_sales_receipts(person):
        return True
    if not person:
        return False
    if _is_office_person(person):
        if _receipt_office_matches(receipt, person):
            return True
        if (receipt.get("office") or "").strip():
            return False
        try:
            user_id = session.get("user_id")
        except RuntimeError:
            user_id = None
        return _is_my_receipt(receipt, person, user_id)
    team = (person.get("team_code") or "").strip()
    if not team or team == plex.SALES_OFFICE_TEAM:
        return False
    if (receipt.get("team_id") or "") == team:
        return True
    if (receipt.get("handling_team_id") or "") == team:
        return True
    # 수출 정규직이 내수로 접수한 본인 건은 목록·수정에서 볼 수 있게 한다.
    try:
        if _is_export_response_chooser(person) and _is_receipt_owner(receipt):
            return True
    except NameError:
        pass
    return False


def _filter_receipts_for_sales_list(rows, person=None):
    """영업접수 목록 조회범위 (요구사항 부록 D).

    - 사무소 계정: 자기 사무소만
    - 일반 사원: 자기 팀만
    - 부서장/경영진/최고기술경영실/admin: 전체
    AS보고서 조회에는 적용하지 않는다(전사 조회 유지).
    """
    if person is None:
        try:
            person = _get_person_for_session()
        except Exception:
            person = None
    if not rows:
        return []
    if _can_view_all_sales_receipts(person):
        return list(rows)
    return [r for r in rows if _can_view_sales_receipt(r, person)]


def _can_manage_receipt(receipt):
    """접수 등록정보 수정(완료 전) — 하위 호환 alias."""
    return _can_edit_receipt_registration(receipt)


def _is_receipt_completed(receipt):
    return (receipt or {}).get("status") == "completed"


def _is_as_assignee(receipt, person=None):
    """지정 AS담당자(이력관리 AS내용 수정 권한). 공동 지정(최대 3명) 포함."""
    if not receipt or not person:
        return False
    pid = (person.get("id") or person.get("code") or "").strip()
    as_id = (receipt.get("as_staff") or "").strip()
    if pid and as_id and as_id == pid:
        return True
    for code in receipt.get("as_staff_ids") or []:
        if pid and code == pid:
            return True
    as_name = (receipt.get("as_staff_name") or "").strip()
    person_name = (person.get("name") or "").strip()
    if as_name and person_name and as_name.startswith(person_name):
        return True
    for name in receipt.get("as_staff_names") or []:
        if person_name and (name or "").startswith(person_name):
            return True
    return False


def _can_edit_receipt_registration(receipt, person=None):
    """영업접수 등록정보(고객·장비·증상 등) 수정 — 완료 전 등록자·부서장·admin."""
    if not receipt:
        return False
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if person is None:
        person = _get_person_for_session()
    if _is_receipt_completed(receipt):
        return False
    try:
        if _is_sales_head(person):
            return True
    except Exception:
        if person and (
            person.get("team_code") == "sales_head"
            or person.get("id") == "sales_head"
        ):
            return True
    return _is_receipt_registrant(receipt, person)


def _can_delete_receipt(receipt, person=None):
    """접수 삭제 — 완료 후 admin만, 진행 중에는 등록자·부서장·admin."""
    if not receipt:
        return False
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if _is_receipt_completed(receipt):
        return False
    return _can_edit_receipt_registration(receipt, person)


def _can_edit_billing_ops(receipt, person=None):
    """유·무상·청구·부품(비용) — 완료 후에도 등록자·부서장·admin. 진행 중 AS담당·팀원 포함."""
    if not receipt:
        return False
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if person is None:
        person = _get_person_for_session()
    try:
        if _is_sales_head(person):
            return True
    except Exception:
        if person and (
            person.get("team_code") == "sales_head"
            or person.get("id") == "sales_head"
        ):
            return True
    if _is_receipt_registrant(receipt, person):
        return True
    if not _is_receipt_completed(receipt):
        if _is_as_assignee(receipt, person):
            return True
        if person and (person.get("team_code") or "") == (receipt.get("team_id") or ""):
            return True
    return False


def _can_edit_history_as_content(receipt, person=None):
    """이력관리 AS내용(점검·수리·부품·사진 등) — 진행 중 AS담당·부서장·admin만.
    이관 승인 후 담당자가 받는 팀으로 바뀌면 원 팀은 수정 권한을 잃는다."""
    if not receipt:
        return False
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if _is_receipt_completed(receipt):
        return False
    if person is None:
        person = _get_person_for_session()
    try:
        if _is_sales_head(person):
            return True
    except Exception:
        if person and (
            person.get("team_code") == "sales_head"
            or person.get("id") == "sales_head"
        ):
            return True
    return _is_as_assignee(receipt, person)


def _apply_receipt_billing_from_request(receipt, request, parts=None):
    """유상 건 청구비용만 반영."""
    if not receipt or receipt.get("charge_type") != "paid":
        return
    parts_total = v32.calc_parts_total(parts or [])
    billing_mode = (request.form.get("billing_mode") or "").strip()
    if not billing_mode:
        billing_mode = "manual" if int(receipt.get("billing_manual") or 0) else "margin"
    person = _get_person_for_session()
    setter = (person or {}).get("name") or session.get("user_name") or ""
    if billing_mode == "manual":
        raw = request.form.get("billing_cost_manual", "")
        try:
            amount = int(str(raw).replace(",", "").strip() or "0")
        except ValueError:
            amount = int(receipt.get("billing_cost") or 0)
        receipt["billing_cost"] = amount
        receipt["billing_manual"] = 1
    else:
        raw_rate = request.form.get("billing_margin_rate", "")
        try:
            rate_f = float(raw_rate) if raw_rate not in (None, "") else None
        except ValueError:
            rate_f = None
        calc = v32.calc_billing_from_margin(parts_total, rate_f)
        receipt["billing_cost"] = calc["billing_cost"]
        receipt["billing_margin_rate"] = str(calc["margin_rate"])
        receipt["billing_manual"] = 0
    if setter:
        receipt["billing_set_by"] = setter


def _is_waiting(receipt):
    return receipt.get("status") in (
        "received",
        "draft",
        "in_progress",
        "parts_waiting",
        "on_hold",
    )


# 수출영업팀 접수의 AS 처리·지정 후보
EXPORT_AS_HANDLER_CODE = "jang_hb"
EXPORT_AS_HANDLER_NAME = "장호빈"
EXPORT_AS_SELF_CODE = "kim_ys_ex"  # 김윤식(수출·본인)
EXPORT_AS_SELF_NAME = "김윤식"
# 영업접수 AS 담당자 지정 시 수출팀에서만 선택 가능
EXPORT_AS_ALLOWED_CODES = frozenset({EXPORT_AS_HANDLER_CODE, EXPORT_AS_SELF_CODE})
# AS 담당 지정 후보에 넣을 영업팀(경영진·지원부서 제외)
AS_ASSIGN_SALES_TEAMS = frozenset({"domestic", "affiliate", "export"})


def _as_assign_team_codes():
    """제품군별 AS 담당 지정 가능 팀."""
    if _is_rods():
        try:
            return frozenset({plex.RODS_AS_TEAM})
        except AttributeError:
            return frozenset({"rods_as"})
    return AS_ASSIGN_SALES_TEAMS


def _as_assignable_people(person=None):
    """AS 담당자 지정 후보.

    - 영업부 부서장·대표이사·총괄·admin: 영업 3팀(또는 용접봉 AS팀) can_as 전원
    - 내수·관계사 팀장/팀원: 본인 팀만
    - 수출팀: 김윤식·장호빈만
    - 경영진(PNS관리자)·지원부서 등은 후보에서 제외
    """
    if person is None:
        try:
            person = _get_person_for_session()
        except Exception:
            person = None
    allow_teams = _as_assign_team_codes()
    people = [
        p
        for p in _get_org_people()
        if p.get("can_as")
        and plex.person_visible_on_line(p)
        and (p.get("team_code") or "") in allow_teams
    ]
    try:
        is_admin = session.get("user_role") == "admin"
    except RuntimeError:
        is_admin = False
    try:
        is_head = bool(_is_sales_head(person))
    except Exception:
        is_head = bool(
            person
            and (
                person.get("id") == "sales_head"
                or person.get("team_code") == "sales_head"
            )
        )
    if is_admin or is_head:
        return people
    viewer_team = (person.get("team_code") or "").strip() if person else ""
    if viewer_team == "export":
        return [
            p for p in people if (p.get("id") or "") in EXPORT_AS_ALLOWED_CODES
        ]
    if viewer_team in allow_teams:
        return [p for p in people if (p.get("team_code") or "") == viewer_team]
    return []


def _as_assign_org_tree(person=None):
    """영업접수 폼 AS 담당자 드롭다운용 조직 트리(권한·팀 필터 적용)."""
    allowed_ids = {p["id"] for p in _as_assignable_people(person) if p.get("id")}
    allow_teams = _as_assign_team_codes()
    tree = []
    for dept in _get_org_tree():
        if (dept.get("id") or "") not in allow_teams:
            continue
        members = [
            p
            for p in (dept.get("people") or [])
            if p.get("can_as") and p.get("id") in allowed_ids
        ]
        if members:
            tree.append({**dept, "people": members})
    return tree


def _sanitize_as_staff_choice(as_staff_id, person):
    """폼으로 넘어온 AS 담당이 권한 후보에 없으면 미지정으로 처리."""
    code = (as_staff_id or "").strip()
    if not code:
        return ""
    allowed = {p["id"] for p in _as_assignable_people(person) if p.get("id")}
    return code if code in allowed else ""


def _register_notify_codes(receipt):
    as_id = (receipt.get("as_staff") or "").strip()
    if as_id:
        return [as_id]
    team = receipt.get("team_id") or ""
    if team == "export":
        return list(EXPORT_AS_ALLOWED_CODES)
    return [
        p["id"]
        for p in _get_org_people()
        if p.get("team_code") == team and p.get("can_as") and p.get("id")
    ]


def _is_export_as_handler(person=None):
    """수출 접수를 내수 이력관리에서 처리·조회할 수 있는 인원(장호빈)."""
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if person is None:
        person = _get_person_for_session()
    if not person:
        return False
    if person.get("id") == EXPORT_AS_HANDLER_CODE:
        return True
    name = (person.get("name") or "").strip()
    return name == EXPORT_AS_HANDLER_NAME or name.startswith(EXPORT_AS_HANDLER_NAME)


def _can_see_export_history(person=None):
    """수출 접수를 이력관리 대기목록에서 볼 수 있는지.

    - 장호빈(내수 AS 처리자) / admin: 내수 탭 포함 조회·처리
    - 수출영업팀: 자팀(수출) 탭에서만 목록 확인
    - 그 외 내수영업(최영호·박용수 등): 노출 없음
    """
    if _is_export_as_handler(person):
        return True
    if person is None:
        person = _get_person_for_session()
    return bool(person and person.get("team_code") == "export")


def _match_history_tab(receipt, tab):
    """이력관리 AS 대응팀 탭 필터.

    원 접수팀(team_id) 또는 이관 후 담당팀(handling_team_id) 기준.
    수출영업팀 접수는 내수영업팀 탭 대기 목록에도 노출한다.
    """
    tid = receipt.get("team_id") or ""
    hid = (receipt.get("handling_team_id") or tid).strip()
    if tid == tab or hid == tab:
        return True
    if tab == "domestic" and tid == "export":
        return True
    return False


def _waiting_receipts(tab="domestic", person=None):
    if person is None:
        person = _get_person_for_session()
    items = []
    for receipt in _line_receipts():
        if not _is_waiting(receipt):
            continue
        if not _match_history_tab(receipt, tab):
            continue
        # 수출 건: 수출팀·장호빈(admin)만 노출. 다른 내수 인원 비노출
        if receipt.get("team_id") == "export":
            if not _can_see_export_history(person):
                continue
            # 내수 탭의 수출 건은 처리자(장호빈)만
            if tab == "domestic" and not _is_export_as_handler(person):
                continue
        if _is_jang_hb(person) and not _jang_hb_assigned(receipt, person):
            continue
        items.append(receipt)
    items.sort(key=lambda x: x["created_at"])
    return items


def _can_process_export_as(receipt, person=None):
    """수출 건 AS 처리(이력 저장) 가능 여부."""
    if not receipt or receipt.get("team_id") != "export":
        return True
    return _is_export_as_handler(person)


def _export_as_staff_label(code):
    if code == EXPORT_AS_HANDLER_CODE:
        return (
            _staff_name("as_staff", EXPORT_AS_HANDLER_CODE) or EXPORT_AS_HANDLER_NAME
        )
    if code == EXPORT_AS_SELF_CODE:
        return _staff_name("as_staff", EXPORT_AS_SELF_CODE) or EXPORT_AS_SELF_NAME
    return _staff_name("as_staff", code) or ""


def _assign_export_as_handler(receipt, *, default_if_empty=False):
    """수출 접수 AS 담당자: 장호빈·김윤식·미지정만 허용.

    default_if_empty=True 이면 미지정 시 장호빈을 기본 지정(신규 접수).
    """
    if not receipt or receipt.get("team_id") != "export":
        return
    as_id = (receipt.get("as_staff") or "").strip()
    if not as_id:
        if default_if_empty:
            receipt["as_staff"] = EXPORT_AS_HANDLER_CODE
            receipt["as_staff_name"] = _export_as_staff_label(EXPORT_AS_HANDLER_CODE)
        else:
            receipt["as_staff"] = ""
            receipt["as_staff_name"] = ""
        return
    if as_id in EXPORT_AS_ALLOWED_CODES:
        receipt["as_staff"] = as_id
        receipt["as_staff_name"] = _export_as_staff_label(as_id)
        return
    # 다른 내수 인원 등 허용 외 → 장호빈으로 교정
    receipt["as_staff"] = EXPORT_AS_HANDLER_CODE
    receipt["as_staff_name"] = _export_as_staff_label(EXPORT_AS_HANDLER_CODE)


def _sync_export_as_handlers():
    """기존 수출 접수 중 허용 외 AS 담당자만 교정한다."""
    changed = False
    for receipt in RECEIPTS:
        if receipt.get("team_id") != "export":
            continue
        before = (receipt.get("as_staff"), receipt.get("as_staff_name"))
        _assign_export_as_handler(receipt)
        after = (receipt.get("as_staff"), receipt.get("as_staff_name"))
        if before != after:
            changed = True
    if changed:
        _save_service_data()


_sync_export_as_handlers()


def _history_actor_label(person=None):
    name = ((person or {}).get("name") or "").strip()
    if name:
        return name
    try:
        return (
            (session.get("user_name") or session.get("username") or "담당자").strip()
            or "담당자"
        )
    except RuntimeError:
        return "담당자"


def _stamp_history_modified(detail, person=None):
    """이력 마지막 수정 시각·수정자를 기록하고 낙관적 잠금 토큰을 갱신한다."""
    if detail is None:
        return ""
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    detail["history_updated_at"] = now
    detail["history_updated_by"] = _history_actor_label(person)
    actor_id = ""
    if person:
        actor_id = str(person.get("id") or person.get("code") or "").strip()
    if not actor_id:
        try:
            actor_id = str(session.get("user_id") or "")
        except RuntimeError:
            actor_id = ""
    detail["history_updated_by_id"] = actor_id
    return now


def _history_lock_token(detail):
    return str((detail or {}).get("history_updated_at") or "").strip()


def _history_lock_info(detail):
    token = _history_lock_token(detail)
    by = str((detail or {}).get("history_updated_by") or "").strip()
    if token and by:
        display = f"{token} · {by}"
    elif token:
        display = token
    elif by:
        display = by
    else:
        display = "아직 저장된 이력 없음"
    return {
        "history_updated_at": token,
        "history_updated_by": by,
        "history_updated_display": display,
    }


def _history_lock_conflict(detail, submitted_token):
    """submitted_token is None 이면 이력 화면이 아닌 호출 → 검사 생략."""
    if submitted_token is None:
        return False
    return _history_lock_token(detail) != str(submitted_token).strip()


def _history_conflict_message(detail):
    who = _history_lock_info(detail)["history_updated_display"]
    return (
        f"다른 담당자가 먼저 저장했습니다. ({who}) "
        "내용을 덮어쓰지 않았습니다. 접수 건을 다시 불러온 뒤 작업하세요."
    )


def _share_conflict_result(detail):
    return {
        "ok": False,
        "conflict": True,
        "error": _history_conflict_message(detail),
        **_history_lock_info(detail),
    }


def _with_history_lock_fields(result, detail):
    out = dict(result or {})
    out.update(_history_lock_info(detail))
    return out


def _history_payload(receipt):
    """이력관리 폼용 데이터"""
    _recover_orphan_upload_photos(receipt, persist=True)
    detail = REPORT_DETAILS.get(receipt["receipt_no"], {})
    parts = detail.get("parts") or []
    photos = _unique_photos(detail.get("photos") or [], "before")
    # 고객전달 사진도 이력 사진 영역에 함께 표시(URL 중복 제외)
    seen = {_photo_url_key(p.get("url")) for p in photos}
    for p in _unique_photos(receipt.get("photos") or [], "customer"):
        key = _photo_url_key(p.get("url"))
        if key and key not in seen:
            photos.append(p)
            seen.add(key)

    loc = detail.get("process_location", "field")
    if loc in ("공장", "factory"):
        loc = "factory"
    else:
        loc = "field"

    re_repair_saved = detail.get("re_repair", "")
    suggest = v32.suggest_re_repair(
        receipt.get("customer_name") or "",
        receipt.get("equipment_name") or "",
        detail.get("serial_no") or "",
        [
            r
            for r in RECEIPTS
            if plex.normalize_receipt_line(r.get("product_line"))
            == plex.normalize_receipt_line(receipt.get("product_line"))
        ],
        REPORT_DETAILS,
        exclude_receipt_no=receipt.get("receipt_no") or "",
    )
    re_repair_matched = suggest.get("matched") or []
    re_repair_auto = bool(suggest.get("suggested") == "Y" and re_repair_matched)
    if re_repair_saved in ("Y", "1회", "예"):
        re_repair = "Y"
    elif re_repair_saved in ("N", "아니오", "해당없음"):
        re_repair = "N"
    else:
        # 담당자 미확정 → 자동 제안값 적용
        re_repair = suggest.get("suggested") or "N"

    ship_date = detail.get("ship_date") or receipt.get("ship_date") or ""
    inspect_date = detail.get("inspection_date") or date.today().isoformat()
    warranty = v32.warranty_verdict(ship_date, inspect_date)

    viewer = _get_person_for_session()
    # 완료 건: AS내용 잠금. 청구 비용은 등록자·부서장·admin이 계속 수정 가능.
    is_completed = _is_receipt_completed(receipt)
    is_locked = is_completed and session.get("user_role") != "admin"
    can_edit_as_content = _can_edit_history_as_content(receipt, viewer)
    can_edit_billing = _can_edit_billing_ops(receipt, viewer)
    if is_completed and session.get("user_role") == "admin":
        can_edit_as_content = True
    support_quality_on = bool(
        receipt.get("support_quality") or detail.get("support_quality")
    )
    can_share_tech = False
    is_headish = False
    is_lead = False
    try:
        is_headish = bool(
            _is_sales_head(viewer) or session.get("user_role") == "admin"
        )
        is_lead = bool(_is_sales_team_leader(viewer))
        can_share_tech = bool(
            _can_share_tech_support(viewer)
            and not is_locked
            and (
                is_headish
                or (_person_is_quality_support(viewer) and support_quality_on)
                or (is_lead and support_quality_on)
            )
        )
    except NameError:
        can_share_tech = bool(
            _can_share_tech_support(viewer)
            and not is_locked
            and (
                (viewer and (viewer.get("id") == "sales_head" or viewer.get("team_code") == "sales_head"))
                or session.get("user_role") == "admin"
                or (_person_is_quality_support(viewer) and support_quality_on)
            )
        )

    payload = {
        "receipt_no": receipt["receipt_no"],
        "customer_name": receipt["customer_name"],
        "customer_contact": receipt.get("customer_contact") or "",
        "equipment_name": receipt["equipment_name"],
        "symptoms": receipt["symptoms"],
        "charge_type": receipt["charge_type"],
        "charge_label": "무상" if receipt["charge_type"] == "free" else "유상",
        "team": receipt.get("team", ""),
        "team_id": receipt.get("team_id", ""),
        "status": receipt.get("status", "received"),
        "is_locked": is_locked,
        "can_edit_as_content": can_edit_as_content,
        "can_edit_billing": can_edit_billing,
        "can_support_field_draft": _can_support_field_draft(receipt, viewer),
        "billing_only_mode": bool(is_completed and can_edit_billing and not can_edit_as_content),
        "customer_photos": _flatten_photo_urls(receipt.get("photos") or []),
        "inspection_date": inspect_date,
        "affiliation": _normalize_affiliation(
            detail.get("affiliation") or receipt.get("team") or "",
            receipt,
        ),
        "as_technician": detail.get("as_technician")
        or receipt.get("as_staff_name")
        or session.get("user_name")
        or "",
        "serial_no": detail.get("serial_no") or "",
        "product_line": plex.normalize_receipt_line(receipt.get("product_line")),
        "spec": receipt.get("spec") or "",
        "lot": receipt.get("lot") or "",
        "product_identity": plex.product_identity_text(receipt),
        "ship_date": ship_date,
        "process_location": loc,
        "ship_out_date": detail.get("ship_out_date") or "",
        "inspection_result": detail.get("inspection_result") or "",
        "repair_content": detail.get("repair_content") or "",
        "detailed_work": detail.get("detailed_work") or "",
        "parts": parts,
        "photos": photos,
        "re_repair": re_repair,
        "re_repair_auto": re_repair_auto,
        "re_repair_matched": re_repair_matched,
        "warranty_code": warranty.get("code"),
        "warranty_label": warranty.get("label"),
        "assigned_as_staff": receipt.get("as_staff") or "",
        "assigned_as_name": receipt.get("as_staff_name") or "",
        "can_claim_unassigned": _can_claim_unassigned_as(receipt, viewer),
        "remarks": detail.get("remarks") or "",
        "support_quality": bool(receipt.get("support_quality") or detail.get("support_quality")),
        "support_tech": bool(receipt.get("support_tech") or detail.get("support_tech")),
        "support_tech_mgmt": bool(detail.get("support_tech_mgmt")),
        "support_feedback": detail.get("support_feedback") or {},
        "quality_overdue": False,
        "can_share_tech": can_share_tech,
        "tech_dept_label": _support_tech_dept_label(receipt),
        "billing_cost": receipt.get("billing_cost"),
        "billing_margin_rate": str(_effective_margin_rate(receipt.get("billing_margin_rate"))),
        "billing_manual": int(receipt.get("billing_manual") or 0),
        "billing_set_by": receipt.get("billing_set_by") or "",
        "billing_note": (
            "이 건은 무상 처리 건으로 고객 청구비용이 없습니다."
            if receipt["charge_type"] == "free"
            else (
                "영업 담당자가 이미 책정함."
                + (
                    f" ({receipt.get('billing_set_by')})"
                    if receipt.get("billing_set_by")
                    else ""
                )
                + " AS담당자가 필요 시 아래에서 조정할 수 있습니다."
                if (
                    receipt.get("billing_set_by")
                    and receipt.get("billing_cost") is not None
                )
                else (
                    f"현재 청구비용: {int(receipt['billing_cost']):,}원. 마진율 또는 직접입력으로 조정할 수 있습니다."
                    if receipt.get("billing_cost") is not None
                    else "유상 건입니다. 부품 합계에 마진율(기본 30%)을 적용하거나 협의가를 직접 입력하세요."
                )
            )
        ),
    }
    try:
        qfb = (payload.get("support_feedback") or {}).get("quality") or {}
        if qfb.get("status") in ("waiting", "confirmed") and qfb.get("shared_at"):
            shared = str(qfb.get("shared_at") or "")[:19]
            try:
                dt = datetime.strptime(shared, "%Y-%m-%d %H:%M:%S")
            except ValueError:
                dt = datetime.strptime(shared[:10], "%Y-%m-%d")
            try:
                from as_ext_0905 import SUPPORT_SLA_DAYS as _sla
            except Exception:
                _sla = 2
            payload["quality_overdue"] = (
                datetime.now().date() - dt.date()
            ).days >= _sla
        else:
            payload["quality_overdue"] = False
    except Exception:
        payload["quality_overdue"] = False
    tfb = (payload.get("support_feedback") or {}).get("tech") or {}
    payload["tech_urgent"] = bool(tfb.get("urgent"))
    # 4-6: 품질 공유 + SLA 지연 시에만 팀장 일반 기술공유
    payload["can_share_tech_delayed"] = bool(
        is_lead
        and not is_headish
        and payload.get("support_quality")
        and payload.get("quality_overdue")
        and not payload.get("support_tech")
        and not is_locked
    )
    # 4-7: 지연 경로로 아직 공유 불가할 때만 긴급 직행 카드
    payload["can_urgent_tech"] = bool(
        is_lead
        and not is_headish
        and not payload.get("support_tech")
        and not is_locked
        and not payload.get("can_share_tech_delayed")
    )
    # 팀장도 기술공유 래퍼 노출(긴급/지연 분기용)
    if is_lead and not is_locked:
        payload["can_share_tech"] = True
    payload.update(_history_lock_info(detail))
    return payload





def _date_key(receipt_date_str):
    """접수날짜 → YYMMDD"""
    return datetime.strptime(receipt_date_str, "%Y-%m-%d").strftime("%y%m%d")


def _max_seq_for_date(date_key, prefix="HWDS"):
    """해당 날짜의 기존 접수번호 중 최대 순번"""
    max_seq = 0
    for r in RECEIPTS:
        parts = r["receipt_no"].split("-")
        if len(parts) == 3 and parts[0] == prefix and parts[1] == date_key:
            try:
                max_seq = max(max_seq, int(parts[2]))
            except ValueError:
                pass
    return max_seq


def _preview_receipt_no(receipt_date_str):
    """다음 발급될 접수번호 미리보기"""
    date_key = _date_key(receipt_date_str)
    prefix = plex.receipt_prefix()
    return f"{prefix}-{date_key}-{_max_seq_for_date(date_key, prefix) + 1:03d}"


def _generate_receipt_no(receipt_date_str):
    """접수번호 자동생성: HWDS-YYMMDD-NNN / RD-YYMMDD-NNN"""
    date_key = _date_key(receipt_date_str)
    prefix = plex.receipt_prefix()
    next_seq = _max_seq_for_date(date_key, prefix) + 1
    if next_seq > 999:
        raise ValueError(f"해당 날짜({date_key}) 일일 접수 한도(999건)를 초과했습니다.")
    return f"{prefix}-{date_key}-{next_seq:03d}"





def _can_edit_sales_ops(receipt, person):
    """유·무상·청구·부품 편집(영업접수 목록)."""
    return _can_edit_billing_ops(receipt, person)


SALES_TEAM_LEADER_CODES = {
    "choi_yh",  # 내수영업팀 최영호
    "kim_jh_aff",  # 관계사영업팀 김정훈
    "kim_ys_ex",  # 수출영업팀 김윤식
}


def _is_sales_team_leader(person=None):
    """영업 3팀 팀장 또는 용접봉 AS담당관."""
    if not person:
        return False
    code = (person.get("id") or person.get("code") or "").strip()
    if code in SALES_TEAM_LEADER_CODES:
        return True
    try:
        officer = plex.RODS_OFFICER_CODE
    except AttributeError:
        officer = "rods_officer"
    if code == officer:
        return True
    title = person.get("title") or ""
    if "팀장" in title:
        return True
    name = (person.get("name") or "").strip()
    team = person.get("team_code") or ""
    try:
        rods_team = plex.RODS_AS_TEAM
    except AttributeError:
        rods_team = "rods_as"
    if team == rods_team and ("AS담당관" in name or name == "AS담당관"):
        return True
    return False


def _as_reassign_notify_message(receipt):
    """재지정 인앱 알림 문구: [담당자 지정] OOO건 AS담당자로 지정되었습니다"""
    cust = (receipt.get("customer_name") or "").strip()
    no = (receipt.get("receipt_no") or "").strip()
    ooo = cust or no or "접수"
    return f"[담당자 지정] {ooo}건 AS담당자로 지정되었습니다"


def _notify_as_staff_reassigned(receipt, staff_code, *, co_note=""):
    """신규 AS 담당자에게 메일·인앱 알림 발송."""
    if not receipt or not staff_code:
        return
    msg = _as_reassign_notify_message(receipt)
    if co_note:
        msg = f"{msg}{co_note}"
    try:
        notify_staff_assigned(
            receipt,
            staff_code,
            "AS 담당자",
            _get_staff_email,
            _mail_base_url(),
        )
    except Exception:
        pass
    try:
        _push_notification(
            user_code=staff_code,
            title=msg,
            body=msg,
            link="/history",
            receipt_no=(receipt.get("receipt_no") or ""),
        )
    except Exception:
        pass


def _viewer_on_as_handling_team(receipt, person):
    """현재 담당 팀(원 접수팀·이관 후 처리팀·수출→내수 라우팅) 소속 여부."""
    if not receipt or not person:
        return False
    viewer_team = (person.get("team_code") or "").strip()
    if not viewer_team:
        return False
    receipt_team = (receipt.get("team_id") or "").strip()
    handling = (receipt.get("handling_team_id") or receipt_team).strip()
    if viewer_team in {receipt_team, handling}:
        return True
    # 수출 접수를 내수 팀이 홈·목록에서 보는 경우 → 내수를 담당 팀으로 취급
    if viewer_team == "domestic" and receipt_team == "export":
        return True
    return False


def _can_reassign_as_staff(receipt, person):
    """AS 담당 지정·변경 (영업접수·홈 팀목록 동일 기준).
    - 부서장/총괄/대표이사/시스템관리자: 팀 무관 항상 가능
    - 그 외 I·IA·A(접수전담·접수+AS·AS전담): 해당 건 담당 팀 소속이면 팀장 아니어도 가능
    - 완료 건·지역 영업사무소: 불가
    - 실제 저장 이력(서비스담당자) 유무와 무관하게 계속 재지정 가능
    """
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if _is_office_person(person):
        return False
    if receipt and _is_receipt_completed(receipt):
        return False
    try:
        if _is_sales_head(person):
            return True
    except NameError:
        if person and (
            person.get("team_code") == "sales_head" or person.get("id") == "sales_head"
        ):
            return True
    try:
        if _is_exec_ops_privileged(person):
            return True
    except NameError:
        pass
    if not person:
        return False
    # I·IA·A 캡만 (지원부서·사무소 등은 can_as/can_receipt 없거나 위에서 차단)
    if not (person.get("can_as") or person.get("can_receipt")):
        return False
    # 접수 등록 시점(폼) — 본인 팀 후보 선택 가능
    if not receipt:
        return True
    if _viewer_on_as_handling_team(receipt, person):
        return True
    try:
        if _is_receipt_registrant(receipt, person):
            return True
    except RuntimeError:
        pass
    return False


def _person_role_ia_or_a(person):
    """접수+AS(IA) 또는 AS전담(A) — 이력 저장 시 서비스담당자 갱신 대상.
    품질팀·기술부·부서장·경영진·관리자 저장은 서비스담당자를 바꾸지 않는다.
    """
    if not person or not person.get("can_as"):
        return False
    try:
        if session.get("user_role") == "admin":
            return False
    except RuntimeError:
        pass
    try:
        if _is_sales_head(person) or _is_executive_person(person):
            return False
    except Exception:
        pass
    try:
        if _person_is_quality_support(person) or _person_is_tech_support(person):
            return False
    except NameError:
        pass
    if (person.get("team_code") or "") in (
        "field_support",
        "tech_mgmt",
        "executive",
        "sales_head",
    ):
        return False
    return True


def _is_jang_hb(person=None):
    """현장직 AS 전담(장호빈) — 미지정 자율 지정 제외, 팀장 지정 건만 진행."""
    if not person:
        return False
    code = (person.get("id") or person.get("code") or "").strip()
    if code == EXPORT_AS_HANDLER_CODE:
        return True
    name = (person.get("name") or "").strip()
    return name == EXPORT_AS_HANDLER_NAME or name.startswith(EXPORT_AS_HANDLER_NAME)


def _jang_hb_assigned(receipt, person=None):
    """장호빈에게 지정된 건(수출 접수 지정·최영호 팀장 지정 등)."""
    if not receipt:
        return False
    as_id = (receipt.get("as_staff") or "").strip()
    if not as_id:
        return False
    if as_id == EXPORT_AS_HANDLER_CODE:
        return True
    if person:
        pid = (person.get("id") or person.get("code") or "").strip()
        if pid and as_id == pid:
            return True
    as_name = (receipt.get("as_staff_name") or "").strip()
    return bool(
        as_name
        and (
            as_name == EXPORT_AS_HANDLER_NAME
            or as_name.startswith(EXPORT_AS_HANDLER_NAME)
        )
    )


def _can_claim_unassigned_as(receipt, person):
    """내수·관계사 등 팀원이 미지정 건을 클릭해 본인 담당으로 지정할 수 있는지.
    장호빈·사무소·수출팀(별도 라우팅) 제외."""
    if not receipt or not person:
        return False
    if _is_receipt_completed(receipt):
        return False
    if (receipt.get("as_staff") or "").strip():
        return False
    if _is_jang_hb(person) or _is_office_person(person):
        return False
    if not person.get("can_as"):
        return False
    team = (person.get("team_code") or "").strip()
    receipt_team = (receipt.get("team_id") or "").strip()
    if not team or not receipt_team or team != receipt_team:
        return False
    if team == "export":
        return False
    return True


def _claim_unassigned_as(receipt, person):
    """미지정 건을 본인 AS담당으로 지정. 이미 지정됐으면 False."""
    if not _can_claim_unassigned_as(receipt, person):
        return False
    if (receipt.get("as_staff") or "").strip():
        return False
    pid = (person.get("id") or person.get("code") or "").strip()
    pname = (
        (person.get("display_name") or person.get("name") or "").strip()
        or pid
    )
    if not pid:
        return False
    receipt["as_staff"] = pid
    receipt["as_staff_name"] = pname
    _save_service_data()
    _notify_as_staff_reassigned(receipt, pid)
    return True


def _filter_rows_for_field_support(rows, person):
    """v32 설치 전 stub. 설치 후 품질/기술 지원요청 필터로 교체됨."""
    return rows


def _support_request_kind(person=None):
    return ""


def _clear_billing_on_free(receipt):
    """유상→무상 전환 시 청구비용 관련 필드 초기화."""
    if not receipt:
        return
    receipt["billing_cost"] = None
    receipt["billing_manual"] = 0
    receipt["billing_margin_rate"] = None
    receipt["billing_set_by"] = ""


def _person_is_quality_support(person=None):
    try:
        return bool(_is_quality_support_person(person))
    except NameError:
        if not person:
            return False
        name = person.get("name") or ""
        return person.get("id") == "quality_acc" or "품질팀" in name


def _person_is_tech_support(person=None):
    try:
        return bool(_is_tech_support_person(person))
    except NameError:
        if not person:
            return False
        name = person.get("name") or ""
        return person.get("id") in ("tech_acc", "rods_lab") or "기술부" in name or "기술연구" in name


def _can_support_field_draft(receipt, person=None):
    """지원부서(품질/기술): 공유된 미완료 건에 점검·수리·상세·사진 임시저장만 허용."""
    if not receipt:
        return False
    if person is None:
        person = _get_person_for_session()
    if not person or person.get("team_code") != "field_support":
        return False
    if _is_receipt_completed(receipt):
        return False
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    kind = ""
    try:
        kind = _support_request_kind(person) or ""
    except Exception:
        kind = ""
    if not kind:
        if _person_is_quality_support(person):
            kind = "quality"
        elif _person_is_tech_support(person):
            kind = "tech"
        else:
            kind = "all"
    try:
        return bool(_receipt_matches_support(receipt, kind))
    except NameError:
        if kind == "quality":
            return bool(receipt.get("support_quality"))
        if kind == "tech":
            return bool(receipt.get("support_tech"))
        return bool(receipt.get("support_quality") or receipt.get("support_tech"))


def _support_tech_dept_label(receipt=None):
    try:
        if receipt is not None:
            return "기술연구소" if plex.receipt_is_rod(receipt) else "기술부"
    except Exception:
        pass
    return "기술연구소" if _is_rods() else "기술부"


def _can_share_tech_support(person=None):
    """기술부/기술연구소 공유: 품질팀 또는 영업부 부서장(·관리자). 영업팀원·팀장은 불가."""
    if person is None:
        try:
            person = _get_person_for_session()
        except Exception:
            person = None
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if not person:
        return False
    if _person_is_quality_support(person):
        return True
    try:
        return bool(_is_sales_head(person))
    except NameError:
        return person.get("id") == "sales_head" or person.get("team_code") == "sales_head"


def _can_request_quality_support(person=None):
    """품질팀 지원요청: 정직원(내수·관계사·수출·부서장·AS)만. 사무소 공용계정은 불가."""
    if person is None:
        try:
            person = _get_person_for_session()
        except Exception:
            person = None
    try:
        if session.get("user_role") == "admin":
            return True
    except RuntimeError:
        pass
    if not person:
        return False
    if _is_office_person(person):
        return False
    if person.get("team_code") == "field_support":
        return False
    try:
        if _is_tech_mgmt_person(person):
            return False
        if _is_executive_person(person) and not _is_sales_head(person):
            return False
    except NameError:
        pass
    try:
        return bool(_can_access_sales(person) or _can_access_history(person))
    except NameError:
        return bool(person.get("can_receipt") or person.get("can_as"))


def _append_support_note(receipt, note):
    receipt_no = (receipt or {}).get("receipt_no") or ""
    if not receipt_no or not note:
        return
    detail = REPORT_DETAILS.setdefault(receipt_no, {})
    prev = (detail.get("support_history") or "").strip()
    if not prev or prev == "지원 요청 이력 없음":
        detail["support_history"] = note
    elif note not in prev:
        detail["support_history"] = f"{prev} / {note}"


def _auto_confirm_support_on_view(receipt, person):
    """v32/0905 설치 전 stub. 설치 후 품질·기술 자동확인으로 교체."""
    return False


def _share_support_kind(receipt, kind, person, submitted_token=None):
    """체크박스 옆 「공유하기/보고하기」 즉시 반영·알림."""
    if not receipt:
        return {"ok": False, "error": "접수 건을 찾을 수 없습니다."}
    if receipt.get("status") == "completed":
        return {"ok": False, "error": "완료 처리된 건은 수정할 수 없습니다."}
    kind = (kind or "").strip()
    receipt_no = receipt.get("receipt_no") or ""
    detail = REPORT_DETAILS.get(receipt_no) or {}
    requester = (
        (person or {}).get("name")
        or session.get("user_name")
        or "담당자"
    )
    today = date.today().isoformat()

    def finish_share(mark, message):
        detail_row = REPORT_DETAILS.setdefault(receipt_no, {})
        _stamp_history_modified(detail_row, person)
        _save_service_data()
        return _with_history_lock_fields(
            {
                "ok": True,
                "already": False,
                "mark": mark,
                "kind": kind,
                "message": message,
            },
            detail_row,
        )

    if kind == "quality":
        if not _can_request_quality_support(person):
            return {"ok": False, "error": "품질팀 지원 요청 권한이 없습니다."}
        if receipt.get("support_quality"):
            return {"ok": True, "already": True, "mark": "공유됨", "kind": kind}
        if _history_lock_conflict(detail, submitted_token):
            return _share_conflict_result(detail)
        receipt["support_quality"] = True
        detail_row = REPORT_DETAILS.setdefault(receipt_no, {})
        detail_row["support_quality"] = True
        _append_support_note(receipt, f"품질팀 지원 요청 ({today})")
        result = finish_share("공유됨", f"품질팀에 공유했습니다. (접수번호: {receipt_no})")
        try:
            _notify_support_request(
                {**receipt, "support_quality": True, "support_tech": False},
                requester=requester,
            )
        except Exception:
            pass
        return result

    if kind == "tech":
        if not _can_share_tech_support(person):
            return {
                "ok": False,
                "error": "기술 지원 부서 공유 권한이 없습니다.",
            }
        is_head = False
        try:
            is_head = bool(_is_sales_head(person) or session.get("user_role") == "admin")
        except NameError:
            is_head = bool(
                session.get("user_role") == "admin"
                or (person and (person.get("id") == "sales_head" or person.get("team_code") == "sales_head"))
            )
        if not is_head and not receipt.get("support_quality"):
            return {
                "ok": False,
                "error": "품질팀에 공유된 건만 기술 지원 부서로 이관할 수 있습니다.",
            }
        tech_label = _support_tech_dept_label(receipt)
        if receipt.get("support_tech"):
            return {"ok": True, "already": True, "mark": "공유됨", "kind": kind}
        if _history_lock_conflict(detail, submitted_token):
            return _share_conflict_result(detail)
        receipt["support_tech"] = True
        detail_row = REPORT_DETAILS.setdefault(receipt_no, {})
        detail_row["support_tech"] = True
        _append_support_note(receipt, f"{tech_label} 지원 요청 ({today})")
        result = finish_share("공유됨", f"{tech_label}에 공유했습니다. (접수번호: {receipt_no})")
        try:
            _notify_support_request(
                {**receipt, "support_quality": False, "support_tech": True},
                requester=requester,
            )
        except Exception:
            pass
        return result

    if kind == "tech_mgmt":
        if person and person.get("team_code") == "field_support":
            return {"ok": False, "error": "최고기술경영실 보고 권한이 없습니다."}
        if detail.get("support_tech_mgmt"):
            return {"ok": True, "already": True, "mark": "보고됨", "kind": kind}
        if _history_lock_conflict(detail, submitted_token):
            return _share_conflict_result(detail)
        detail_row = REPORT_DETAILS.setdefault(receipt_no, {})
        detail_row["support_tech_mgmt"] = True
        _append_support_note(receipt, f"최고기술경영실 보고 ({today})")
        if v32.normalize_escalation(receipt) == "normal":
            receipt["escalation_level"] = "monitor"
            receipt["severity"] = v32.ESCALATION_TO_SEVERITY["monitor"]
            receipt["intervention_status"] = "watching"
        result = finish_share("보고됨", f"최고기술경영실에 보고했습니다. (접수번호: {receipt_no})")
        try:
            _push_notification(
                team_code="tech_mgmt",
                title=f"[지원요청] {receipt_no} 최고기술경영실 보고",
                body=f"요청자: {requester} / 등급 모니터링 상향",
                link="/tech-monitor",
                receipt_no=receipt_no,
            )
            notify_tech_mgmt_reported(
                receipt,
                _get_staff_email,
                _mail_base_url(),
                requester=requester,
            )
        except Exception:
            pass
        return result

    return {"ok": False, "error": "알 수 없는 공유 유형입니다."}


def _staged_support_flags(person, receipt, form):
    """영업팀(내수·수출·관계사)은 품질팀만. 기술부 이관은 품질팀 또는 영업부 부서장."""
    current_quality = bool(receipt.get("support_quality")) if receipt else False
    current_tech = bool(receipt.get("support_tech")) if receipt else False
    is_head = False
    try:
        is_head = bool(_is_sales_head(person) or session.get("user_role") == "admin")
    except NameError:
        is_head = bool(
            session.get("user_role") == "admin"
            or (person and (person.get("id") == "sales_head" or person.get("team_code") == "sales_head"))
        )
    if _person_is_quality_support(person):
        quality = current_quality
        tech = bool(form.get("support_tech")) if current_quality else current_tech
        return quality, tech
    if is_head:
        quality = bool(form.get("support_quality")) if _can_request_quality_support(person) else current_quality
        tech = bool(form.get("support_tech"))
        return quality, tech
    if _is_office_person(person) or not _can_request_quality_support(person):
        return current_quality, current_tech
    quality = bool(form.get("support_quality"))
    return quality, current_tech


def _notify_support_request(receipt, requester=""):
    """6.4: 품질팀/기술부 지원요청 알림 (해당 계정만)."""
    if not receipt:
        return
    no = receipt.get("receipt_no") or ""
    body_tail = f"{(receipt.get('customer_name') or '')} / {(receipt.get('symptoms') or '')[:80]}"
    req = requester or "담당자"
    try:
        push = _push_notification
    except NameError:
        return
    if receipt.get("support_quality"):
        push(
            user_code="quality_acc",
            title=f"[지원요청] {no} 품질팀 지원 필요",
            body=f"요청자: {req} / {body_tail}",
            link="/history",
            receipt_no=no,
        )
        try:
            notify_support_requested(
                receipt,
                _get_staff_email,
                _mail_base_url(),
                target_code="quality_acc",
                target_label="품질팀",
                requester=req,
            )
        except Exception:
            pass
    if receipt.get("support_tech"):
        tech_code = plex.RODS_LAB_CODE if plex.receipt_is_rod(receipt) else "tech_acc"
        tech_label = "기술연구소" if plex.receipt_is_rod(receipt) else "기술부"
        push(
            user_code=tech_code,
            title=f"[지원요청] {no} {tech_label} 지원 필요",
            body=f"요청자: {req} / {body_tail}",
            link="/history",
            receipt_no=no,
        )
        try:
            notify_support_requested(
                receipt,
                _get_staff_email,
                _mail_base_url(),
                target_code=tech_code,
                target_label=tech_label,
                requester=req,
            )
        except Exception:
            pass


def _get_queue_info(receipt, all_receipts):
    return v32.queue_info(receipt, all_receipts)


def _queue_summary_team_scope(person):
    """대기열 요약 집계 팀 범위. None이면 목록 전체."""
    if not person:
        return None
    team_code = person.get("team_code") or ""
    if team_code == "sales_head":
        return None
    try:
        if _can_view_all_sales_receipts(person):
            return None
    except Exception:
        pass
    return team_code or None


def _compute_queue_summary(receipts, person):
    return v32.count_queue_summary(
        receipts,
        team_id=_queue_summary_team_scope(person),
        person_id=(person or {}).get("id") or "",
    )


def _split_assignee_display(name):
    """AS담당 표시: '이름(역할)' → 이름 / (역할) 분리."""
    raw = (name or "").strip() or "미지정"
    match = re.match(r"^(.+?)\s*(\([^)]+\))\s*$", raw)
    if match:
        return match.group(1).strip(), match.group(2).strip()
    return raw, ""


def _receipt_queue_team_flags(receipt, team_scope):
    """팀 대기열 요약 배지 클릭 필터용."""
    status = receipt.get("status") or ""
    as_staff = (receipt.get("as_staff") or "").strip()
    in_team = v32.in_team_queue_scope(receipt, team_scope)
    return {
        "queue_team_unassigned_wait": bool(
            in_team and not as_staff and v32.is_wait_queue(status)
        ),
        "queue_team_in_progress": bool(
            in_team and v32.is_progress_queue(status)
        ),
    }




def _enrich_receipt(receipt, all_receipts, *, skip_queue=False, queue_team_scope=None):

    if skip_queue:
        queue = {"text": "—", "type": "none"}
    else:
        queue = _get_queue_info(receipt, all_receipts)

    if receipt["charge_type"] == "free":

        billing = "—"

    elif receipt["billing_cost"]:

        billing = f"{int(receipt['billing_cost']):,}원"

    else:

        billing = "미책정"

    detail = REPORT_DETAILS.get(receipt.get("receipt_no") or "", {}) or {}
    parts = detail.get("parts") or []
    parts_count = len(parts)
    parts_total = v32.calc_parts_total(parts)
    assignee_main, assignee_sub = _split_assignee_display(receipt["as_staff_name"] or "미지정")

    return {

        **receipt,

        "queue_text": queue["text"],

        "queue_type": queue["type"],

        "queue_main": queue.get("main") or queue["text"],

        "queue_label": queue.get("label") or "",

        "queue_rank": queue.get("rank") or "",

        "queue_scope": queue.get("scope") or "",

        "as_display": receipt["as_staff_name"] or "미지정",

        "assignee_main": assignee_main,

        "assignee_sub": assignee_sub,

        "as_unassigned": not receipt.get("as_staff"),

        "charge_label": "무상" if receipt["charge_type"] == "free" else "유상",

        "billing_display": billing,

        "completed_display": receipt["completed_date"] or "—",

        "parts_count": parts_count,

        "parts_total": parts_total,

        "parts_summary": (
            f"{parts_count}건 / {parts_total:,}원" if parts_count else "—"
        ),

        "status_label": v32.status_label(receipt.get("status")),
        "status_class": _dashboard_status_class(receipt.get("status")),
        "filter_status": "in_progress" if receipt["status"] == "draft" else receipt["status"],

        **_receipt_queue_team_flags(receipt, queue_team_scope),

    }





def _get_status_counts(receipts):

    counts = {"all": len(receipts), "received": 0, "in_progress": 0, "completed": 0}

    for r in receipts:

        if r["status"] == "draft":
            counts["in_progress"] += 1
        elif r["status"] in counts:

            counts[r["status"]] += 1

    return counts





@app.context_processor

def inject_globals():
    if "csrf_token" not in session:
        session["csrf_token"] = uuid.uuid4().hex
    person = _get_person_for_session() if session.get("user_id") else None
    can_sales = True
    can_history = True
    is_sales_head = False
    unread = 0
    if person or session.get("user_id"):
        try:
            can_sales = _can_access_sales(person)
            can_history = _can_access_history(person)
            is_sales_head = _is_sales_head(person)
            unread = _unread_notification_count(person)
        except NameError:
            # v32 설치 전 부트스트랩 단계
            can_sales = bool(not person or person.get("can_receipt") or session.get("user_role") == "admin")
            can_history = bool(not person or person.get("can_as") or session.get("user_role") == "admin")
    return {
        "current_user": session.get("user_name"),
        "current_user_id": session.get("user_id"),
        "current_username": session.get("username"),
        "current_user_role": session.get("user_role"),
        "current_user_is_executive": bool(
            _is_executive_person(person) and not is_sales_head
        ),
        "current_user_is_tech_mgmt": _is_tech_mgmt_person(person),
        "current_user_is_sales_head": is_sales_head,
        "current_user_can_sales": can_sales,
        "current_user_can_history": can_history,
        "notification_unread": unread,
        "csrf_token": session["csrf_token"],
        "login_dropdown": _get_login_dropdown() if not session.get("user_id") else [],
        "product_line": session.get("product_line") or "",
        "ui_labels": _ui_labels(),
        "is_rods": _is_rods(),
        "is_office_account": _is_office_person(person),
    }





def _get_person_for_session():
    """로그인 사용자와 매칭되는 소속 인원 정보."""
    username = (session.get("username") or "").strip()
    display_name = (session.get("user_name") or "").strip()
    if not username and not display_name:
        return None

    people = _get_org_people()
    for person in people:
        if username and person["name"].casefold() == username.casefold():
            return person
    for person in people:
        if display_name and (
            person["name"].casefold() == display_name.casefold()
            or person["display_name"].casefold() == display_name.casefold()
        ):
            return person
    return None


def _account_type_label(person):
    if not person:
        return "담당자 계정"
    if _is_office_person(person):
        return "사무소 접수 계정"
    if _is_tech_mgmt_person(person):
        return "전 팀 모니터링·개입 전용"
    if person.get("team_code") == "sales_head" or person.get("id") == "sales_head":
        return "영업부 부서장 · 전권 의뢰"
    if _is_exec_ops_privileged(person):
        return "경영진 · 접수·의뢰 총괄"
    if person.get("team_code") == "executive":
        return "전체 현황 조회 전용"
    can_receipt = person.get("can_receipt")
    can_as = person.get("can_as")
    if can_receipt and can_as:
        return "접수·AS 겸용 계정"
    if can_receipt:
        return "접수 전담 계정"
    if can_as:
        return "AS 담당 계정"
    return "조회 계정"


def _is_tech_mgmt_team_code(team_code):
    """최고기술경영실 팀 여부 (6.13 통계·개입관리 전용)."""
    if not team_code:
        return False
    if team_code == "tech_mgmt":
        return True
    # 레거시 코드·명칭 매칭 (느슨한 '모니터링' 키워드 매칭은 제외 — 타 권한 오인 방지)
    if team_code == "dept001":
        return True
    for team in _get_teams():
        if team["id"] != team_code:
            continue
        name = team.get("name") or ""
        if "최고기술경영" in name:
            return True
    return False


def _is_tech_mgmt_person(person):
    """최고기술경영실 소속만 True. 경영진·부서장·admin 역할만으로는 False."""
    return bool(person and _is_tech_mgmt_team_code(person.get("team_code")))


EXEC_OPS_PRIVILEGED_CODES = frozenset({"exec_chief", "exec_ceo"})
EXEC_OPS_PRIVILEGED_NAMES = frozenset({"대표이사", "총괄"})


def _is_exec_ops_privileged(person):
    """대표이사·총괄: 영업부 부서장과 동등한 접수·이력·공유 권한."""
    if not person:
        return False
    code = (person.get("id") or person.get("code") or "").strip()
    if code in EXEC_OPS_PRIVILEGED_CODES:
        return True
    name = (person.get("name") or "").strip()
    return name in EXEC_OPS_PRIVILEGED_NAMES


def _is_executive_person(person):
    return bool(person and person.get("team_code") == "executive")


def _is_all_team_viewer(person):
    """전 팀 현황을 볼 수 있는 계정(경영진·최고기술경영실)."""
    return _is_executive_person(person) or _is_tech_mgmt_person(person)


SEVERITY_OPTIONS = [
    ("low", "낮음"),
    ("medium", "보통"),
    ("high", "높음"),
    ("critical", "긴급"),
]
SEVERITY_RANK = {"low": 1, "medium": 2, "high": 3, "critical": 4}
INTERVENTION_OPTIONS = [
    ("none", "미개입"),
    ("watching", "주시"),
    ("intervening", "개입중"),
    ("resolved", "조치완료"),
]
COMPENSATION_OPTIONS = [
    ("none", "해당없음"),
    ("reviewing", "검토중"),
    ("negotiating", "협의중"),
    ("agreed", "합의완료"),
    ("dispute", "분쟁·소송"),
    ("paid", "보상완료"),  # 레거시
    ("rejected", "불성립"),  # 레거시
]


def _severity_label(code):
    return dict(SEVERITY_OPTIONS).get(code or "medium", code or "보통")


def _intervention_label(code):
    return dict(INTERVENTION_OPTIONS).get(code or "none", code or "미개입")


def _compensation_label(code):
    return dict(COMPENSATION_OPTIONS).get(code or "none", code or "해당없음")


def _is_high_severity(severity):
    return SEVERITY_RANK.get(severity or "medium", 2) >= 3


def _team_scope_codes(team_code, person=None):
    """팀 대시보드에 포함할 접수 팀 코드. None이면 전체."""
    if team_code == "executive" or _is_tech_mgmt_team_code(team_code):
        return None
    if team_code == "domestic":
        # 수출 접수는 내수 AS 처리자(장호빈)에게만 홈/목록 노출
        if _is_export_as_handler(person):
            return {"domestic", "export"}
        return {"domestic"}
    if team_code:
        return {team_code}
    return set()


def _receipt_dashboard_teams():
    """AS 접수가 있는 영업·서비스 팀 목록."""
    skip = {"executive", "field_support", "tech_mgmt", "dept001", plex.SALES_OFFICE_TEAM}
    if _is_rods():
        skip.update({"domestic", "affiliate", "export", "sales_head", "service"})
    else:
        skip.add(plex.RODS_AS_TEAM)
    teams = []
    for team in _get_teams():
        if team["id"] in skip or _is_tech_mgmt_team_code(team["id"]):
            continue
        teams.append(team)
    return teams


def _dashboard_stats_for_rows(rows, today):
    today_count = sum(
        1 for r in rows if (r.get("receipt_date") or "")[:10] == today
    )
    in_progress_count = sum(1 for r in rows if _is_pending_receipt(r))
    completed_count = sum(1 for r in rows if r.get("status") == "completed")
    total_count = len(rows)
    completion_rate = (
        int(round((completed_count / total_count) * 100)) if total_count else 0
    )
    return {
        "today": today_count,
        "in_progress": in_progress_count,
        "completed": completed_count,
        "rate": completion_rate,
        "total": total_count,
    }


def _is_pending_receipt(receipt):
    return receipt.get("status") in (
        "received",
        "draft",
        "in_progress",
        "parts_waiting",
        "on_hold",
    )


def _pending_item_row(receipt):
    return {
        "receipt_no": receipt["receipt_no"],
        "receipt_date": receipt.get("receipt_date") or "",
        "customer_name": receipt.get("customer_name") or "",
        "equipment_name": receipt.get("equipment_name") or "",
        "status_label": receipt.get("status_label") or "",
        "status_class": _dashboard_status_class(receipt.get("status")),
        "assignee": receipt.get("as_display") or "미지정",
    }


def _build_pending_by_assignee(pending_receipts):
    """팀 미처리 건을 AS 담당자별로 집계."""
    groups = {}
    for receipt in pending_receipts:
        assignee = receipt.get("as_display") or receipt.get("as_staff_name") or "미지정"
        bucket = groups.setdefault(
            assignee,
            {
                "assignee": assignee,
                "total": 0,
                "received": 0,
                "draft": 0,
                "in_progress": 0,
                "pending_rows": [],
            },
        )
        bucket["total"] += 1
        status = receipt.get("status") or "received"
        if status in bucket:
            bucket[status] += 1
        bucket["pending_rows"].append(_pending_item_row(receipt))

    rows = list(groups.values())
    for row in rows:
        row["pending_rows"].sort(key=lambda item: item.get("receipt_date") or "")
    rows.sort(
        key=lambda row: (
            row["assignee"] == "미지정",
            -row["total"],
            row["assignee"],
        )
    )
    return rows


def _month_key(date_str):
    return (date_str or "")[:7]


def _available_report_years(receipts):
    years = set()
    for receipt in receipts:
        for field in ("receipt_date", "completed_date", "created_at"):
            value = receipt.get(field) or ""
            if len(value) >= 4 and value[:4].isdigit():
                years.add(int(value[:4]))
    years.add(date.today().year)
    return sorted(years, reverse=True)


def _effective_margin_rate(value, default=None):
    """마진율: 미설정(None/'')만 기본 30%, 0은 0%로 유지."""
    if default is None:
        default = v32.DEFAULT_MARGIN_RATE
    if value is None:
        return float(default)
    if isinstance(value, str) and not value.strip():
        return float(default)
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default)


def _receipt_billing_amount(receipt):
    if (receipt or {}).get("charge_type") == "free":
        return 0
    raw = (receipt or {}).get("billing_cost")
    if raw in (None, ""):
        return 0
    try:
        return int(raw)
    except (TypeError, ValueError):
        try:
            return int(str(raw).replace(",", ""))
        except (TypeError, ValueError):
            return 0


def _period_bucket_key(date_str, period, year):
    """날짜 → 기간 버킷 키. 해당 없으면 None."""
    text = (date_str or "").strip()[:10]
    if len(text) < 7:
        return None
    try:
        d = datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return None
    if period == "year":
        return str(d.year)
    if d.year != int(year):
        return None
    if period == "quarter":
        return f"{year}-Q{(d.month - 1) // 3 + 1}"
    return f"{year}-{d.month:02d}"


def _build_period_axis(period, year, year_options):
    """기간 축(키·라벨) 구성."""
    if period == "quarter":
        keys = [f"{year}-Q{q}" for q in range(1, 5)]
        labels = {
            f"{year}-Q1": "1분기(1~3월)",
            f"{year}-Q2": "2분기(4~6월)",
            f"{year}-Q3": "3분기(7~9월)",
            f"{year}-Q4": "4분기(10~12월)",
        }
        short = {f"{year}-Q{q}": f"Q{q}" for q in range(1, 5)}
        return keys, labels, short
    if period == "year":
        keys = [str(y) for y in sorted(year_options)]
        labels = {k: f"{k}년" for k in keys}
        return keys, labels, labels
    keys = [f"{year}-{month:02d}" for month in range(1, 13)]
    labels = {k: f"{int(k.split('-')[1])}월" for k in keys}
    return keys, labels, labels


def _empty_period_bucket():
    return {
        "received": 0,
        "completed": 0,
        "parts_cost": 0,
        "billing_cost": 0,
        "re_repair_count": 0,
        "warranty_count": 0,
    }


def _build_status_report(year=None, period="month"):
    """경영진용 기간별(월/분기/연) AS 현황보고 데이터."""
    year = year or date.today().year
    if period not in ("month", "quarter", "year"):
        period = "month"
    year_text = str(year)
    line_rows = _line_receipts()
    year_options = _available_report_years(line_rows)
    teams = _receipt_dashboard_teams()
    enriched = [_enrich_receipt(r, line_rows) for r in line_rows]
    period_keys, period_labels, period_short = _build_period_axis(
        period, year, year_options
    )
    current_month = date.today().strftime("%Y-%m")
    current_quarter = f"{date.today().year}-Q{(date.today().month - 1) // 3 + 1}"

    period_title = {
        "month": f"월별 추이 ({year}년)",
        "quarter": f"분기별 추이 ({year}년)",
        "year": "연도별 추이",
    }[period]
    period_unit = {"month": "월", "quarter": "분기", "year": "연도"}[period]

    def blank_map():
        return {key: _empty_period_bucket() for key in period_keys}

    overall_map = blank_map()
    team_reports = []

    for team in teams:
        team_receipts = [r for r in enriched if r.get("team_id") == team["id"]]  # 7-3: 원 접수팀 1건=1회
        year_receipts = [
            r for r in team_receipts if (r.get("receipt_date") or "").startswith(year_text)
        ]
        bucket_map = blank_map()

        for receipt in team_receipts:
            detail = REPORT_DETAILS.get(receipt.get("receipt_no") or "", {}) or {}
            parts_total = int(receipt.get("parts_total") or v32.calc_parts_total(detail.get("parts") or []))
            billing = _receipt_billing_amount(receipt)
            re_repair = (detail.get("re_repair") or "")
            is_re_repair = re_repair in ("Y", "1회", "예")
            ship_date = detail.get("ship_date") or receipt.get("ship_date") or ""
            inspect_date = detail.get("inspection_date") or ""
            warranty = v32.warranty_verdict(ship_date, inspect_date or None)
            is_warranty = warranty.get("code") == "warranty"

            recv_key = _period_bucket_key(receipt.get("receipt_date"), period, year)
            if recv_key in bucket_map:
                bucket_map[recv_key]["received"] += 1
                bucket_map[recv_key]["parts_cost"] += parts_total
                bucket_map[recv_key]["billing_cost"] += billing
                if is_re_repair:
                    bucket_map[recv_key]["re_repair_count"] += 1
                if is_warranty:
                    bucket_map[recv_key]["warranty_count"] += 1

            done_key = _period_bucket_key(receipt.get("completed_date"), period, year)
            if done_key in bucket_map and receipt.get("status") == "completed":
                bucket_map[done_key]["completed"] += 1

        pending_items = [r for r in team_receipts if _is_pending_receipt(r)]
        pending_by_assignee = _build_pending_by_assignee(pending_items)
        pending_by_status = {"received": 0, "draft": 0, "in_progress": 0}
        for receipt in pending_items:
            status = receipt.get("status") or "received"
            if status in pending_by_status:
                pending_by_status[status] += 1

        period_rows = []
        for key in period_keys:
            b = bucket_map[key]
            overall = overall_map[key]
            overall["received"] += b["received"]
            overall["completed"] += b["completed"]
            overall["parts_cost"] += b["parts_cost"]
            overall["billing_cost"] += b["billing_cost"]
            overall["re_repair_count"] += b["re_repair_count"]
            overall["warranty_count"] += b["warranty_count"]
            period_rows.append(
                {
                    "key": key,
                    "label": period_labels[key],
                    "short_label": period_short[key],
                    "received": b["received"],
                    "completed": b["completed"],
                    "net": b["received"] - b["completed"],
                    "parts_cost": b["parts_cost"],
                    "billing_cost": b["billing_cost"],
                    "re_repair_count": b["re_repair_count"],
                    "warranty_count": b["warranty_count"],
                }
            )

        # 월별 키 호환(기존 템플릿/계산용)
        received_this_month = sum(
            1
            for r in team_receipts
            if _month_key(r.get("receipt_date")) == current_month
        )
        completed_this_month = sum(
            1
            for r in team_receipts
            if r.get("status") == "completed"
            and _month_key(r.get("completed_date")) == current_month
        )
        year_completed = sum(
            1
            for r in team_receipts
            if r.get("status") == "completed"
            and (r.get("completed_date") or "").startswith(year_text)
        )

        completed_count = sum(1 for r in team_receipts if r.get("status") == "completed")
        total_count = len(team_receipts)
        chart_max = max(
            [row["received"] for row in period_rows]
            + [row["completed"] for row in period_rows]
            + [1]
        )
        team_reports.append(
            {
                "team_code": team["id"],
                "team_name": team["name"],
                "total": total_count,
                "year_received": len(year_receipts),
                "year_completed": year_completed,
                "received_this_month": received_this_month,
                "completed_this_month": completed_this_month,
                "pending_count": len(pending_items),
                "completed_count": completed_count,
                "completion_rate": int(round((completed_count / total_count) * 100))
                if total_count
                else 0,
                "pending_by_status": pending_by_status,
                "pending_by_assignee": pending_by_assignee,
                "period_rows": period_rows,
                "monthly": period_rows,  # 하위 호환
                "chart_max": chart_max,
                "pending_items": [
                    _pending_item_row(receipt)
                    for receipt in sorted(
                        pending_items,
                        key=lambda item: item.get("receipt_date") or "",
                    )
                ],
            }
        )

    overall_period = []
    for key in period_keys:
        b = overall_map[key]
        overall_period.append(
            {
                "key": key,
                "label": period_labels[key],
                "short_label": period_short[key],
                "received": b["received"],
                "completed": b["completed"],
                "net": b["received"] - b["completed"],
                "parts_cost": b["parts_cost"],
                "billing_cost": b["billing_cost"],
                "re_repair_count": b["re_repair_count"],
                "warranty_count": b["warranty_count"],
            }
        )

    total_all = len(enriched)
    pending_all = sum(team["pending_count"] for team in team_reports)
    completed_all = sum(team["completed_count"] for team in team_reports)
    chart_max = max(
        [row["received"] for row in overall_period]
        + [row["completed"] for row in overall_period]
        + [1]
    )
    year_received = sum(
        1 for r in enriched if (r.get("receipt_date") or "").startswith(year_text)
    )
    year_completed = sum(
        1
        for r in enriched
        if r.get("status") == "completed"
        and (r.get("completed_date") or "").startswith(year_text)
    )
    received_this_month = sum(
        1 for r in enriched if _month_key(r.get("receipt_date")) == current_month
    )
    completed_this_month = sum(
        1
        for r in enriched
        if r.get("status") == "completed"
        and _month_key(r.get("completed_date")) == current_month
    )

    period_parts_total = sum(row["parts_cost"] for row in overall_period)
    period_billing_total = sum(row["billing_cost"] for row in overall_period)

    teams_total = {
        "team_name": "합계",
        "total": sum(t["total"] for t in team_reports),
        "pending_count": sum(t["pending_count"] for t in team_reports),
        "pending_by_status": {
            "received": sum(t["pending_by_status"].get("received", 0) for t in team_reports),
            "draft": sum(t["pending_by_status"].get("draft", 0) for t in team_reports),
            "in_progress": sum(
                t["pending_by_status"].get("in_progress", 0) for t in team_reports
            ),
        },
        "completed_count": sum(t["completed_count"] for t in team_reports),
        "completion_rate": int(round((completed_all / total_all) * 100)) if total_all else 0,
        "year_received": sum(t["year_received"] for t in team_reports),
        "year_completed": sum(t["year_completed"] for t in team_reports),
        "received_this_month": sum(t["received_this_month"] for t in team_reports),
        "completed_this_month": sum(t["completed_this_month"] for t in team_reports),
    }

    return {
        "year": year,
        "period": period,
        "period_title": period_title,
        "period_unit": period_unit,
        "year_options": year_options,
        "teams": team_reports,
        "teams_total": teams_total,
        "overall": {
            "total": total_all,
            "pending": pending_all,
            "completed": completed_all,
            "completion_rate": int(round((completed_all / total_all) * 100)) if total_all else 0,
            "received_this_month": received_this_month,
            "completed_this_month": completed_this_month,
            "year_received": year_received,
            "year_completed": year_completed,
            "period_rows": overall_period,
            "monthly": overall_period,
            "chart_max": chart_max,
            "parts_cost_total": period_parts_total,
            "billing_cost_total": period_billing_total,
            "current_quarter": current_quarter,
        },
    }


def _name_matches_person(name, person):
    if not name or not person:
        return False
    name = (name or "").strip()
    person_name = (person.get("name") or "").strip()
    if not name or not person_name:
        return False
    return (
        name.casefold() == person_name.casefold()
        or name.casefold().startswith(person_name.casefold())
        or person_name.casefold() in name.casefold()
    )


def _is_my_as_work(receipt, person, user_id=None):
    """내가 AS 처리한(또는 담당인) 건인지 판별."""
    if user_id is None:
        user_id = session.get("user_id")
    receipt_no = (receipt or {}).get("receipt_no") or ""
    detail = REPORT_DETAILS.get(receipt_no, {}) if receipt_no else {}

    if user_id and detail.get("processed_by_user_id") == user_id:
        return True
    if person:
        if receipt.get("as_staff") == person["id"]:
            return True
        if _name_matches_person(receipt.get("as_staff_name"), person):
            return True
        if _name_matches_person(detail.get("as_technician"), person):
            return True

    user_name = (session.get("user_name") or "").strip()
    username = (session.get("username") or "").strip()
    for candidate in (
        receipt.get("as_staff_name"),
        detail.get("as_technician"),
    ):
        text = (candidate or "").strip()
        if not text:
            continue
        if user_name and text.casefold().startswith(user_name.casefold()):
            return True
        if username and text.casefold().startswith(username.casefold()):
            return True
    return False


def _dashboard_customer_label(receipt, viewer_team_code):
    name = receipt.get("customer_name") or ""
    if (
        viewer_team_code == "domestic"
        and receipt.get("team_id") == "export"
        and not name.startswith("[")
    ):
        return f"[수출영업팀접수] {name}"
    return name


def _dashboard_status_class(status):
    return {
        "received": "received",
        "draft": "progress",
        "in_progress": "progress",
        "parts_waiting": "progress",
        "on_hold": "progress",
        "completed": "completed",
    }.get(status, "received")


def _build_home_dashboard():
    """로그인 담당자용 홈 대시보드 데이터."""
    person = _get_person_for_session()
    user_id = session.get("user_id")
    user_name = session.get("user_name") or session.get("username") or "담당자"

    team_code = person["team_code"] if person else ""
    is_executive = _is_executive_person(person)
    is_tech_mgmt = _is_tech_mgmt_person(person)
    is_sales_head = False
    try:
        is_sales_head = _is_sales_head(person)
    except Exception:
        is_sales_head = team_code == "sales_head"
    is_all_viewer = is_executive or is_tech_mgmt or is_sales_head
    team_name = "전체" if is_all_viewer else (person["team"] if person else "담당")
    account_label = _account_type_label(person)
    scope_codes = _team_scope_codes(team_code, person)

    # 경영진·최고기술경영실: 대기열 O(n²) 계산 생략. 부서장은 본인 접수 순번 표시.
    skip_queue = bool((is_executive or is_tech_mgmt) and not is_sales_head)
    # 최고기술경영실: 통계·카드만 필요 → 전체 enrich/정렬 생략 (홈 체감 속도)
    line_rows = _line_receipts()
    if is_tech_mgmt:
        enriched = []
        team_rows = []
        for item in line_rows:
            if scope_codes is not None and item.get("team_id") not in scope_codes:
                continue
            team_rows.append(item)
    else:
        enriched = [
            _enrich_receipt(
                r,
                line_rows,
                skip_queue=skip_queue,
                queue_team_scope=_queue_summary_team_scope(person),
            )
            for r in line_rows
        ]
        enriched.sort(key=v32.recent_board_sort_key)

        today = date.today().isoformat()
        email_map = _build_staff_email_map()

        team_rows = []
        for item in enriched:
            if scope_codes is not None:
                tid = item.get("team_id")
                hid = item.get("handling_team_id") or tid
                if tid not in scope_codes and hid not in scope_codes:
                    continue
            staff_emails = _receipt_staff_emails(item, email_map=email_map)
            notice_sent = _receipt_notice_sent(item)
            severity = item.get("severity") or "medium"
            esc = v32.normalize_escalation(item)
            try:
                can_reassign = _can_reassign_as_staff(item, person)
            except Exception:
                can_reassign = False
            row = {
                **item,
                "customer_display": _dashboard_customer_label(item, team_code),
                "assignee_display": item.get("as_display") or "미지정",
                "status_class": _dashboard_status_class(item.get("status")),
                "notice_sent": notice_sent,
                "can_send_email": bool(staff_emails) and not notice_sent,
                "severity": severity,
                "severity_label": _severity_label(severity),
                "is_high_severity": _is_high_severity(severity),
                "escalation_level": esc,
                "escalation_label": dict(v32.ESCALATION_OPTIONS).get(esc, esc),
                "intervention_label": _intervention_label(item.get("intervention_status")),
                "compensation_label": _compensation_label(item.get("compensation_status")),
                "can_claim_as": _can_claim_unassigned_as(item, person),
                "can_reassign_as": can_reassign,
            }
            team_rows.append(row)

    today = date.today().isoformat()

    try:
        team_rows = _filter_team_rows_for_domestic(team_rows, person)
    except Exception:
        pass
    try:
        team_rows = _filter_rows_for_field_support(team_rows, person)
    except Exception:
        pass

    if team_code == "field_support":
        try:
            kind = _support_request_kind(person)
        except Exception:
            kind = ""
        if kind == "quality":
            team_name = "품질팀 지원요청"
        elif kind == "tech":
            team_name = "기술연구소 지원요청" if _is_rods() else "기술부 지원요청"
        else:
            team_name = "지원요청"

    my_receipts = []
    my_as_history = []
    # 부서장은 본인 접수목록 + 전사 현황
    if is_sales_head or not is_all_viewer:
        for item in enriched:
            if not _is_my_receipt(item, person, user_id):
                continue
            try:
                my_can_reassign = _can_reassign_as_staff(item, person)
            except Exception:
                my_can_reassign = False
            my_receipts.append(
                {
                    **item,
                    "customer_display": item.get("customer_name") or "",
                    "status_class": _dashboard_status_class(item.get("status")),
                    "can_claim_as": _can_claim_unassigned_as(item, person),
                    "can_reassign_as": my_can_reassign,
                }
            )
        my_receipts.sort(key=v32.recent_board_sort_key)

    if not is_all_viewer:
        for item in enriched:
            receipt_no = item.get("receipt_no")
            if not receipt_no or receipt_no not in REPORT_DETAILS:
                continue
            if not _is_my_as_work(item, person, user_id):
                continue
            my_as_history.append(
                {
                    **item,
                    "customer_display": item.get("customer_name") or "",
                    "status_class": _dashboard_status_class(item.get("status")),
                }
            )
        my_as_history.sort(key=v32.recent_board_sort_key)

    my_as_open = sum(
        1 for r in my_as_history if v32.is_pending(r.get("status"))
    )
    my_as_done = sum(
        1 for r in my_as_history if r.get("status") == "completed"
    )

    stats = _dashboard_stats_for_rows(team_rows, today)
    high_severity_count = sum(
        1
        for r in team_rows
        if r.get("is_high_severity")
        or _is_high_severity(r.get("severity") or "medium")
    )
    intervening_count = sum(
        1 for r in team_rows if (r.get("intervention_status") or "none") == "intervening"
    )
    negotiating_count = sum(
        1
        for r in team_rows
        if (r.get("compensation_status") or "none") == "negotiating"
    )

    team_breakdown = []
    if is_all_viewer:
        breakdown_source = team_rows if is_tech_mgmt else enriched
        for team in _receipt_dashboard_teams():
            subset = [r for r in breakdown_source if r.get("team_id") == team["id"]]
            team_breakdown.append(
                {
                    "team_code": team["id"],
                    "team_name": team["name"],
                    "stats": _dashboard_stats_for_rows(subset, today),
                }
            )

    auto_flagged = []
    auto_cleared_count = 0
    # 6.13: 정확도 통계는 최고기술경영실 전용 — 타 권한에는 데이터 자체를 넣지 않음
    auto_accuracy = None
    monitor_rows = []
    monitor_priority_total = 0
    panels_lazy = False
    if is_tech_mgmt:
        # 홈 첫 화면은 통계만 즉시 표시. 카드/목록은 API 지연 로딩
        panels_lazy = True
        try:
            _cards, auto_cleared_count, auto_accuracy = _build_auto_detect_cards(
                card_limit=0
            )
        except TypeError:
            auto_flagged, auto_cleared_count, auto_accuracy = _build_auto_detect_cards()
            auto_flagged = []
        except Exception:
            auto_accuracy = {
                "total_flagged": 0,
                "confirmed": 0,
                "false_positive": 0,
                "accuracy": None,
                "unreviewed": 0,
            }
        for r in team_rows:
            if v32.normalize_escalation(r) in ("monitor", "critical"):
                monitor_priority_total += 1
        auto_flagged = []
        monitor_rows = []
    # 경영진·임원·영업부 부서장: 최근 접수 현황
    # 접수·진행중=접수일 과거순, 완료=접수일 최신순
    recent_receipts = []
    recent_stats = None
    if (is_executive or is_sales_head) and not is_tech_mgmt:
        week_ago = (date.today() - timedelta(days=6)).isoformat()
        month_key = date.today().strftime("%Y-%m")
        recent_stats = {
            "today": stats["today"],
            "week": sum(
                1
                for r in team_rows
                if (r.get("receipt_date") or "")[:10] >= week_ago
            ),
            "month": sum(
                1
                for r in team_rows
                if (r.get("receipt_date") or "").startswith(month_key)
            ),
            "pending": sum(1 for r in team_rows if _is_pending_receipt(r)),
            "completed_week": sum(
                1
                for r in team_rows
                if r.get("status") == "completed"
                and (r.get("completed_date") or "")[:10] >= week_ago
            ),
            "today_key": today,
            "week_ago": week_ago,
            "month_key": month_key,
            "preview_limit": 25,
        }
        # 카드 필터(금일·7일·당월·미처리·7일완료)와 건수가 맞도록 전체 행을 내려보냄
        recent_receipts = sorted(team_rows, key=v32.recent_board_sort_key)

    # 홈 테이블: 경영진은 상위 N건, 최고기술경영실은 탭 진입 시 API 지연 로딩
    HOME_LIST_LIMIT = 100
    team_receipts_total = len(team_rows)
    team_receipts_lazy = False
    if is_tech_mgmt:
        team_receipts_view = []
        team_receipts_lazy = True
    else:
        team_receipts_view = team_rows
        if is_all_viewer and team_receipts_total > HOME_LIST_LIMIT:
            team_receipts_view = team_rows[:HOME_LIST_LIMIT]
        monitor_priority_total = len(monitor_rows)

    queue_summary = None
    if (
        not is_tech_mgmt
        and not (is_executive and not is_sales_head)
        and not _is_office_person(person)
    ):
        queue_summary = _compute_queue_summary(team_rows, person)

    return {
        "user_name": user_name,
        "team_name": team_name,
        "team_code": team_code,
        "account_label": account_label,
        "is_executive": is_executive,
        "is_tech_mgmt": is_tech_mgmt,
        "is_sales_head": is_sales_head,
        "is_all_viewer": is_all_viewer and not is_sales_head,
        "stats": {
            **stats,
            "high_severity": high_severity_count,
            "intervening": intervening_count,
            "negotiating": negotiating_count,
        },
        "team_breakdown": team_breakdown,
        "team_receipts": team_receipts_view,
        "team_receipts_total": team_receipts_total,
        "team_receipts_limited": (not team_receipts_lazy)
        and team_receipts_total > len(team_receipts_view),
        "team_receipts_lazy": team_receipts_lazy,
        "panels_lazy": panels_lazy,
        "home_teams": _receipt_dashboard_teams(),
        "home_status_options": plex.status_options(),
        "my_receipts": my_receipts,
        "is_office_account": _is_office_person(person),
        "my_receipts_title": (
            "우리 사무소 접수 AS 목록" if _is_office_person(person) else "내가 접수한 AS 목록"
        ),
        "my_receipts_tab_label": (
            "우리 사무소 접수 AS" if _is_office_person(person) else "내가 접수한 AS"
        ),
        "my_receipts_completed": sum(
            1 for r in my_receipts if (r.get("status") or "") == "completed"
        ),
        "my_receipts_pending": sum(
            1 for r in my_receipts if v32.is_pending(r.get("status"))
        ),
        "my_as_history": my_as_history,
        "my_as_open_count": my_as_open,
        "my_as_done_count": my_as_done,
        "mail_enabled": mail_configured(),
        "auto_flagged": auto_flagged[:24],
        "auto_cleared_count": auto_cleared_count,
        "auto_accuracy": auto_accuracy,
        "auto_flagged_total": (
            (auto_accuracy or {}).get("unreviewed")
            if auto_accuracy and auto_accuracy.get("unreviewed") is not None
            else len(auto_flagged)
        ),
        "monitor_priority": monitor_rows,
        "monitor_priority_total": monitor_priority_total,
        "escalation_options": v32.ESCALATION_OPTIONS,
        "recent_receipts": recent_receipts,
        "recent_stats": recent_stats,
        "queue_summary": queue_summary,
        "today": today,
        "as_reassign_options": [
            {
                "id": p.get("id") or "",
                "label": p.get("display_name") or p.get("name") or "",
                "team_code": p.get("team_code") or "",
            }
            for p in _as_assignable_people(person)
            if p.get("id")
        ],
        "export_as_allowed_codes": list(EXPORT_AS_ALLOWED_CODES),
    }


@app.route("/")
def index():
    session.clear()
    return render_template(
        "select_product.html",
        active_page="select",
    )


@app.route("/home")
def home():
    _enter_product_line("welding")
    dashboard = None
    if session.get("user_id"):
        dashboard = _build_home_dashboard()
    return render_template(
        "index.html",
        active_page="home",
        dashboard=dashboard,
    )


@app.route("/switch-product")
def switch_product():
    session.clear()
    return redirect(url_for("index"))


@app.route("/api/home/team-receipts")
@_login_required
def api_home_team_receipts():
    """최고기술경영실 홈 — 팀 AS 목록 지연 로딩용 경량 JSON."""
    person = _get_person_for_session()
    if not _is_tech_mgmt_person(person):
        return jsonify({"error": "권한 없음"}), 403

    email_map = _build_staff_email_map()
    esc_labels = dict(v32.ESCALATION_OPTIONS)
    items = []
    for receipt in sorted(_line_receipts(), key=v32.recent_board_sort_key):
        esc = v32.normalize_escalation(receipt)
        emails = _receipt_staff_emails(receipt, email_map=email_map)
        notice_sent = _receipt_notice_sent(receipt)
        status = receipt.get("status") or ""
        symptoms = (receipt.get("symptoms") or "")[:100]
        assignee_raw = receipt.get("as_staff_name") or "미지정"
        assignee_main, assignee_sub = _split_assignee_display(assignee_raw)
        items.append(
            {
                "receipt_no": receipt.get("receipt_no") or "",
                "receipt_date": receipt.get("receipt_date") or "",
                "team": receipt.get("team") or "",
                "customer_name": receipt.get("customer_name") or "",
                "equipment_name": receipt.get("equipment_name") or "",
                "symptoms": symptoms,
                "assignee": assignee_raw,
                "assignee_main": assignee_main,
                "assignee_sub": assignee_sub,
                "status": status,
                "status_label": v32.status_label(status),
                "status_class": _dashboard_status_class(status),
                "escalation_level": esc,
                "escalation_label": esc_labels.get(esc, esc),
                "notice_sent": notice_sent,
                "can_send_email": bool(emails) and not notice_sent,
            }
        )
    return jsonify(
        {
            "items": items,
            "total": len(items),
            "mail_enabled": mail_configured(),
            "csrf_token": session.get("csrf_token") or "",
        }
    )


def _home_list_csv_payload(kind):
    """홈 목록 CSV용 행·파일명 구성. kind: team | my-receipts | my-history."""
    person = _get_person_for_session()
    user_id = session.get("user_id")
    team_code = person["team_code"] if person else ""
    is_executive = _is_executive_person(person)
    is_tech_mgmt = _is_tech_mgmt_person(person)
    is_sales_head = False
    try:
        is_sales_head = _is_sales_head(person)
    except Exception:
        is_sales_head = team_code == "sales_head"
    is_all_viewer = is_executive or is_tech_mgmt or is_sales_head
    scope_codes = _team_scope_codes(team_code, person)

    q = (request.args.get("q") or "").strip().casefold()
    status_filter = (request.args.get("status") or "").strip()
    team_filter = (request.args.get("team") or "").strip()

    def match_filters(row, *, search_parts):
        if status_filter and (row.get("status") or "") != status_filter:
            return False
        if team_filter and (row.get("team") or "") != team_filter:
            return False
        if q:
            hay = " ".join(str(p or "") for p in search_parts).casefold()
            if q not in hay:
                return False
        return True

    if kind == "team":
        line_rows = _line_receipts()
        if is_tech_mgmt:
            # 지연 로딩 계정: 현재 제품군 접수만 사용
            source = []
            for r in line_rows:
                if scope_codes is not None and r.get("team_id") not in scope_codes:
                    continue
                source.append(r)
        else:
            skip_queue = True
            enriched = [
                _enrich_receipt(r, line_rows, skip_queue=skip_queue) for r in line_rows
            ]
            source = []
            for item in enriched:
                if scope_codes is not None and item.get("team_id") not in scope_codes:
                    continue
                source.append(item)
            try:
                source = _filter_team_rows_for_domestic(source, person)
            except Exception:
                pass
            try:
                source = _filter_rows_for_field_support(source, person)
            except Exception:
                pass
        source = sorted(source, key=v32.recent_board_sort_key)
        fields = [
            "접수번호",
            "접수일",
            "AS 대응팀",
            "고객명",
            "장비명",
            "증상",
            "담당자",
            "상태",
        ]
        rows = []
        for r in source:
            status = r.get("status") or ""
            assignee = r.get("as_display") or r.get("as_staff_name") or "미지정"
            customer = r.get("customer_display") or r.get("customer_name") or ""
            if not match_filters(
                r,
                search_parts=[
                    r.get("receipt_no"),
                    customer,
                    r.get("equipment_name"),
                    r.get("symptoms"),
                    assignee,
                    r.get("team"),
                    v32.status_label(status),
                ],
            ):
                continue
            rows.append(
                {
                    "접수번호": r.get("receipt_no") or "",
                    "접수일": r.get("receipt_date") or "",
                    "AS 대응팀": r.get("team") or "",
                    "고객명": customer,
                    "장비명": r.get("equipment_name") or "",
                    "증상": r.get("symptoms") or "",
                    "담당자": assignee,
                    "상태": v32.status_label(status),
                }
            )
        return rows, fields, "home_team_as"

    # my-receipts / my-history: 전사 전용 계정(경영진·최고기술)은 제외
    if is_all_viewer and not is_sales_head:
        return None, None, None

    line_rows = _line_receipts()
    enriched = [_enrich_receipt(r, line_rows, skip_queue=False) for r in line_rows]
    enriched.sort(key=v32.recent_board_sort_key)

    if kind == "my-receipts":
        fields = [
            "접수번호",
            "접수일",
            "고객명",
            "장비명",
            "증상",
            "AS 대응팀",
            "접수담당자",
            "접수사무소",
            "대기순번",
            "상태",
        ]
        rows = []
        for item in enriched:
            if not _is_my_receipt(item, person, user_id):
                continue
            status = item.get("status") or ""
            if not match_filters(
                item,
                search_parts=[
                    item.get("receipt_no"),
                    item.get("customer_name"),
                    item.get("equipment_name"),
                    item.get("symptoms"),
                    item.get("team"),
                    item.get("receipt_staff_name"),
                    item.get("office"),
                    item.get("queue_text"),
                    v32.status_label(status),
                ],
            ):
                continue
            rows.append(
                {
                    "접수번호": item.get("receipt_no") or "",
                    "접수일": item.get("receipt_date") or "",
                    "고객명": item.get("customer_name") or "",
                    "장비명": item.get("equipment_name") or "",
                    "증상": item.get("symptoms") or "",
                    "AS 대응팀": item.get("team") or "",
                    "접수담당자": item.get("receipt_staff_name") or "",
                    "접수사무소": item.get("office") or "",
                    "대기순번": item.get("queue_text") or "",
                    "상태": v32.status_label(status),
                }
            )
        return rows, fields, "home_my_receipts"

    if kind == "my-history":
        fields = [
            "접수번호",
            "접수일",
            "고객명",
            "장비명",
            "증상",
            "AS 대응팀",
            "상태",
        ]
        rows = []
        for item in enriched:
            receipt_no = item.get("receipt_no")
            if not receipt_no or receipt_no not in REPORT_DETAILS:
                continue
            if not _is_my_as_work(item, person, user_id):
                continue
            status = item.get("status") or ""
            if not match_filters(
                item,
                search_parts=[
                    item.get("receipt_no"),
                    item.get("customer_name"),
                    item.get("equipment_name"),
                    item.get("symptoms"),
                    item.get("team"),
                    v32.status_label(status),
                ],
            ):
                continue
            rows.append(
                {
                    "접수번호": item.get("receipt_no") or "",
                    "접수일": item.get("receipt_date") or "",
                    "고객명": item.get("customer_name") or "",
                    "장비명": item.get("equipment_name") or "",
                    "증상": item.get("symptoms") or "",
                    "AS 대응팀": item.get("team") or "",
                    "상태": v32.status_label(status),
                }
            )
        return rows, fields, "home_my_as_history"

    return None, None, None


@app.route("/api/home/export.csv")
@_login_required
def api_home_export_csv():
    """홈 — 팀 AS / 내가 접수한 AS / 내 AS 처리 이력 CSV."""
    kind = (request.args.get("kind") or "").strip().lower()
    if kind not in ("team", "my-receipts", "my-history"):
        return "잘못된 요청입니다.", 400

    rows, fields, prefix = _home_list_csv_payload(kind)
    if rows is None:
        return "권한이 없습니다.", 403

    buf = io.StringIO()
    writer = csv.DictWriter(
        buf,
        fieldnames=fields,
        extrasaction="ignore",
        lineterminator="\n",
    )
    writer.writeheader()
    writer.writerows(rows)

    stamp = date.today().strftime("%Y%m%d")
    filename = f"{prefix}_{stamp}.csv"
    payload = buf.getvalue().encode("utf-8-sig")
    return Response(
        payload,
        mimetype="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )


@app.route("/api/home/tech-cards")
@_login_required
def api_home_tech_cards():
    """최고기술경영실 홈 — 심각·모니터링 / 자동감지 카드 지연 로딩."""
    person = _get_person_for_session()
    if not _is_tech_mgmt_person(person):
        return jsonify({"error": "권한 없음"}), 403

    kind = (request.args.get("kind") or "").strip().lower()
    if kind == "priority":
        esc_labels = dict(v32.ESCALATION_OPTIONS)
        items = []
        for receipt in _line_receipts():
            esc = v32.normalize_escalation(receipt)
            if esc not in ("monitor", "critical"):
                continue
            items.append(
                {
                    "receipt_no": receipt.get("receipt_no") or "",
                    "customer_name": receipt.get("customer_name") or "",
                    "equipment_name": receipt.get("equipment_name") or "",
                    "team": receipt.get("team") or "",
                    "escalation_level": esc,
                    "escalation_label": esc_labels.get(esc, esc),
                }
            )
            if len(items) >= 60:
                break
        return jsonify({"kind": "priority", "items": items, "total": len(items)})

    if kind == "auto":
        try:
            flagged, cleared, _acc = _build_auto_detect_cards(card_limit=24)
        except TypeError:
            flagged, cleared, _acc = _build_auto_detect_cards()
            flagged = flagged[:24]
        except Exception:
            flagged, cleared = [], 0
        items = []
        for r in flagged:
            items.append(
                {
                    "receipt_no": r.get("receipt_no") or "",
                    "customer_name": r.get("customer_name") or "",
                    "equipment_name": r.get("equipment_name") or "",
                    "team": r.get("team") or "",
                    "status_label": r.get("status_label") or "",
                    "days_open": r.get("days_open"),
                    "auto_flag_labels": list(r.get("auto_flag_labels") or []),
                    "escalation_level": r.get("escalation_level") or "normal",
                }
            )
        return jsonify(
            {
                "kind": "auto",
                "items": items,
                "total": len(items),
                "cleared_count": cleared,
            }
        )

    return jsonify({"error": "invalid kind"}), 400


@app.route("/status-report")
@_executive_required
def status_report():
    year_param = request.args.get("year", "").strip()
    period = (request.args.get("period") or "month").strip().lower()
    if period not in ("month", "quarter", "year"):
        period = "month"
    try:
        year = int(year_param) if year_param else date.today().year
    except ValueError:
        year = date.today().year

    report = _build_status_report(year, period)
    return render_template(
        "status_report.html",
        active_page="status_report",
        report=report,
    )


def _csv_period_parts(value):
    """날짜 문자열 → (연도, 분기, 월키, 일자) — 피벗용."""
    text = (value or "").strip()[:10]
    if not text:
        return "", "", "", ""
    try:
        d = datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return "", "", "", ""
    quarter = (d.month - 1) // 3 + 1
    return str(d.year), f"Q{quarter}", f"{d.year}-{d.month:02d}", d.isoformat()


def _csv_yes_no(flag):
    return "Y" if flag else "N"


def _csv_re_repair_label(value):
    text = (value or "").strip()
    if text in ("Y", "1회", "예"):
        return "Y"
    if text in ("N", "아니오", "해당없음"):
        return "N"
    return text or ""


def _csv_parts_detail(parts):
    rows = []
    for p in parts or []:
        name = (p.get("name") or "").strip()
        if not name:
            continue
        qty = p.get("qty") or 0
        unit = p.get("unit_price") or 0
        sub = p.get("subtotal")
        if sub is None:
            try:
                sub = int(qty) * int(unit)
            except (TypeError, ValueError):
                sub = 0
        spec = (p.get("spec") or "").strip()
        spec_part = f"|규격:{spec}" if spec else ""
        rows.append(f"{name}{spec_part}|수량:{qty}|단가:{unit}|소계:{sub}")
    return "; ".join(rows)


def _build_history_csv_rows(year=None, all_lines=False):
    """경영진 CSV용 전체(또는 연도 필터) 이력 행 목록."""
    year_text = str(year) if year else ""
    status_map = {
        "received": "접수",
        "draft": "임시저장",
        "in_progress": "진행중",
        "parts_waiting": "부품대기",
        "on_hold": "보류",
        "completed": "완료",
    }
    rows = []
    source = RECEIPTS if all_lines else _line_receipts()
    for receipt in sorted(
        source,
        key=lambda r: (r.get("receipt_date") or "", r.get("receipt_no") or ""),
    ):
        receipt_date = receipt.get("receipt_date") or ""
        if year_text and not receipt_date.startswith(year_text):
            continue

        detail = REPORT_DETAILS.get(receipt.get("receipt_no") or "", {}) or {}
        parts = detail.get("parts") or []
        parts_total = v32.calc_parts_total(parts)
        ship_date = detail.get("ship_date") or receipt.get("ship_date") or ""
        inspect_date = detail.get("inspection_date") or ""
        warranty = v32.warranty_verdict(ship_date, inspect_date or None)
        field_photos = detail.get("photos") or []
        customer_photos = receipt.get("photos") or []

        ry, rq, rm, rd = _csv_period_parts(receipt_date)
        cy, cq, cm, cd = _csv_period_parts(receipt.get("completed_date") or "")
        iy, iq, im, idate = _csv_period_parts(inspect_date)

        try:
            billing_cost = int(receipt["billing_cost"]) if receipt.get("billing_cost") not in (None, "") else ""
        except (TypeError, ValueError):
            billing_cost = receipt.get("billing_cost") or ""

        margin_rate = _effective_margin_rate(receipt.get("billing_margin_rate"))

        esc = v32.normalize_escalation(receipt)
        rows.append(
            {
                "접수번호": receipt.get("receipt_no") or "",
                "접수연도": ry,
                "접수분기": rq,
                "접수월": rm,
                "접수일": rd,
                "완료연도": cy,
                "완료분기": cq,
                "완료월": cm,
                "완료일": cd,
                "점검연도": iy,
                "점적분기": iq,
                "점검월": im,
                "점검일": idate,
                "팀코드": receipt.get("team_id") or "",
                "팀명": receipt.get("team") or "",
                "고객명": receipt.get("customer_name") or "",
                "고객담당자": receipt.get("customer_contact") or "",
                "연락처": receipt.get("phone") or detail.get("phone") or "",
                "장비명": receipt.get("equipment_name") or "",
                "규격": receipt.get("spec") or "",
                "로트번호": receipt.get("lot") or "",
                "제품군": receipt.get("product_line") or "equipment",
                "시리얼번호": detail.get("serial_no") or "",
                "증상": receipt.get("symptoms") or "",
                "유무상": "무상" if receipt.get("charge_type") == "free" else "유상",
                "유무상코드": receipt.get("charge_type") or "",
                "청구비용": billing_cost,
                "청구마진율": margin_rate,
                "청구수동책정": _csv_yes_no(int(receipt.get("billing_manual") or 0)),
                "청구책정자": receipt.get("billing_set_by") or "",
                "부품건수": len(parts),
                "부품합계금액": parts_total,
                "부품내역": _csv_parts_detail(parts),
                "워런티코드": warranty.get("code") or "",
                "워런티": warranty.get("label") or "",
                "워런티경과일": "" if warranty.get("days") is None else warranty.get("days"),
                "출하일": ship_date,
                "재수리여부": _csv_re_repair_label(detail.get("re_repair")),
                "AS상태코드": receipt.get("status") or "",
                "AS상태": status_map.get(receipt.get("status"), receipt.get("status") or ""),
                "접수담당자": receipt.get("receipt_staff_name") or detail.get("receipt_staff_name") or "",
                "접수사무소": receipt.get("office") or "",
                "AS배정담당자": receipt.get("as_staff_name") or "",
                "AS처리담당자": detail.get("as_technician") or "",
                "소속": detail.get("affiliation") or receipt.get("team") or "",
                "처리위치": detail.get("process_location") or "",
                "출고일": detail.get("ship_out_date") or "",
                "점검결과": detail.get("inspection_result") or "",
                "수리내용": detail.get("repair_content") or "",
                "상세작업내용": detail.get("detailed_work") or "",
                "비고": detail.get("remarks") or "",
                "지원요청이력": detail.get("support_history") or "",
                "품질지원요청": _csv_yes_no(receipt.get("support_quality") or detail.get("support_quality")),
                "기술지원요청": _csv_yes_no(receipt.get("support_tech") or detail.get("support_tech")),
                "최고기술경영실보고": _csv_yes_no(detail.get("support_tech_mgmt")),
                "에스컬레이션": dict(v32.ESCALATION_OPTIONS).get(esc, esc),
                "개입상태": _intervention_label(receipt.get("intervention_status")),
                "보상상태": _compensation_label(receipt.get("compensation_status")),
                "보상금액": receipt.get("compensation_amount") or "",
                "고객전달사진수": len(customer_photos or []),
                "현장사진수": len(field_photos or []),
                "보고서작성일": detail.get("report_date") or "",
                "등록일시": receipt.get("created_at") or "",
            }
        )
    return rows


HISTORY_CSV_FIELDS = [
    "접수번호",
    "접수연도",
    "접수분기",
    "접수월",
    "접수일",
    "완료연도",
    "완료분기",
    "완료월",
    "완료일",
    "점검연도",
    "점적분기",
    "점검월",
    "점검일",
    "팀코드",
    "팀명",
    "고객명",
    "고객담당자",
    "연락처",
    "장비명",
    "시리얼번호",
    "증상",
    "유무상",
    "유무상코드",
    "청구비용",
    "청구마진율",
    "청구수동책정",
    "청구책정자",
    "부품건수",
    "부품합계금액",
    "부품내역",
    "워런티코드",
    "워런티",
    "워런티경과일",
    "출하일",
    "재수리여부",
    "AS상태코드",
    "AS상태",
    "접수담당자",
    "접수사무소",
    "AS배정담당자",
    "AS처리담당자",
    "소속",
    "처리위치",
    "출고일",
    "점검결과",
    "수리내용",
    "상세작업내용",
    "비고",
    "지원요청이력",
    "품질지원요청",
    "기술지원요청",
    "최고기술경영실보고",
    "에스컬레이션",
    "개입상태",
    "보상상태",
    "보상금액",
    "고객전달사진수",
    "현장사진수",
    "보고서작성일",
    "등록일시",
]


@app.route("/status-report/export.csv")
@_executive_required
def status_report_export_csv():
    """경영진·현황보고 권한: 전체(또는 연도) AS 이력 CSV 다운로드."""
    year_param = request.args.get("year", "").strip().lower()
    year = None
    if year_param and year_param not in ("all", "전체"):
        try:
            year = int(year_param)
        except ValueError:
            year = None

    rows = _build_history_csv_rows(year=year)
    buf = io.StringIO()
    writer = csv.DictWriter(
        buf,
        fieldnames=HISTORY_CSV_FIELDS,
        extrasaction="ignore",
        lineterminator="\n",
    )
    writer.writeheader()
    writer.writerows(rows)

    stamp = date.today().strftime("%Y%m%d")
    scope = f"y{year}" if year else "full"
    filename = f"AS_history_{scope}_{stamp}.csv"
    payload = buf.getvalue().encode("utf-8-sig")
    return Response(
        payload,
        mimetype="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )


def _tech_monitor_case_row(receipt, *, auto_labels=None, is_auto_candidate=False):
    """개입·보상 카드용 경량 행 (enrich/대기열 계산 없음)."""
    esc = v32.normalize_escalation(receipt)
    severity = receipt.get("severity") or v32.ESCALATION_TO_SEVERITY.get(esc, "medium")
    logs = receipt.get("intervention_log") or []
    if isinstance(logs, str):
        logs = _parse_json_list(logs)
    return {
        "receipt_no": receipt.get("receipt_no") or "",
        "receipt_date": receipt.get("receipt_date") or "",
        "team": receipt.get("team") or "",
        "team_id": receipt.get("team_id") or "",
        "customer_name": receipt.get("customer_name") or "",
        "equipment_name": receipt.get("equipment_name") or "",
        "symptoms": receipt.get("symptoms") or "",
        "status": receipt.get("status") or "",
        "status_label": v32.status_label(receipt.get("status")),
        "status_class": _dashboard_status_class(receipt.get("status")),
        "assignee_display": receipt.get("as_staff_name") or "미지정",
        "severity": severity,
        "severity_label": _severity_label(severity),
        "is_high_severity": esc in ("monitor", "critical"),
        "is_auto_candidate": bool(is_auto_candidate),
        "intervention_status": receipt.get("intervention_status") or "none",
        "intervention_label": _intervention_label(receipt.get("intervention_status")),
        "compensation_status": receipt.get("compensation_status") or "none",
        "compensation_label": _compensation_label(receipt.get("compensation_status")),
        "escalation_level": esc,
        "escalation_label": dict(v32.ESCALATION_OPTIONS).get(esc, esc),
        "compensation_amount": receipt.get("compensation_amount") or "",
        "intervention_note": receipt.get("intervention_note") or "",
        "compensation_note": receipt.get("compensation_note") or "",
        "intervention_log": logs,
        "intervention_by": receipt.get("intervention_by") or "",
        "intervention_at": receipt.get("intervention_at"),
        "created_at": receipt.get("created_at") or "",
        "auto_flag_labels": list(auto_labels or []),
    }


def _build_tech_monitor(
    filter_escalation="all",
    filter_team="",
    focus_no="",
    q="",
    filter_status="",
):
    """최고기술경영실용 전 팀 모니터링·개입 목록 (에스컬레이션 기준)."""
    legacy_map = {
        "high": "active",
        "critical": "critical",
        "all": "all",
        "low": "normal",
        "medium": "normal",
    }
    if filter_escalation in legacy_map:
        filter_escalation = legacy_map[filter_escalation]
    if filter_escalation not in (
        "active",
        "all",
        "normal",
        "monitor",
        "critical",
        "cleared",
        "closed",
    ):
        filter_escalation = "all"

    focus_no = (focus_no or "").strip()
    q = (q or "").strip().casefold()
    filter_status = (filter_status or "").strip()
    allowed_status = {
        "received",
        "draft",
        "in_progress",
        "parts_waiting",
        "on_hold",
        "completed",
    }
    if filter_status and filter_status not in allowed_status:
        filter_status = ""

    line_rows = _line_receipts()
    stats = {
        "total": len(line_rows),
        "monitor": 0,
        "critical": 0,
        "high_severity": 0,
        "intervening": 0,
        "negotiating": 0,
        "watching": 0,
    }
    matched = []
    focus_row = None

    for receipt in line_rows:
        esc = v32.normalize_escalation(receipt)
        if esc == "monitor":
            stats["monitor"] += 1
            stats["high_severity"] += 1
        elif esc == "critical":
            stats["critical"] += 1
            stats["high_severity"] += 1
        if (receipt.get("intervention_status") or "none") == "intervening":
            stats["intervening"] += 1
        if (receipt.get("compensation_status") or "none") == "negotiating":
            stats["negotiating"] += 1
        if (receipt.get("intervention_status") or "none") == "watching":
            stats["watching"] += 1

        no = receipt.get("receipt_no") or ""
        if focus_no and no == focus_no:
            focus_row = _tech_monitor_case_row(receipt)

        if filter_team and receipt.get("team_id") != filter_team:
            continue
        if filter_status and (receipt.get("status") or "") != filter_status:
            continue
        if filter_escalation == "active":
            if esc not in ("monitor", "critical"):
                continue
        elif filter_escalation != "all" and esc != filter_escalation:
            continue

        row = _tech_monitor_case_row(receipt)
        if q:
            hay = " ".join(
                [
                    str(row.get("receipt_no") or ""),
                    str(row.get("customer_name") or ""),
                    str(row.get("equipment_name") or ""),
                    str(row.get("symptoms") or ""),
                    str(row.get("assignee_display") or ""),
                    str(row.get("team") or ""),
                    str(row.get("escalation_label") or ""),
                    str(row.get("status_label") or ""),
                ]
            ).casefold()
            if q not in hay:
                continue
        matched.append(row)

    def _esc_sort_key(r):
        rank = {"critical": 0, "monitor": 1}.get(r.get("escalation_level"), 2)
        pending = 0 if v32.is_pending(r.get("status")) else 1
        return (rank, pending, r.get("created_at") or "")

    matched.sort(key=_esc_sort_key)

    # 편집 폼은 포커스 1건만 (전체 목록 렌더 부하 방지)
    detail_rows = [focus_row] if focus_row else []

    return {
        "rows": detail_rows,
        "list_rows": matched,
        "focus_no": focus_no,
        "q": q,
        "filter_status": filter_status,
        "total_matched": len(matched),
        "filter_escalation": filter_escalation,
        "filter_severity": filter_escalation,
        "filter_team": filter_team,
        "teams": _receipt_dashboard_teams(),
        "status_options": [
            ("received", "접수"),
            ("in_progress", "진행중"),
            ("parts_waiting", "부품대기"),
            ("on_hold", "보류"),
            ("completed", "완료"),
        ],
        "severity_options": SEVERITY_OPTIONS,
        "escalation_options": v32.ESCALATION_OPTIONS,
        "intervention_options": INTERVENTION_OPTIONS,
        "compensation_options": v32.COMPENSATION_V32,
        "stats": stats,
    }


@app.route("/tech-monitor")
@_tech_mgmt_required
def tech_monitor():
    filter_escalation = (
        request.args.get("escalation")
        or request.args.get("severity")
        or "all"
    ).strip() or "all"
    filter_team = request.args.get("team", "").strip()
    filter_status = request.args.get("status", "").strip()
    q = request.args.get("q", "").strip()
    focus_no = (request.args.get("focus") or request.args.get("receipt_no") or "").strip()
    monitor = _build_tech_monitor(
        filter_escalation,
        filter_team,
        focus_no=focus_no,
        q=q,
        filter_status=filter_status,
    )
    auto_flagged = []
    auto_cleared_count = 0
    auto_accuracy = {
        "total_flagged": 0,
        "confirmed": 0,
        "false_positive": 0,
        "accuracy": None,
    }
    try:
        auto_flagged, auto_cleared_count, auto_accuracy = _build_auto_detect_cards()
    except Exception:
        pass

    # 자동감지 후보는 카드 요약만 (무거운 편집 폼으로 병합하지 않음)
    monitor["auto_flagged"] = auto_flagged[:12]
    monitor["auto_flagged_total"] = len(auto_flagged)
    monitor["auto_cleared_count"] = auto_cleared_count
    monitor["auto_accuracy"] = auto_accuracy
    monitor["auto_rules"] = [
        ("repeat_equipment", "동일장비 반복접수", "같은 고객사·장비 누적 ≥ 3회"),
        ("high_billing", "고액청구", "청구 ≥ 300만원 또는 부품합계 ≥ 200만원 (유상)"),
        ("long_open", "장기미해결", "접수 후 14일 이상 미완료"),
        ("many_parts", "부품다량교체", "교체 부품 종류 ≥ 4종"),
        ("repeat_customer", "동일고객 집중접수", "최근 90일 내 동일 고객 ≥ 3건"),
    ]
    return render_template(
        "tech_monitor.html",
        active_page="tech_monitor",
        monitor=monitor,
        csrf_token=session.get("csrf_token"),
    )


@app.route("/tech-monitor/<receipt_no>/update", methods=["POST"])
@_tech_mgmt_required
@_service_write_locked
def tech_monitor_update(receipt_no):
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("tech_monitor"))

    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        flash("접수 건을 찾을 수 없습니다.", "error")
        return redirect(url_for("tech_monitor"))

    severity = request.form.get("severity", receipt.get("severity") or "medium").strip()
    intervention = request.form.get(
        "intervention_status", receipt.get("intervention_status") or "none"
    ).strip()
    compensation = request.form.get(
        "compensation_status", receipt.get("compensation_status") or "none"
    ).strip()
    allowed_sev = {k for k, _ in SEVERITY_OPTIONS}
    allowed_int = {k for k, _ in INTERVENTION_OPTIONS}
    allowed_comp = {k for k, _ in COMPENSATION_OPTIONS}
    if severity not in allowed_sev:
        severity = "medium"
    if intervention not in allowed_int:
        intervention = "none"
    if compensation not in allowed_comp:
        compensation = "none"

    receipt["severity"] = severity
    receipt["intervention_status"] = intervention
    receipt["compensation_status"] = compensation
    receipt["intervention_note"] = (request.form.get("intervention_note") or "").strip()
    receipt["compensation_note"] = (request.form.get("compensation_note") or "").strip()
    receipt["compensation_amount"] = (request.form.get("compensation_amount") or "").strip()
    old_esc = v32.normalize_escalation(receipt)
    esc = (request.form.get("escalation_level") or old_esc).strip()
    if esc not in v32.ESCALATION_RANK:
        esc = old_esc
    receipt["escalation_level"] = esc
    # 에스컬레이션이 주 등급 — severity 동기화
    receipt["severity"] = v32.ESCALATION_TO_SEVERITY.get(esc, severity)
    if esc == "critical" and receipt["intervention_status"] in ("none", "watching"):
        receipt["intervention_status"] = "intervening"
    if esc == "monitor" and receipt["intervention_status"] == "none":
        receipt["intervention_status"] = "watching"
    log_note = (request.form.get("intervention_log_entry") or "").strip()
    actor = session.get("user_name") or session.get("username") or ""
    receipt["intervention_by"] = actor
    receipt["intervention_at"] = datetime.now().isoformat(timespec="seconds")
    if log_note:
        logs = receipt.get("intervention_log") or []
        if isinstance(logs, str):
            logs = _parse_json_list(logs)
        logs.append(
            {
                "date": date.today().isoformat(),
                "by": actor,
                "note": log_note,
            }
        )
        receipt["intervention_log"] = logs
    _save_service_data()
    if esc != old_esc:
        try:
            target = receipt.get("as_staff") or receipt.get("receipt_staff") or ""
            old_label = dict(v32.ESCALATION_OPTIONS).get(old_esc, old_esc)
            new_label = dict(v32.ESCALATION_OPTIONS).get(esc, esc)
            _push_notification(
                user_code=target,
                title=(
                    f"[최고기술경영실] {receipt_no} 등급 변경: "
                    f"{old_label} → {new_label}"
                ),
                body=f"고객·보상협의: {_compensation_label(receipt.get('compensation_status'))}",
                link=f"/report?receipt_no={receipt_no}",
                receipt_no=receipt_no,
            )
            notify_escalation_changed(
                receipt,
                _get_staff_email,
                _mail_base_url(),
                old_label=old_label,
                new_label=new_label,
            )
        except Exception:
            pass

    flash(f"개입·보상 상태를 저장했습니다. ({receipt_no})", "success")
    return_to = (request.form.get("return_to") or "").strip()
    if return_to == "home":
        return redirect(url_for("home", _anchor="auto-detect"))
    return redirect(
        url_for(
            "tech_monitor",
            escalation=request.form.get("return_escalation") or "active",
            team=request.form.get("return_team") or None,
            status=request.form.get("return_status") or None,
            q=request.form.get("return_q") or None,
            tab=request.form.get("return_tab") or "list",
            focus=request.form.get("return_focus") or receipt_no,
            _anchor="case-edit",
        )
    )


@app.route("/receipt/<receipt_no>/send-email", methods=["POST"])
@_login_required
def receipt_send_email(receipt_no):
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("home"))

    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        flash("접수 건을 찾을 수 없습니다.", "error")
        return redirect(url_for("home"))

    if _receipt_notice_sent(receipt):
        flash("이미 안내 메일이 발송된 건입니다.", "error")
        return redirect(url_for("home"))

    ok, reason = send_receipt_notice(receipt, _get_staff_email, _mail_base_url())
    if ok:
        _mark_receipt_notice_sent(receipt)
        flash(f"이메일을 발송했습니다. (접수번호: {receipt_no})", "success")
    else:
        flash(mail_error_message(reason), "error")
    return redirect(url_for("home"))




@app.route("/sales")
@_login_required
def sales():
    person = _get_person_for_session()
    if not _can_access_sales(person):
        flash("영업접수 화면 접근 권한이 없습니다.", "error")
        return redirect(url_for("home"))

    line_rows = _line_receipts()
    visible_rows = _filter_receipts_for_sales_list(line_rows, person)
    queue_team_scope = _queue_summary_team_scope(person)
    enriched = [
        _enrich_receipt(r, visible_rows, queue_team_scope=queue_team_scope)
        for r in visible_rows
    ]
    for item in enriched:
        try:
            item["can_ops"] = _can_edit_billing_ops(item, person)
        except Exception:
            item["can_ops"] = False
        try:
            item["can_reassign_as"] = _can_reassign_as_staff(item, person)
        except Exception:
            item["can_reassign_as"] = False
        item["can_claim_as"] = _can_claim_unassigned_as(item, person)
        item["can_manage"] = _can_edit_receipt_registration(item, person)
        item["can_delete"] = _can_delete_receipt(item, person)
        completed = _is_receipt_completed(item)
        # 진행 중: 기존 부품관리 유지. 완료 후: 교체부품(AS기록)은 admin만.
        item["can_edit_parts"] = (
            _can_edit_history_as_content(item, person)
            if completed
            else bool(item["can_ops"])
        )
        try:
            is_head = bool(_is_sales_head(person))
        except Exception:
            is_head = bool(
                person
                and (
                    person.get("team_code") == "sales_head"
                    or person.get("id") == "sales_head"
                )
            )
        item["delete_locked"] = bool(
            completed
            and not item["can_delete"]
            and (
                is_head
                or _is_receipt_registrant(item, person)
            )
        )
    enriched.sort(key=v32.recent_board_sort_key)
    edit_no = request.args.get("edit", "").strip()
    edit_receipt = _get_receipt_by_no(edit_no) if edit_no else None
    if edit_receipt and not _can_view_sales_receipt(edit_receipt, person):
        flash("조회 권한이 없는 접수입니다.", "error")
        edit_receipt = None
    elif edit_receipt and not _can_edit_receipt_registration(edit_receipt):
        try:
            if session.get("user_role") != "admin":
                flash("완료된 건이거나 수정 권한이 없습니다. 청구 비용은 목록에서 조정할 수 있습니다.", "error")
                edit_receipt = None
        except RuntimeError:
            flash("수정·삭제는 등록자 또는 영업부 부서장만 가능합니다.", "error")
            edit_receipt = None
    if edit_receipt:
        _recover_orphan_upload_photos(edit_receipt, persist=True)

    as_options = [
        {"id": p["id"], "label": p["display_name"], "team_code": p["team_code"]}
        for p in _as_assignable_people(person)
    ]
    org_tree = _as_assign_org_tree(person)

    scope_hint = ""
    if _can_view_all_sales_receipts(person):
        scope_hint = "전체 접수 조회"
    elif _is_office_person(person):
        scope_hint = f"{_office_label(person) or '소속 사무소'} 접수만 표시"
    elif person:
        scope_hint = f"{person.get('team') or person.get('team_code') or '소속 팀'} 접수만 표시"

    form_teams, team_select_enabled = _sales_form_teams(person)
    is_export_chooser = _is_export_response_chooser(person)
    can_reassign_form = (
        False
        if (is_export_chooser and not edit_receipt)
        else _can_reassign_as_staff(edit_receipt, person)
    )

    return render_template(
        "sales.html",
        active_page="sales",
        today=date.today().isoformat(),
        teams=form_teams,
        team_select_enabled=team_select_enabled,
        can_choose_as_response_team=is_export_chooser,
        receipt_staff=_get_staff_master("receipt_staff"),
        as_staff=_get_staff_master("as_staff"),
        org_tree=org_tree,
        receipts=enriched,
        status_counts=_get_status_counts(visible_rows),
        next_receipt_no=_preview_receipt_no(date.today().isoformat()),
        edit_receipt=edit_receipt,
        edit_detail=(REPORT_DETAILS.get(edit_receipt["receipt_no"]) or {}) if edit_receipt else {},
        export_as_allowed_codes=sorted(EXPORT_AS_ALLOWED_CODES),
        export_as_default=EXPORT_AS_HANDLER_CODE,
        export_as_handler_name=EXPORT_AS_HANDLER_NAME,
        export_as_self_code=EXPORT_AS_SELF_CODE,
        export_as_self_name=EXPORT_AS_SELF_NAME,
        is_sales_head=_is_sales_head(person),
        can_edit_all_ops=_is_sales_head(person) or session.get("user_role") == "admin",
        as_reassign_options=as_options,
        current_person_team=(person or {}).get("team_code") or "",
        sales_scope_hint=scope_hint,
        can_request_quality_support=_can_request_quality_support(person),
        can_choose_receipt_staff=_can_choose_receipt_staff(person),
        receipt_staff_choice_tree=_receipt_staff_choice_tree(),
        locked_receipt_staff=_locked_receipt_staff(person),
        can_reassign_form=can_reassign_form,
        can_pick_assigned_office=_can_pick_assigned_office(person),
        office_choices=_sales_office_choices(
            (edit_receipt or {}).get("office") if edit_receipt else None
        ),
        queue_summary=(
            None
            if _is_office_person(person)
            else _compute_queue_summary(visible_rows, person)
        ),
    )





@app.route("/api/receipt-no-preview")
@_login_required
def receipt_no_preview():

    receipt_date = request.args.get("date", date.today().isoformat())

    try:

        return {"receipt_no": _preview_receipt_no(receipt_date)}

    except ValueError:

        return {"error": "유효하지 않은 날짜입니다."}, 400





@app.route("/sales/register", methods=["POST"])
@_login_required
@_service_write_locked
def sales_register():
    global _next_id

    receipt_date = request.form.get("receipt_date", date.today().isoformat())
    person = _get_person_for_session()
    team_id = _normalize_sales_team_id(person, request.form.get("team", "domestic"))
    receipt_staff_id = request.form.get("receipt_staff", "")
    receipt_staff_typed = (request.form.get("receipt_staff_name") or "").strip()
    as_staff_id = request.form.get("as_staff", "")
    if _is_export_response_chooser(person):
        as_staff_id = _apply_export_response_as_staff(person, team_id, as_staff_id)
    elif not _can_reassign_as_staff({"team_id": team_id}, person):
        as_staff_id = ""
    else:
        as_staff_id = _sanitize_as_staff_choice(as_staff_id, person)
    charge_type = request.form.get("charge_type", "free")
    spec = (request.form.get("spec") or "").strip()
    lot = (request.form.get("lot") or "").strip()
    if _is_rods() and not spec:
        flash("용접봉 접수는 규격을 필수 입력해야 합니다.", "error")
        return redirect(url_for("sales"))

    office_name = ""
    receipt_staff_id, receipt_staff_name = _resolve_receipt_staff(
        person, receipt_staff_id, receipt_staff_typed
    )
    if _is_office_person(person):
        office_name = _office_label(person)
        if not receipt_staff_name:
            flash("응대한 접수 담당자 이름을 입력해 주세요.", "error")
            return redirect(url_for("sales"))
    elif not receipt_staff_id:
        flash("접수 담당자를 지정해 주세요.", "error")
        return redirect(url_for("sales"))
    elif _can_pick_assigned_office(person):
        office_name = _normalize_assigned_office(request.form.get("office", ""))

    receipt_no = _generate_receipt_no(receipt_date)
    photo_files = request.files.getlist("photos")
    photo_urls = _save_uploaded_photos(
        receipt_no,
        photo_files,
        tags=["customer"] * len(photo_files),
    )

    new_receipt = {
        "id": _next_id,
        "receipt_no": receipt_no,
        "receipt_date": receipt_date,
        "team_id": team_id,
        "team": _team_name(team_id),
        "customer_name": request.form.get("customer_name", ""),
        "customer_contact": request.form.get("customer_contact", ""),
        "equipment_name": request.form.get("equipment_name", ""),
        "symptoms": request.form.get("symptoms", ""),
        "as_staff": as_staff_id,
        "as_staff_name": _staff_name("as_staff", as_staff_id),
        "receipt_staff": receipt_staff_id,
        "receipt_staff_name": receipt_staff_name,
        "office": office_name,
        "status": "received",
        "charge_type": charge_type,
        "billing_cost": None,
        "completed_date": None,
        "created_at": datetime.now().isoformat(),
        "photos": photo_urls,
        "phone": _format_mobile_phone(request.form.get("phone", "")),
        "ship_date": request.form.get("ship_date", ""),
        "support_quality": bool(request.form.get("support_quality"))
        and _can_request_quality_support(person),
        "support_tech": bool(request.form.get("support_tech"))
        and _can_share_tech_support(person)
        and not _person_is_quality_support(person),
        "created_by_user_id": session.get("user_id"),
        "created_by_username": session.get("username", ""),
        "severity": "medium",
        "escalation_level": "normal",
        "intervention_status": "none",
        "compensation_status": "none",
        "intervention_note": "",
        "compensation_note": "",
        "compensation_amount": "",
        "intervention_log": [],
        "intervention_by": "",
        "intervention_at": None,
        "billing_margin_rate": None,
        "billing_manual": 0,
        "billing_set_by": "",
        "product_line": plex.receipt_line_code(),
        "spec": spec,
        "lot": lot,
        "notice_sent": 0,
    }
    serial_no = (request.form.get("serial_no") or "").strip()
    if serial_no and not _is_rods():
        REPORT_DETAILS.setdefault(receipt_no, {})
        REPORT_DETAILS[receipt_no]["serial_no"] = serial_no
    if not _is_export_response_chooser(person):
        _assign_export_as_handler(
            new_receipt,
            default_if_empty=not _can_reassign_as_staff(new_receipt, person),
        )
        # 부서장이 수출 건에 타 인원 지정 시 유지
        if (
            new_receipt.get("team_id") == "export"
            and as_staff_id
            and (_is_sales_head(person) or session.get("user_role") == "admin")
        ):
            new_receipt["as_staff"] = as_staff_id
            new_receipt["as_staff_name"] = _staff_name("as_staff", as_staff_id)
    _next_id += 1
    RECEIPTS.append(new_receipt)
    _save_service_data()
    mailed = notify_receipt_registered(
        new_receipt,
        _get_staff_email,
        _mail_base_url(),
        recipient_codes=_register_notify_codes(new_receipt),
    )
    if mailed:
        _mark_receipt_notice_sent(new_receipt)
    try:
        _notify_export_register(new_receipt)
    except Exception:
        pass
    try:
        _notify_support_request(new_receipt, requester=(person or {}).get("name") or "")
    except Exception:
        pass

    flash(
        (
            f"접수가 완료되었고 품질팀에 공유했습니다. (접수번호: {receipt_no})"
            if new_receipt.get("support_quality")
            else f"접수가 완료되었습니다. (접수번호: {receipt_no})"
        ),
        "success",
    )
    session["last_worked_receipt_no"] = receipt_no
    return redirect(url_for("sales", _anchor="receiptList"))


@app.route("/sales/<receipt_no>/update", methods=["POST"])
@_login_required
@_service_write_locked
def sales_update(receipt_no):
    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        flash("수정할 접수 건을 찾을 수 없습니다.", "error")
        return redirect(url_for("sales"))
    person = _get_person_for_session()
    if not _can_view_sales_receipt(receipt, person):
        flash("조회 권한이 없는 접수입니다.", "error")
        return redirect(url_for("sales", _anchor="receiptList"))
    if not _can_edit_receipt_registration(receipt, person):
        flash(
            "완료된 건이거나 접수 등록정보 수정 권한이 없습니다. "
            "청구 비용은 목록·이력관리에서 조정할 수 있습니다.",
            "error",
        )
        return redirect(url_for("sales", _anchor="receiptList"))
    team_id = _normalize_sales_team_id(
        person, request.form.get("team", receipt.get("team_id", "")), receipt
    )
    spec = (request.form.get("spec") or "").strip()
    lot = (request.form.get("lot") or "").strip()
    if _is_rods() and not spec:
        flash("용접봉 접수는 규격을 필수 입력해야 합니다.", "error")
        return redirect(url_for("sales", edit=receipt_no))
    receipt_staff_id = request.form.get("receipt_staff", "")
    receipt_staff_typed = (request.form.get("receipt_staff_name") or "").strip()
    receipt_staff_id, receipt_staff_name = _resolve_receipt_staff(
        person, receipt_staff_id, receipt_staff_typed
    )
    if _is_office_person(person):
        if not receipt_staff_name:
            flash("응대한 접수 담당자 이름을 입력해 주세요.", "error")
            return redirect(url_for("sales", edit=receipt_no))
        office_name = receipt.get("office") or _office_label(person)
    elif not receipt_staff_id:
        flash("접수 담당자를 지정해 주세요.", "error")
        return redirect(url_for("sales", edit=receipt_no))
    elif _can_pick_assigned_office(person):
        office_name = _normalize_assigned_office(request.form.get("office", ""))
    else:
        office_name = receipt.get("office") or ""
    as_staff_id = request.form.get("as_staff", "")
    old_receipt_staff = receipt.get("receipt_staff", "")
    old_as_staff = receipt.get("as_staff", "")
    old_as_staff_name = receipt.get("as_staff_name", "")
    old_team_id = receipt.get("team_id") or ""
    can_reassign = _can_reassign_as_staff(receipt, person)
    team_changed_by_export = _is_export_response_chooser(person) and team_id != old_team_id
    if not can_reassign and not team_changed_by_export:
        as_staff_id = old_as_staff
    if team_changed_by_export:
        as_staff_id = _apply_export_response_as_staff(person, team_id, as_staff_id)
    elif can_reassign:
        as_staff_id = _sanitize_as_staff_choice(as_staff_id, person)
    was_support_quality = bool(receipt.get("support_quality"))
    was_support_tech = bool(receipt.get("support_tech"))
    photo_files = request.files.getlist("photos")
    new_photos = _save_uploaded_photos(
        receipt_no,
        photo_files,
        tags=["customer"] * len(photo_files),
    )
    removed_photos = set(request.form.getlist("removed_photos"))
    remaining_photos = [
        _normalize_photo(photo, "customer")
        for photo in (receipt.get("photos") or [])
        if (_normalize_photo(photo, "customer") or {}).get("url") not in removed_photos
    ]
    remaining_photos = [p for p in remaining_photos if p]

    receipt.update(
        {
            "receipt_date": request.form.get("receipt_date", receipt["receipt_date"]),
            "team_id": team_id,
            "team": _team_name(team_id),
            "customer_name": request.form.get("customer_name", "").strip(),
            "customer_contact": request.form.get("customer_contact", "").strip(),
            "phone": _format_mobile_phone(request.form.get("phone", "")),
            "equipment_name": request.form.get("equipment_name", "").strip(),
            "spec": spec,
            "lot": lot,
            "ship_date": request.form.get("ship_date", ""),
            "symptoms": request.form.get("symptoms", "").strip(),
            "receipt_staff": receipt_staff_id,
            "receipt_staff_name": receipt_staff_name,
            "office": office_name,
            "as_staff": as_staff_id,
            "as_staff_name": _staff_name("as_staff", as_staff_id) if as_staff_id else "",
            "charge_type": request.form.get("charge_type", "free"),
            "support_quality": (
                bool(request.form.get("support_quality"))
                if _can_request_quality_support(person)
                else bool(receipt.get("support_quality"))
            ),
            "support_tech": (
                bool(request.form.get("support_tech"))
                if _can_share_tech_support(person) and not _person_is_quality_support(person)
                else bool(receipt.get("support_tech"))
            ),
            "photos": remaining_photos + new_photos,
        }
    )
    if receipt.get("charge_type") == "free":
        _clear_billing_on_free(receipt)
    if not can_reassign and not team_changed_by_export:
        receipt["as_staff"] = old_as_staff
        receipt["as_staff_name"] = old_as_staff_name
    elif not (_is_sales_head(person) or session.get("user_role") == "admin"):
        if not _is_export_response_chooser(person):
            # 부서장·admin은 수출 건도 자유 지정. 그 외는 장호빈·김윤식·미지정만
            _assign_export_as_handler(receipt)

    detail = REPORT_DETAILS.get(receipt_no)
    if detail is not None:
        detail["phone"] = receipt["phone"]
        detail["ship_date"] = receipt["ship_date"]
        detail["affiliation"] = receipt["team"]
        detail["receipt_staff_name"] = receipt["receipt_staff_name"] or "—"
        detail["as_technician"] = receipt["as_staff_name"] or "미지정"
    if not _is_rods():
        serial_no = (request.form.get("serial_no") or "").strip()
        detail_row = REPORT_DETAILS.setdefault(receipt_no, {})
        detail_row["serial_no"] = serial_no

    _save_service_data()
    base_url = _mail_base_url()
    if receipt_staff_id and receipt_staff_id != old_receipt_staff:
        notify_staff_assigned(receipt, receipt_staff_id, "접수 담당자", _get_staff_email, base_url)
    final_as = receipt.get("as_staff") or ""
    if final_as and final_as != old_as_staff:
        _notify_as_staff_reassigned(receipt, final_as)
    try:
        notify_receipt = {
            "receipt_no": receipt_no,
            "customer_name": receipt.get("customer_name"),
            "symptoms": receipt.get("symptoms"),
            "support_quality": bool(receipt.get("support_quality")) and not was_support_quality,
            "support_tech": bool(receipt.get("support_tech")) and not was_support_tech,
        }
        if notify_receipt["support_quality"] or notify_receipt["support_tech"]:
            _notify_support_request(
                notify_receipt, requester=(person or {}).get("name") or ""
            )
    except Exception:
        pass
    flash(f"접수 내용을 수정했습니다. (접수번호: {receipt_no})", "success")
    return redirect(url_for("sales", _anchor="receiptList"))


@app.route("/sales/<receipt_no>/delete", methods=["POST"])
@_login_required
@_service_write_locked
def sales_delete(receipt_no):
    if not _valid_admin_csrf():
        flash("잘못된 삭제 요청입니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("sales", _anchor="receiptList"))

    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        flash("삭제할 접수 건을 찾을 수 없습니다.", "error")
        return redirect(url_for("sales", _anchor="receiptList"))
    person = _get_person_for_session()
    if not _can_view_sales_receipt(receipt, person):
        flash("조회 권한이 없는 접수입니다.", "error")
        return redirect(url_for("sales", _anchor="receiptList"))
    if not _can_delete_receipt(receipt, person):
        flash(
            "완료된 건은 admin만 삭제할 수 있습니다. 진행 중 건은 등록자·부서장만 삭제 가능합니다.",
            "error",
        )
        return redirect(url_for("sales", _anchor="receiptList"))

    RECEIPTS.remove(receipt)
    REPORT_DETAILS.pop(receipt_no, None)
    _save_service_data()
    flash(f"접수 건 전체를 삭제했습니다. (접수번호: {receipt_no})", "success")
    return redirect(url_for("sales", _anchor="receiptList"))





@app.route("/history")
@_login_required
def history():
    person = _get_person_for_session()
    if not _can_access_history(person):
        flash("이력관리 화면 접근 권한이 없습니다.", "error")
        return redirect(url_for("home"))
    history_teams = _teams_for_history()
    if not history_teams:
        history_teams = _get_teams()
    team_codes = {team["id"] for team in history_teams}
    # 내수영업팀 로그인 시 기본 탭을 내수로
    if person and person.get("team_code") == "domestic":
        default_tab = "domestic"
    elif person and person.get("team_code") == "export":
        default_tab = "export"
    else:
        default_tab = history_teams[0]["id"] if history_teams else ""
    tab = request.args.get("tab", default_tab)
    if tab not in team_codes:
        tab = default_tab

    for team in history_teams:
        team["waiting_count"] = len(_waiting_receipts(team["id"], person))

    waiting = _waiting_receipts(tab, person)
    options = []
    for r in waiting:
        label = f"{r['receipt_no']} · {r['customer_name']} · {r['equipment_name']}"
        if not (r.get("as_staff") or "").strip():
            label += " · 미지정"
        if tab == "domestic" and r.get("team_id") == "export":
            label = f"[수출] {label}"
        options.append({"receipt_no": r["receipt_no"], "label": label})
    selected_receipt_no = request.args.get("receipt_no", "").strip()
    explicit_receipt_no = selected_receipt_no
    option_nos = {item["receipt_no"] for item in options}
    if selected_receipt_no and selected_receipt_no not in option_nos:
        selected = _get_receipt_by_no(selected_receipt_no)
        if selected and _is_jang_hb(person) and not _jang_hb_assigned(selected, person):
            flash(
                "장호빈은 최영호 팀장이 지정한 건만 진행할 수 있습니다. 미지정 건은 해당 팀원이 클릭해 담당하거나, 팀장이 지정해야 합니다.",
                "error",
            )
            selected_receipt_no = ""
        elif selected and _can_process_export_as(selected, person):
            # 홈 등에서 완료/다른 상태 건을 열어온 경우에도 선택 가능하도록 추가
            prefix = "[수출] " if selected.get("team_id") == "export" else ""
            options.insert(
                0,
                {
                    "receipt_no": selected["receipt_no"],
                    "label": (
                        f"{prefix}{selected['receipt_no']} · "
                        f"{selected['customer_name']} · {selected['equipment_name']}"
                    ),
                },
            )
            # 수출 건은 내수 처리자(장호빈)에게 내수 탭으로 유지
            if selected.get("team_id") == "export" and _is_export_as_handler(person):
                tab = "domestic"
            elif selected.get("team_id") in team_codes:
                tab = selected["team_id"]
            for team in history_teams:
                team["waiting_count"] = len(_waiting_receipts(team["id"], person))
            waiting = _waiting_receipts(tab, person)
        else:
            if selected and selected.get("team_id") == "export":
                flash(
                    "수출영업팀 접수는 내수영업팀 장호빈 담당자만 이력관리에서 처리할 수 있습니다.",
                    "error",
                )
            selected_receipt_no = ""
    elif selected_receipt_no not in option_nos:
        selected_receipt_no = ""
    # 홈·목록에서 미지정 접수번호를 클릭해 들어온 경우: 팀원 본인 담당으로 지정
    if explicit_receipt_no and selected_receipt_no == explicit_receipt_no:
        picked = _get_receipt_by_no(explicit_receipt_no)
        if picked and _is_jang_hb(person) and not _jang_hb_assigned(picked, person):
            flash(
                "장호빈은 최영호 팀장이 지정한 건만 진행할 수 있습니다. 미지정 건은 해당 팀원이 클릭해 담당하거나, 팀장이 지정해야 합니다.",
                "error",
            )
            selected_receipt_no = ""
        elif picked and _can_claim_unassigned_as(picked, person):
            actor = (person.get("display_name") or person.get("name") or "").strip()
            if _claim_unassigned_as(picked, person):
                flash(
                    f"미지정 건을 본인({actor}) 담당으로 지정했습니다. AS 이력을 진행할 수 있습니다.",
                    "success",
                )
    # 미선택 시: 최근 작업 건(대기목록에 있으면) → 없으면 최신(목록 첫) 건 자동 선택
    if not selected_receipt_no and options:
        last_worked = (session.get("last_worked_receipt_no") or "").strip()
        if last_worked and last_worked in option_nos:
            selected_receipt_no = last_worked
        else:
            selected_receipt_no = options[0]["receipt_no"]
    return render_template(
        "history.html",
        active_page="history",
        today=date.today().isoformat(),
        teams=_get_teams(),
        affiliation_options=_affiliation_options(),
        default_affiliation=_default_affiliation(),
        history_teams=history_teams,
        history_tab=tab,
        waiting_count=len(waiting),
        waiting_options=options,
        selected_receipt_no=selected_receipt_no,
        current_user_name=session.get("user_name") or "",
        current_person_id=(person or {}).get("id") or "",
        is_support_viewer=bool(person and person.get("team_code") == "field_support"),
        is_quality_viewer=_person_is_quality_support(person),
        can_request_quality_support=_can_request_quality_support(person),
        can_share_tech_support=_can_share_tech_support(person),
        is_sales_head=bool(_is_sales_head(person)) if person else False,
        is_sales_team_leader=bool(_is_sales_team_leader(person)) if person else False,
    )


@app.route("/api/history/list")
@_login_required
def api_history_list():
    person = _get_person_for_session()
    teams = _get_teams()
    default_tab = teams[0]["id"] if teams else ""
    tab = request.args.get("tab", default_tab)
    waiting = _waiting_receipts(tab, person)
    items = []
    for r in waiting:
        label = f"{r['receipt_no']} · {r['customer_name']} · {r['equipment_name']}"
        if not (r.get("as_staff") or "").strip():
            label += " · 미지정"
        if tab == "domestic" and r.get("team_id") == "export":
            label = f"[수출] {label}"
        items.append({"receipt_no": r["receipt_no"], "label": label})
    return {
        "count": len(waiting),
        "items": items,
    }


@app.route("/api/history/<receipt_no>")
@_login_required
def api_history_detail(receipt_no):
    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        return {"error": "접수번호를 찾을 수 없습니다."}, 404
    person = _get_person_for_session()
    if not _can_process_export_as(receipt, person):
        return {
            "error": "수출영업팀 접수는 내수영업팀 장호빈 담당자만 처리할 수 있습니다."
        }, 403
    if _is_jang_hb(person) and not _jang_hb_assigned(receipt, person):
        return {
            "error": "장호빈은 최영호 팀장이 지정한 건만 진행할 수 있습니다. 미지정 건은 해당 팀원이 클릭해 담당하거나, 팀장이 지정해야 합니다."
        }, 403
    # 지원부서 불러오기 시 대기중 → 확인함 자동 전환
    try:
        if _auto_confirm_support_on_view(receipt, person):
            _save_service_data()
    except NameError:
        pass
    except Exception:
        pass
    return _history_payload(receipt)


@app.route("/api/history/claim", methods=["POST"])
@_login_required
@_service_write_locked
def api_history_claim():
    """미지정 건을 클릭한 팀원 본인 AS담당으로 지정."""
    receipt_no = (request.form.get("receipt_no") or "").strip()
    receipt = _get_receipt_by_no(receipt_no)
    person = _get_person_for_session()
    if not receipt:
        return jsonify({"error": "접수번호를 찾을 수 없습니다."}), 404
    if _is_jang_hb(person):
        return jsonify({
            "error": "장호빈은 최영호 팀장이 지정한 건만 진행할 수 있습니다."
        }), 403
    current = (receipt.get("as_staff") or "").strip()
    if current:
        if current == (person.get("id") or person.get("code") or "").strip():
            return jsonify({
                "ok": True,
                "as_staff": receipt.get("as_staff"),
                "as_staff_name": receipt.get("as_staff_name"),
            })
        return jsonify({
            "error": f"이미 {receipt.get('as_staff_name') or '다른 담당자'}에게 지정된 건입니다."
        }), 409
    if not _can_claim_unassigned_as(receipt, person):
        return jsonify({"error": "미지정 건을 지정할 권한이 없습니다."}), 403
    if not _claim_unassigned_as(receipt, person):
        return jsonify({
            "error": "지정에 실패했습니다. 이미 다른 담당자가 지정했을 수 있습니다."
        }), 409
    return jsonify({
        "ok": True,
        "as_staff": receipt.get("as_staff"),
        "as_staff_name": receipt.get("as_staff_name"),
        "message": f"미지정 건을 본인({receipt.get('as_staff_name')}) 담당으로 지정했습니다.",
    })


def _history_support_field_draft(receipt, person):
    """지원부서 임시저장: 점검·수리·상세·사진만 반영. 완료/청구/부품/담당자 변경 없음."""
    import json

    receipt_no = receipt.get("receipt_no") or ""
    tab = request.form.get("tab") or "domestic"
    redirect_to = url_for(
        "history", tab=tab, receipt_no=receipt_no, _anchor="historyForm"
    )
    existing_detail = REPORT_DETAILS.get(receipt_no) or {}
    if _history_lock_conflict(existing_detail, request.form.get("history_updated_at", "")):
        flash(_history_conflict_message(existing_detail), "error")
        return redirect(redirect_to)

    existing_raw = request.form.get("existing_photos_json", "[]")
    try:
        existing_photos = json.loads(existing_raw)
    except json.JSONDecodeError:
        existing_photos = []
    existing_photos = [
        p for p in (_normalize_photo(x, "before") for x in existing_photos) if p
    ]
    new_tags = request.form.getlist("photo_tags")
    new_photos = _save_uploaded_photos(
        receipt_no,
        request.files.getlist("field_photos"),
        subdir="field",
        tags=new_tags,
    )
    all_field_photos = existing_photos + new_photos

    detail = REPORT_DETAILS.setdefault(receipt_no, {})
    detail.update(
        {
            "inspection_result": request.form.get("inspection_result", ""),
            "repair_content": request.form.get("repair_content", ""),
            "detailed_work": request.form.get("detailed_work", ""),
            "photos": all_field_photos,
            "inspection_date": request.form.get("inspection_date")
            or detail.get("inspection_date")
            or date.today().isoformat(),
        }
    )
    # 미완료 건은 임시저장 시 진행중으로만 유지(완료 금지)
    status = (receipt.get("status") or "received").strip()
    if status in ("received", "draft", ""):
        receipt["status"] = "in_progress"
        detail["as_status"] = "in_progress"
    elif status != "completed":
        detail["as_status"] = status
    _stamp_history_modified(detail, person)
    _save_service_data()
    flash("지원 현장 내용(점검·수리·상세·사진)을 임시 저장했습니다. 완료 처리는 AS담당자가 합니다.", "success")
    return redirect(redirect_to)


def _quality_share_tech(receipt, person):
    """품질팀: 이미 공유된 건만 기술부/기술연구소로 이관."""
    tab = request.form.get("tab") or request.args.get("tab") or ""
    receipt_no = receipt.get("receipt_no") or ""
    redirect_to = url_for("history", tab=tab, receipt_no=receipt_no) if tab else url_for(
        "history", receipt_no=receipt_no
    )
    if not receipt.get("support_quality"):
        flash("품질팀에 공유된 건만 기술 지원 부서로 이관할 수 있습니다.", "error")
        return redirect(redirect_to)
    if receipt.get("status") == "completed":
        flash("완료 처리된 건은 수정할 수 없습니다.", "error")
        return redirect(redirect_to)
    existing_detail = REPORT_DETAILS.get(receipt_no) or {}
    if _history_lock_conflict(existing_detail, request.form.get("history_updated_at", "")):
        flash(_history_conflict_message(existing_detail), "error")
        return redirect(redirect_to)
    want_tech = bool(request.form.get("support_tech"))
    was_tech = bool(receipt.get("support_tech"))
    receipt["support_tech"] = want_tech
    tech_label = _support_tech_dept_label(receipt)
    detail = REPORT_DETAILS.get(receipt_no, {})
    notes = []
    if receipt.get("support_quality"):
        notes.append(f"품질팀 지원 요청 ({date.today().isoformat()})")
    if want_tech:
        notes.append(f"{tech_label} 지원 요청 ({date.today().isoformat()})")
    if detail.get("support_tech_mgmt"):
        notes.append(f"최고기술경영실 보고 ({date.today().isoformat()})")
    detail["support_tech"] = want_tech
    detail["support_quality"] = True
    detail["support_history"] = " / ".join(notes) if notes else "지원 요청 이력 없음"
    _stamp_history_modified(detail, person)
    REPORT_DETAILS[receipt_no] = detail
    _save_service_data()
    if want_tech and not was_tech:
        _notify_support_request(
            {
                **receipt,
                "support_quality": False,
                "support_tech": True,
            },
            requester=(person or {}).get("name") or "품질팀",
        )
        flash(f"{tech_label}에 공유했습니다. (접수번호: {receipt_no})", "success")
    elif not want_tech and was_tech:
        flash(f"{tech_label} 공유를 해제했습니다. (접수번호: {receipt_no})", "success")
    else:
        flash("지원 공유 상태를 저장했습니다.", "success")
    return redirect(redirect_to)


@app.route("/api/support-share", methods=["POST"])
@_login_required
@_service_write_locked
def api_support_share():
    data = request.get_json(silent=True) or request.form
    receipt_no = (data.get("receipt_no") or "").strip()
    kind = (data.get("kind") or "").strip()
    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        return jsonify({"ok": False, "error": "접수번호를 찾을 수 없습니다."}), 404
    person = _get_person_for_session()
    can_sales_share = False
    try:
        can_sales_share = bool(
            _can_view_sales_receipt(receipt, person)
            and (
                _can_manage_receipt(receipt)
                or _is_sales_head(person)
                or session.get("user_role") == "admin"
            )
        )
    except Exception:
        can_sales_share = False
    can_history_share = False
    try:
        can_history_share = bool(_can_access_history(person))
    except Exception:
        can_history_share = False
    if not can_sales_share and not can_history_share:
        return jsonify({"ok": False, "error": "공유 권한이 없습니다."}), 403
    submitted = None
    if "history_updated_at" in data:
        submitted = data.get("history_updated_at") or ""
    result = _share_support_kind(receipt, kind, person, submitted)
    if result.get("conflict"):
        return jsonify(result), 409
    return jsonify(result), (200 if result.get("ok") else 400)


@app.route("/api/session-status")
def api_session_status():
    """다른 탭에서 계정이 바뀌었는지 감지하기 위한 현재 세션 식별자."""
    return jsonify(
        {
            "user_id": session.get("user_id"),
            "username": session.get("username") or "",
            "user_name": session.get("user_name") or "",
        }
    )


@app.route("/history/save", methods=["POST"])
@_login_required
@_service_write_locked
def history_save():
    import json

    receipt_no = request.form.get("receipt_no", "").strip()
    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        flash("접수번호를 찾을 수 없습니다.", "error")
        return redirect(url_for("history"))
    person = _get_person_for_session()
    if not _can_access_history(person):
        flash("이력관리 화면 접근 권한이 없습니다.", "error")
        return redirect(url_for("home"))
    if person and person.get("team_code") == "field_support":
        if _person_is_quality_support(person) and request.form.get("action") == "share_tech":
            return _quality_share_tech(receipt, person)
        action = (request.form.get("action") or "").strip()
        if action == "draft" and _can_support_field_draft(receipt, person):
            return _history_support_field_draft(receipt, person)
        flash( 
            "지원부서(품질팀/기술부)는 공유된 건의 검토의견 회신과 "
            "점검·수리·사진 임시저장만 가능합니다. 완료 처리는 AS담당자가 합니다.",
            "error",
        )
        return redirect(url_for("history"))
    if not _can_process_export_as(receipt, person):
        flash(
            "수출영업팀 접수는 장호빈 또는 원 접수자(김윤식)만 이력관리에서 처리할 수 있습니다.",
            "error",
        )
        return redirect(url_for("history", tab="domestic"))
    if _is_jang_hb(person) and not _jang_hb_assigned(receipt, person):
        flash(
            "장호빈은 최영호 팀장이 지정한 건만 진행할 수 있습니다.",
            "error",
        )
        return redirect(url_for("history", tab="domestic"))
    if _can_claim_unassigned_as(receipt, person):
        _claim_unassigned_as(receipt, person)

    # 완료 건 · 진행 중 건 권한 분기
    is_completed = _is_receipt_completed(receipt)
    is_admin = session.get("user_role") == "admin"
    action = request.form.get("action", "save")

    if is_completed:
        if is_admin:
            pass  # admin: AS내용 포함 전체 수정 허용
        elif action == "billing" and _can_edit_billing_ops(receipt, person):
            existing_detail = REPORT_DETAILS.get(receipt_no) or {}
            if _history_lock_conflict(existing_detail, request.form.get("history_updated_at", "")):
                flash(_history_conflict_message(existing_detail), "error")
                return redirect(
                    url_for(
                        "history",
                        tab=request.form.get("tab", "domestic"),
                        receipt_no=receipt_no,
                        _anchor="historyForm",
                    )
                )
            parts = existing_detail.get("parts") or []
            _apply_receipt_billing_from_request(receipt, request, parts)
            _stamp_history_modified(existing_detail, person)
            _save_service_data()
            flash("청구 비용을 저장했습니다.", "success")
            return redirect(
                url_for(
                    "history",
                    tab=request.form.get("tab", "domestic"),
                    receipt_no=receipt_no,
                    _anchor="historyForm",
                )
            )
        else:
            flash(
                "완료된 건의 AS 처리 내용은 admin만 수정할 수 있습니다. "
                "청구 비용은 등록자·부서장만 변경 가능합니다.",
                "error",
            )
            return redirect(
                url_for(
                    "history",
                    tab=request.form.get("tab", "domestic"),
                    receipt_no=receipt_no,
                    _anchor="historyForm",
                )
            )
    elif not _can_edit_history_as_content(receipt, person):
        flash("AS 이력은 담당 AS담당자 또는 영업부 부서장만 수정할 수 있습니다.", "error")
        return redirect(
            url_for(
                "history",
                tab=request.form.get("tab", "domestic"),
                receipt_no=receipt_no,
                _anchor="historyForm",
            )
        )

    existing_detail = REPORT_DETAILS.get(receipt_no) or {}
    if _history_lock_conflict(existing_detail, request.form.get("history_updated_at", "")):
        flash(_history_conflict_message(existing_detail), "error")
        return redirect(
            url_for(
                "history",
                tab=request.form.get("tab", "domestic"),
                receipt_no=receipt_no,
                _anchor="historyForm",
            )
        )

    action = request.form.get("action", "save")
    allowed_status = {
        "received",
        "draft",
        "in_progress",
        "parts_waiting",
        "on_hold",
        "completed",
    }
    status = request.form.get("status", "in_progress").strip()
    if status not in allowed_status:
        status = "in_progress"
    if _is_rods() and status == "parts_waiting":
        status = "in_progress"

    # 6.2: 진행중/부품대기/보류 → 임시 저장만 / 완료 → 완료 처리만
    if action == "draft":
        if status == "completed":
            status = "in_progress"
        elif status not in ("in_progress", "parts_waiting", "on_hold", "received", "draft"):
            status = "in_progress"
    elif action == "complete":
        status = "completed"
    else:
        # 일반 저장: 완료가 아니면 처리중 상태 유지, 완료는 complete 액션만 허용
        if status == "completed":
            flash(
                "완료 처리는 '완료 처리 & 저장' 버튼으로만 가능합니다. "
                "진행중·부품대기·보류는 임시 저장을 사용해 주세요.",
                "error",
            )
            return redirect(
                url_for(
                    "history",
                    tab=request.form.get("tab", "domestic"),
                    receipt_no=receipt_no,
                    _anchor="historyForm",
                )
            )

    parts_raw = request.form.get("parts_json", "[]")
    try:
        parts = json.loads(parts_raw)
    except json.JSONDecodeError:
        parts = []

    cleaned_parts = []
    for p in parts:
        name = (p.get("name") or "").strip()
        if not name:
            continue
        try:
            qty = float(p.get("qty") or 0)
        except (TypeError, ValueError):
            qty = 0
        if float(qty).is_integer():
            qty = int(qty)
        try:
            unit_price = int(str(p.get("unit_price") or "0").replace(",", ""))
        except (TypeError, ValueError):
            unit_price = 0
        cleaned_parts.append({
            "name": name,
            "spec": (p.get("spec") or "").strip(),
            "qty": qty,
            "unit_price": unit_price,
            "subtotal": int(round(float(qty) * unit_price)),
        })

    # 기존 현장 사진 유지 + 신규 업로드
    existing_raw = request.form.get("existing_photos_json", "[]")
    try:
        existing_photos = json.loads(existing_raw)
    except json.JSONDecodeError:
        existing_photos = []
    existing_photos = [p for p in (_normalize_photo(x, "before") for x in existing_photos) if p]

    new_tags = request.form.getlist("photo_tags")
    new_photos = _save_uploaded_photos(
        receipt_no,
        request.files.getlist("field_photos"),
        subdir="field",
        tags=new_tags,
    )
    all_field_photos = existing_photos + new_photos

    support_quality, support_tech = _staged_support_flags(person, receipt, request.form)
    support_tech_mgmt = bool(request.form.get("support_tech_mgmt"))
    was_support_quality = bool(receipt.get("support_quality"))
    was_support_tech = bool(receipt.get("support_tech"))
    was_support_tech_mgmt = bool(
        (REPORT_DETAILS.get(receipt_no) or {}).get("support_tech_mgmt")
    )
    support_notes = []
    if support_quality:
        support_notes.append(f"품질팀 지원 요청 ({date.today().isoformat()})")
    if support_tech:
        tech_label = _support_tech_dept_label(receipt)
        support_notes.append(f"{tech_label} 지원 요청 ({date.today().isoformat()})")
    if support_tech_mgmt:
        support_notes.append(f"최고기술경영실 보고 ({date.today().isoformat()})")
        # 에스컬레이션 자동 상향
        if v32.normalize_escalation(receipt) == "normal":
            receipt["escalation_level"] = "monitor"
            receipt["severity"] = v32.ESCALATION_TO_SEVERITY["monitor"]
            receipt["intervention_status"] = "watching"
    support_history = " / ".join(support_notes) if support_notes else "지원 요청 이력 없음"

    loc = request.form.get("process_location", "field")
    loc_label = "현장" if loc == "field" else "공장"
    re_repair = request.form.get("re_repair", "N")
    person = _get_person_for_session()
    # 서비스담당자(as_technician): AS전담·접수+AS 저장 시에만 본인으로 기록.
    # 품질팀·기술부·부서장·관리자 저장·확인은 기존값을 유지(비어 있어도 채우지 않음).
    if person and _person_role_ia_or_a(person):
        as_technician = (
            request.form.get("as_technician")
            or session.get("user_name")
            or (person.get("display_name") or person.get("name") or "")
            or ""
        ).strip()
    else:
        as_technician = (existing_detail.get("as_technician") or "").strip()

    detail = REPORT_DETAILS.get(receipt_no, {})
    detail.update({
        "report_date": date.today().isoformat(),
        "phone": receipt.get("phone") or detail.get("phone", ""),
        "serial_no": request.form.get("serial_no", ""),
        "re_repair": re_repair,
        "receipt_staff_name": detail.get("receipt_staff_name", "—"),
        "as_technician": as_technician,
        "processed_by_user_id": session.get("user_id"),
        "processed_by_username": session.get("username") or "",
        "ship_date": receipt.get("ship_date") or detail.get("ship_date", ""),
        "warranty": detail.get("warranty", ""),
        "inspection_date": request.form.get("inspection_date") or date.today().isoformat(),
        "affiliation": _normalize_affiliation(
            request.form.get("affiliation") or receipt.get("team") or "",
            receipt,
        ),
        "process_location": loc_label,
        "ship_out_date": request.form.get("ship_out_date", ""),
        "inspection_result": request.form.get("inspection_result", ""),
        "repair_content": request.form.get("repair_content", ""),
        "detailed_work": request.form.get("detailed_work", ""),
        "parts": cleaned_parts,
        "support_history": support_history,
        "support_quality": support_quality,
        "support_tech": support_tech,
        "support_tech_mgmt": support_tech_mgmt,
        "photos": all_field_photos,
        "remarks": request.form.get("remarks", ""),
        "as_status": status,
    })
    _stamp_history_modified(detail, person)
    REPORT_DETAILS[receipt_no] = detail

    receipt["status"] = status
    receipt["support_quality"] = support_quality
    receipt["support_tech"] = support_tech
    # AS담당자(receipt): IA·A만 갱신. 비어 있어도 지원부서·부서장·관리자는 손대지 않음.
    if person and _person_role_ia_or_a(person):
        receipt["as_staff"] = person["id"]
        receipt["as_staff_name"] = (
            as_technician or person.get("display_name") or person.get("name") or ""
        )
        try:
            _sync_as_staff_list_primary(receipt, person["id"], receipt["as_staff_name"])
        except NameError:
            pass
    if status == "completed":
        receipt["completed_date"] = request.form.get("ship_out_date") or date.today().isoformat()
    else:
        receipt["completed_date"] = None

    # 유상 건 청구비용: 마진 자동계산 또는 직접입력(협의가)
    if receipt.get("charge_type") == "paid":
        _apply_receipt_billing_from_request(receipt, request, cleaned_parts)

    _save_service_data()
    session["last_worked_receipt_no"] = receipt_no
    status_label = v32.status_label(status)
    try:
        # 새로 켜진 지원요청만 알림
        notify_receipt = {
            **receipt,
            "support_quality": support_quality and not was_support_quality,
            "support_tech": support_tech and not was_support_tech,
        }
        if notify_receipt["support_quality"] or notify_receipt["support_tech"]:
            _notify_support_request(notify_receipt, requester=as_technician)
        if support_tech_mgmt and not was_support_tech_mgmt:
            _push_notification(
                team_code="tech_mgmt",
                title=f"[지원요청] {receipt_no} 최고기술경영실 보고",
                body=f"요청자: {as_technician} / 등급 모니터링 상향",
                link="/tech-monitor",
                receipt_no=receipt_no,
            )
            notify_tech_mgmt_reported(
                receipt,
                _get_staff_email,
                _mail_base_url(),
                requester=as_technician,
            )
    except Exception:
        pass
    if status == "completed":
        office_team = _office_notify_team_code(receipt) if receipt.get("office") else ""
        mail_receipt = {
            **receipt,
            "as_technician": as_technician,
            **({"office_team_code": office_team} if office_team else {}),
        }
        notify_receipt_completed(mail_receipt, _get_staff_email, _mail_base_url())
        try:
            if office_team:
                _push_notification(
                    team_code=office_team,
                    title=f"[AS완료] {receipt_no} 처리 완료",
                    body=f"처리자: {as_technician}",
                    link="/report",
                    receipt_no=receipt_no,
                )
            elif receipt.get("receipt_staff"):
                _push_notification(
                    user_code=receipt["receipt_staff"],
                    title=f"[AS완료] {receipt_no} 처리 완료",
                    body=f"처리자: {as_technician}",
                    link="/report",
                    receipt_no=receipt_no,
                )
        except Exception:
            pass
        flash(f"완료 처리되었습니다. (접수번호: {receipt_no})", "success")
        return redirect(url_for("report", receipt_no=receipt_no))

    flash(f"{status_label} 상태로 저장되었습니다. (접수번호: {receipt_no})", "success")
    return redirect(
        url_for(
            "history",
            tab=request.form.get("tab", "domestic"),
            receipt_no=receipt_no,
            _anchor="historyForm",
        )
    )


@app.route("/report")
@_login_required
def report():
    customer_q = (request.args.get("q") or request.args.get("customer") or "").strip()
    line_rows = list(_line_receipts())
    if customer_q:
        q_cf = customer_q.casefold()
        line_rows = [
            r
            for r in line_rows
            if q_cf in (r.get("customer_name") or "").casefold()
        ]
    receipt_options = [
        {
            "receipt_no": r["receipt_no"],
            "label": f"{r['receipt_no']} · {r['customer_name']}",
            "customer_name": r.get("customer_name") or "",
        }
        for r in sorted(line_rows, key=lambda x: x.get("created_at") or "")
    ]
    selected_no = request.args.get("receipt_no")
    if selected_no and _get_receipt_by_no(selected_no):
        # 검색 결과 밖 번호도 직접 지정 시 허용
        if not any(o["receipt_no"] == selected_no for o in receipt_options):
            receipt = _get_receipt_by_no(selected_no)
            if receipt and plex.receipt_matches_session(receipt):
                receipt_options.insert(
                    0,
                    {
                        "receipt_no": receipt["receipt_no"],
                        "label": f"{receipt['receipt_no']} · {receipt.get('customer_name') or ''}",
                        "customer_name": receipt.get("customer_name") or "",
                    },
                )
        default_no = selected_no
    else:
        last_worked = (session.get("last_worked_receipt_no") or "").strip()
        if (
            last_worked
            and _get_receipt_by_no(last_worked)
            and any(o["receipt_no"] == last_worked for o in receipt_options)
        ):
            default_no = last_worked
        elif receipt_options:
            default_no = receipt_options[0]["receipt_no"]
        else:
            default_no = None
    if default_no:
        session["last_worked_receipt_no"] = default_no
    default_report = None
    if default_no:
        receipt = _get_receipt_by_no(default_no)
        if receipt:
            default_report = _build_report(receipt)

    # 6.9: 작업 화면으로 되돌아가기 (상대경로만 허용)
    return_url = (request.args.get("return_url") or "").strip()
    if not return_url.startswith("/") or return_url.startswith("//"):
        return_url = ""
    return_label = "이전 화면으로 돌아가기"
    if "/history" in return_url:
        return_label = "이력관리로 돌아가기"
    elif "/sales" in return_url:
        return_label = "영업접수로 돌아가기"
    elif return_url in ("/", "/home", "/rods"):
        return_label = "홈으로 돌아가기"
    if not return_url:
        person = _get_person_for_session()
        if _can_access_history(person):
            return_url = url_for("rods_history" if _is_rods() else "history")
            return_label = "이력관리로 돌아가기"
        elif _can_access_sales(person):
            return_url = url_for("rods_sales" if _is_rods() else "sales")
            return_label = "영업접수로 돌아가기"
        else:
            return_url = url_for("rods_home" if _is_rods() else "home")
            return_label = "홈으로 돌아가기"

    return render_template(
        "report.html",
        active_page="report",
        receipt_options=receipt_options,
        default_receipt_no=default_no,
        report=default_report,
        return_url=return_url,
        return_label=return_label,
        customer_q=customer_q,
    )


@app.route("/api/report/<receipt_no>")
@_login_required
def api_report(receipt_no):
    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        return {"error": "접수번호를 찾을 수 없습니다."}, 404
    return _build_report(receipt)


@app.route("/report/<receipt_no>/pdf")
@_login_required
def report_pdf(receipt_no):
    """AS 보고서 PDF 파일 다운로드"""
    from io import BytesIO

    receipt = _get_receipt_by_no(receipt_no)
    if not receipt:
        flash("접수번호를 찾을 수 없습니다.", "error")
        return redirect(url_for("report"))

    report_data = _build_report(receipt)
    try:
        pdf_bytes = build_report_pdf(report_data, BASE_DIR)
    except Exception as exc:
        flash(f"PDF 생성에 실패했습니다: {exc}", "error")
        return redirect(url_for("report", receipt_no=receipt_no))

    filename = f"{receipt_no}_AS_Report.pdf"
    return send_file(
        BytesIO(pdf_bytes),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=filename,
    )





@app.route("/login", methods=["POST"])
def login():
    username = _normalize_login_text(request.form.get("username", ""))
    password = _normalize_login_password(request.form.get("password", ""))
    next_url = request.form.get("next", "").strip()

    with _user_db() as conn:
        user = _find_user_for_login(conn, username)

    if not user or not user["active"] or not check_password_hash(user["password_hash"], password):
        flash("아이디 또는 비밀번호가 올바르지 않습니다.", "error")
        redirect_args = {"login": 1}
        if next_url.startswith("/") and not next_url.startswith("//"):
            redirect_args["next"] = next_url
        return redirect(url_for(_product_home_endpoint(), **redirect_args))

    product_line = session.get("product_line") or "welding"
    session.clear()
    session["product_line"] = product_line
    session["user_id"] = user["id"]
    session["username"] = user["username"]
    session["user_name"] = user["display_name"]
    session["user_role"] = user["role"]
    session["csrf_token"] = uuid.uuid4().hex

    _flash_once(f"{user['display_name']}님, 로그인되었습니다.", "success")
    if user["role"] == "admin":
        return redirect(url_for("admin_users", _anchor="orgMaster"))
    if next_url.startswith("/") and not next_url.startswith("//"):
        return redirect(next_url)
    return redirect(url_for(_product_home_endpoint()))


@app.route("/account/password", methods=["POST"])
@_login_required
def account_password_change():
    if not _valid_session_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return _redirect_with_password_modal(request.form.get("next"))

    current_password = _normalize_login_password(
        request.form.get("current_password", "")
    )
    new_password = _normalize_login_password(request.form.get("new_password", ""))
    confirm_password = _normalize_login_password(
        request.form.get("confirm_password", "")
    )
    next_path = request.form.get("next") or request.referrer

    if not current_password:
        flash("현재 비밀번호를 입력해 주세요.", "error")
        return _redirect_with_password_modal(next_path)
    if len(new_password) < 4:
        flash("새 비밀번호는 4자 이상이어야 합니다.", "error")
        return _redirect_with_password_modal(next_path)
    if new_password != confirm_password:
        flash("새 비밀번호 확인이 일치하지 않습니다.", "error")
        return _redirect_with_password_modal(next_path)
    if current_password == new_password:
        flash("새 비밀번호는 현재 비밀번호와 달라야 합니다.", "error")
        return _redirect_with_password_modal(next_path)

    user_id = session.get("user_id")
    with _user_db() as conn:
        user = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if not user or not user["active"]:
            flash("계정을 찾을 수 없습니다.", "error")
            return redirect(_safe_redirect_path(next_path))
        if not check_password_hash(user["password_hash"], current_password):
            flash("현재 비밀번호가 올바르지 않습니다.", "error")
            return _redirect_with_password_modal(next_path)
        conn.execute(
            """
            UPDATE users
            SET password_hash = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                generate_password_hash(new_password),
                datetime.now().isoformat(timespec="seconds"),
                user_id,
            ),
        )

    flash("비밀번호를 변경했습니다.", "success")
    base = _safe_redirect_path(next_path)
    joiner = "&" if "?" in base else "?"
    return redirect(f"{base}{joiner}password_changed=1")


@app.route("/admin/users")
@_admin_required
def admin_users():
    org_tree, account_count = _get_org_tree_for_admin()
    return render_template(
        "admin_users.html",
        active_page="admin",
        org_tree=org_tree,
        master_teams=_get_teams(),
        org_people_count=len(_get_org_people()),
        account_count=account_count,
    )


def _admin_db_file_stats():
    path = USER_DB_PATH
    size = path.stat().st_size if path.exists() else 0
    mtime = (
        datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S")
        if path.exists()
        else "—"
    )
    if size >= 1024 * 1024:
        size_label = f"{size / (1024 * 1024):.1f} MB"
    elif size >= 1024:
        size_label = f"{size / 1024:.1f} KB"
    else:
        size_label = f"{size} B"
    return {
        "db_name": path.name,
        "db_size_label": size_label,
        "db_mtime_label": mtime,
        "receipt_count": len(RECEIPTS),
        "detail_count": len(REPORT_DETAILS),
    }


def _xlsx_write_sheet(ws, headers, rows):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="7A1F3D")
    header_font = Font(color="FFFFFF", bold=True)
    for col, title in enumerate(headers, 1):
        cell = ws.cell(1, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r_idx, row in enumerate(rows, 2):
        for c_idx, key in enumerate(headers, 1):
            val = row.get(key, "") if isinstance(row, dict) else (row[c_idx - 1] if c_idx - 1 < len(row) else "")
            if val is None:
                val = ""
            elif isinstance(val, (list, dict)):
                val = json.dumps(val, ensure_ascii=False)
            ws.cell(r_idx, c_idx, val)
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        maxlen = len(str(headers[col - 1] or ""))
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col, max_row=min(ws.max_row, 200)):
            for cell in row:
                maxlen = max(maxlen, min(len(str(cell.value or "")), 40))
        ws.column_dimensions[letter].width = min(max(maxlen + 2, 10), 36)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _build_admin_excel_bytes():
    """관리자 엑셀 백업: 접수·이력·계정·조직."""
    from openpyxl import Workbook

    wb = Workbook()

    # 1) 접수목록
    ws1 = wb.active
    ws1.title = "접수목록"
    receipt_headers = [
        "id", "접수번호", "접수일", "팀코드", "팀명", "고객명", "고객담당자", "연락처",
        "장비명", "규격", "로트", "증상", "AS담당코드", "AS담당명", "접수담당코드", "접수담당명",
        "사무소", "상태", "유무상", "청구비용", "완료일", "등록일시",
        "품질지원", "기술지원", "제품군", "알림발송",
    ]
    receipt_rows = []
    for r in sorted(RECEIPTS, key=lambda x: (x.get("receipt_date") or "", x.get("receipt_no") or "")):
        receipt_rows.append({
            "id": r.get("id"),
            "접수번호": r.get("receipt_no"),
            "접수일": r.get("receipt_date"),
            "팀코드": r.get("team_id"),
            "팀명": r.get("team"),
            "고객명": r.get("customer_name"),
            "고객담당자": r.get("customer_contact"),
            "연락처": r.get("phone"),
            "장비명": r.get("equipment_name"),
            "규격": r.get("spec"),
            "로트": r.get("lot"),
            "증상": r.get("symptoms"),
            "AS담당코드": r.get("as_staff"),
            "AS담당명": r.get("as_staff_name"),
            "접수담당코드": r.get("receipt_staff"),
            "접수담당명": r.get("receipt_staff_name"),
            "사무소": r.get("office"),
            "상태": r.get("status"),
            "유무상": r.get("charge_type"),
            "청구비용": r.get("billing_cost"),
            "완료일": r.get("completed_date"),
            "등록일시": r.get("created_at"),
            "품질지원": 1 if r.get("support_quality") else 0,
            "기술지원": 1 if r.get("support_tech") else 0,
            "제품군": r.get("product_line"),
            "알림발송": 1 if r.get("notice_sent") else 0,
        })
    _xlsx_write_sheet(ws1, receipt_headers, receipt_rows)

    # 2) AS이력상세 (기존 CSV 필드 재사용)
    ws2 = wb.create_sheet("AS이력상세")
    hist_rows = _build_history_csv_rows(year=None, all_lines=True)
    _xlsx_write_sheet(ws2, HISTORY_CSV_FIELDS, hist_rows)

    # 3) 로그인계정
    ws3 = wb.create_sheet("로그인계정")
    user_headers = ["id", "username", "display_name", "role", "active", "created_at", "updated_at"]
    user_rows = []
    with _user_db() as conn:
        for row in conn.execute(
            "SELECT id, username, display_name, role, active, created_at, updated_at FROM users ORDER BY id"
        ).fetchall():
            user_rows.append({k: row[k] for k in user_headers})
    _xlsx_write_sheet(ws3, user_headers, user_rows)

    # 4) 조직인원
    ws4 = wb.create_sheet("조직인원")
    people_headers = [
        "id", "code", "name", "title", "role_note", "team_code",
        "can_receipt", "can_as", "sort_order", "emails", "product_line",
    ]
    people_rows = []
    with _user_db() as conn:
        for row in conn.execute(
            """
            SELECT id, code, name, title, role_note, team_code,
                   can_receipt, can_as, sort_order, IFNULL(email, '') AS email,
                   IFNULL(product_line, 'equipment') AS product_line
            FROM org_people
            ORDER BY sort_order, id
            """
        ).fetchall():
            people_rows.append({
                "id": row["id"],
                "code": row["code"],
                "name": row["name"],
                "title": row["title"] or "",
                "role_note": row["role_note"] or "",
                "team_code": row["team_code"] or "",
                "can_receipt": int(row["can_receipt"] or 0),
                "can_as": int(row["can_as"] or 0),
                "sort_order": row["sort_order"],
                "emails": ", ".join(_parse_person_emails(row["email"])),
                "product_line": row["product_line"] or "equipment",
            })
    _xlsx_write_sheet(ws4, people_headers, people_rows)

    # 5) 부서
    ws5 = wb.create_sheet("부서")
    team_headers = ["id", "code", "name", "category", "sort_order", "note"]
    team_rows = []
    with _user_db() as conn:
        cols = [r[1] for r in conn.execute("PRAGMA table_info(teams)").fetchall()]
        has_note = "note" in cols
        sql = "SELECT id, code, name, category, sort_order" + (", note" if has_note else "") + " FROM teams ORDER BY sort_order, id"
        for row in conn.execute(sql).fetchall():
            team_rows.append({
                "id": row["id"],
                "code": row["code"],
                "name": row["name"],
                "category": row["category"],
                "sort_order": row["sort_order"],
                "note": row["note"] if has_note else "",
            })
    _xlsx_write_sheet(ws5, team_headers, team_rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@app.route("/admin/backup")
@_admin_required
def admin_backup():
    stats = _admin_db_file_stats()
    monitor = _backup_monitor_public()
    return render_template(
        "admin_backup.html",
        active_page="admin_backup",
        backup_monitor=monitor,
        backup_files=_list_backup_archive_files(),
        **stats,
    )


@app.route("/admin/backup/db")
@_admin_required
def admin_backup_db():
    """SQLite DB 파일 백업 다운로드."""
    import tempfile

    try:
        _save_service_data()
    except Exception:
        pass
    if not USER_DB_PATH.exists():
        flash("DB 파일을 찾을 수 없습니다.", "error")
        return redirect(url_for("admin_backup"))

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"pns_as_backup_{stamp}.db"
    with tempfile.NamedTemporaryFile(delete=False, suffix=".db") as tf:
        tmp_path = Path(tf.name)
    try:
        with SERVICE_DATA_LOCK:
            src = sqlite3.connect(str(USER_DB_PATH))
            try:
                dst = sqlite3.connect(str(tmp_path))
                try:
                    src.backup(dst)
                finally:
                    dst.close()
            finally:
                src.close()
        payload = io.BytesIO(tmp_path.read_bytes())
        payload.seek(0)
        return send_file(
            payload,
            mimetype="application/octet-stream",
            as_attachment=True,
            download_name=filename,
        )
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass


@app.route("/admin/backup/excel")
@_admin_required
def admin_backup_excel():
    """접수·이력·계정 엑셀(xlsx) 내보내기."""
    try:
        _save_service_data()
    except Exception:
        pass
    try:
        buf = _build_admin_excel_bytes()
    except ImportError:
        flash("엑셀 내보내기에 openpyxl이 필요합니다. (pip install openpyxl)", "error")
        return redirect(url_for("admin_backup"))
    except Exception as exc:
        flash(f"엑셀 생성에 실패했습니다: {exc}", "error")
        return redirect(url_for("admin_backup"))
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"pns_as_export_{stamp}.xlsx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


def _admin_backup_dir():
    path = BASE_DIR / "db_backups"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _backup_monitor_path():
    return _admin_backup_dir() / "backup_monitor.json"


def _backup_daily_dir():
    path = _admin_backup_dir() / "daily"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _backup_manual_dir():
    path = _admin_backup_dir() / "manual"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _format_bytes_label(size):
    size = int(size or 0)
    if size >= 1024 * 1024:
        return f"{size / (1024 * 1024):.1f} MB"
    if size >= 1024:
        return f"{size / 1024:.1f} KB"
    return f"{size} B"


_BACKUP_MONITOR_LOCK = RLock()
_AUTO_BACKUP_THREAD = None
_AUTO_BACKUP_STARTED = False
_BACKUP_MONITOR_META_KEY = "backup_monitor_json"


def _default_backup_monitor_state():
    return {
        "enabled": True,
        "hour": 2,
        "minute": 0,
        "keep_days": 30,
        "last_run_at": "",
        "last_success_at": "",
        "last_error": "",
        "last_file": "",
        "last_size": 0,
        "last_trigger": "",
        "history": [],
    }


def _load_backup_monitor_state():
    state = _default_backup_monitor_state()
    raw = None
    try:
        with _user_db() as conn:
            row = conn.execute(
                "SELECT value FROM app_meta WHERE key = ?",
                (_BACKUP_MONITOR_META_KEY,),
            ).fetchone()
            if row and row["value"]:
                raw = row["value"]
    except Exception:
        raw = None
    # 이전 파일 저장분 마이그레이션
    if raw is None:
        legacy = _backup_monitor_path()
        if legacy.exists():
            try:
                raw = legacy.read_text(encoding="utf-8")
            except OSError:
                raw = None
    if raw:
        try:
            parsed = json.loads(raw)
            if isinstance(parsed, dict):
                state.update(parsed)
        except json.JSONDecodeError:
            pass
    state["enabled"] = bool(state.get("enabled", True))
    try:
        state["hour"] = max(0, min(23, int(state.get("hour", 2))))
    except (TypeError, ValueError):
        state["hour"] = 2
    try:
        state["minute"] = max(0, min(59, int(state.get("minute", 0))))
    except (TypeError, ValueError):
        state["minute"] = 0
    try:
        state["keep_days"] = max(1, min(365, int(state.get("keep_days", 30))))
    except (TypeError, ValueError):
        state["keep_days"] = 30
    if not isinstance(state.get("history"), list):
        state["history"] = []
    return state


def _save_backup_monitor_state(state):
    payload = dict(state)
    payload["history"] = list(payload.get("history") or [])[:60]
    text = json.dumps(payload, ensure_ascii=False)
    with _BACKUP_MONITOR_LOCK:
        with _user_db() as conn:
            conn.execute(
                "INSERT OR REPLACE INTO app_meta (key, value) VALUES (?, ?)",
                (_BACKUP_MONITOR_META_KEY, text),
            )
            conn.commit()


def _append_backup_history(state, entry):
    history = list(state.get("history") or [])
    history.insert(0, entry)
    state["history"] = history[:60]


def _snapshot_users_db_to(target_path):
    """운영 users.db를 target_path로 스냅샷 복사. 실제 저장된 Path 반환."""
    target_path = Path(target_path)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        _save_service_data()
    except Exception:
        pass
    if not USER_DB_PATH.exists():
        raise FileNotFoundError("users.db 파일이 없습니다.")

    tmp_path = target_path.with_name(target_path.name + ".tmp")
    try:
        if tmp_path.exists():
            tmp_path.unlink()
    except OSError:
        tmp_path = target_path.with_name(
            f"{target_path.stem}_{datetime.now().strftime('%H%M%S')}.tmp"
        )

    with SERVICE_DATA_LOCK:
        src = sqlite3.connect(str(USER_DB_PATH))
        try:
            dst = sqlite3.connect(str(tmp_path))
            try:
                src.backup(dst)
            finally:
                dst.close()
        finally:
            src.close()

    final_path = target_path
    try:
        os.replace(tmp_path, final_path)
    except OSError:
        try:
            final_path.unlink(missing_ok=True)
        except OSError:
            pass
        try:
            os.replace(tmp_path, final_path)
        except OSError:
            final_path = target_path.with_name(
                f"{target_path.stem}_{datetime.now().strftime('%H%M%S')}{target_path.suffix}"
            )
            os.replace(tmp_path, final_path)

    return final_path


def _prune_old_backups(keep_days):
    cutoff = datetime.now() - timedelta(days=int(keep_days))
    removed = 0
    root = _admin_backup_dir()
    patterns = (
        root.glob("users_pre_restore_*.db"),
        _backup_daily_dir().glob("pns_as_daily_*.db"),
        _backup_manual_dir().glob("pns_as_manual_*.db"),
    )
    for group in patterns:
        for path in group:
            try:
                mtime = datetime.fromtimestamp(path.stat().st_mtime)
                if mtime < cutoff:
                    path.unlink(missing_ok=True)
                    removed += 1
            except OSError:
                continue
    return removed


def _daily_backup_path_for(day=None):
    day = day or date.today()
    return _backup_daily_dir() / f"pns_as_daily_{day.strftime('%Y%m%d')}.db"


def _run_db_backup(trigger="auto", force=False):
    """자동/수동 DB 백업 실행. 성공 시 파일 Path 반환."""
    trigger = (trigger or "auto").strip().lower()
    if trigger not in ("auto", "manual"):
        trigger = "auto"

    state = _load_backup_monitor_state()
    now = datetime.now()
    now_label = now.strftime("%Y-%m-%d %H:%M:%S")
    state["last_run_at"] = now_label
    state["last_trigger"] = trigger

    if trigger == "auto":
        target = _daily_backup_path_for(now.date())
        if target.exists() and not force:
            state["last_success_at"] = state.get("last_success_at") or now_label
            state["last_error"] = ""
            state["last_file"] = str(target.relative_to(_admin_backup_dir())).replace("\\", "/")
            state["last_size"] = target.stat().st_size
            _save_backup_monitor_state(state)
            return target
    else:
        stamp = now.strftime("%Y%m%d_%H%M%S")
        target = _backup_manual_dir() / f"pns_as_manual_{stamp}.db"

    try:
        target = _snapshot_users_db_to(target)
        size = target.stat().st_size
        rel = str(target.relative_to(_admin_backup_dir())).replace("\\", "/")
        state["last_success_at"] = now_label
        state["last_error"] = ""
        state["last_file"] = rel
        state["last_size"] = size
        _append_backup_history(
            state,
            {
                "at": now_label,
                "ok": True,
                "trigger": trigger,
                "file": rel,
                "size": size,
                "size_label": _format_bytes_label(size),
                "error": "",
            },
        )
        _prune_old_backups(state.get("keep_days", 30))
        _save_backup_monitor_state(state)
        return target
    except Exception as exc:
        state["last_error"] = str(exc)
        _append_backup_history(
            state,
            {
                "at": now_label,
                "ok": False,
                "trigger": trigger,
                "file": "",
                "size": 0,
                "error": str(exc),
            },
        )
        _save_backup_monitor_state(state)
        raise


def _backup_health(state=None):
    state = state or _load_backup_monitor_state()
    now = datetime.now()
    last_ok = (state.get("last_success_at") or "").strip()
    last_err = (state.get("last_error") or "").strip()
    last_run = (state.get("last_run_at") or "").strip()
    enabled = bool(state.get("enabled", True))
    has_today = _daily_backup_path_for(now.date()).exists()
    hour = int(state.get("hour", 2))
    minute = int(state.get("minute", 0))
    scheduled_passed = (now.hour, now.minute) >= (hour, minute)

    if last_err and last_run and (not last_ok or last_run > last_ok):
        return "error", "최근 백업 실패"

    if not enabled:
        return "off", "자동 백업 중지"

    if has_today:
        return "ok", "오늘 일일 백업 완료"

    if scheduled_passed:
        return "warn", "오늘 일일 백업 미실행"

    return "ok", "대기 중 (예정 시각 전)"


def _next_backup_label(state=None):
    state = state or _load_backup_monitor_state()
    if not state.get("enabled", True):
        return "—"
    now = datetime.now()
    target = now.replace(
        hour=int(state.get("hour", 2)),
        minute=int(state.get("minute", 0)),
        second=0,
        microsecond=0,
    )
    if now >= target:
        target = target + timedelta(days=1)
    # 오늘 파일이 아직 없고 예정 시각이 지났으면 '즉시(미실행 보정)'
    if not _daily_backup_path_for(now.date()).exists() and now >= now.replace(
        hour=int(state.get("hour", 2)),
        minute=int(state.get("minute", 0)),
        second=0,
        microsecond=0,
    ):
        return "미실행 보정 대기(수분 내)"
    return target.strftime("%Y-%m-%d %H:%M")


def _backup_monitor_public():
    state = _load_backup_monitor_state()
    health_code, health_label = _backup_health(state)
    return {
        "enabled": bool(state.get("enabled", True)),
        "hour": int(state.get("hour", 2)),
        "minute": int(state.get("minute", 0)),
        "keep_days": int(state.get("keep_days", 30)),
        "schedule_label": f"{int(state.get('hour', 2)):02d}:{int(state.get('minute', 0)):02d}",
        "last_run_at": state.get("last_run_at") or "—",
        "last_success_at": state.get("last_success_at") or "—",
        "last_error": state.get("last_error") or "",
        "last_file": state.get("last_file") or "—",
        "last_size_label": _format_bytes_label(state.get("last_size") or 0)
        if state.get("last_file")
        else "—",
        "last_trigger": state.get("last_trigger") or "—",
        "next_run_label": _next_backup_label(state),
        "health_code": health_code,
        "health_label": health_label,
        "today_exists": _daily_backup_path_for().exists(),
        "history": state.get("history") or [],
    }


def _list_backup_archive_files(limit=40):
    root = _admin_backup_dir().resolve()
    rows = []
    for path in root.rglob("*.db"):
        try:
            if path.name == "backup_monitor.json":
                continue
            rel = str(path.relative_to(root)).replace("\\", "/")
            st = path.stat()
            kind = "기타"
            name = path.name
            if name.startswith("pns_as_daily_"):
                kind = "일일자동"
            elif name.startswith("pns_as_manual_"):
                kind = "수동"
            elif name.startswith("users_pre_restore_"):
                kind = "복원전"
            rows.append(
                {
                    "rel": rel,
                    "name": name,
                    "kind": kind,
                    "size": st.st_size,
                    "size_label": _format_bytes_label(st.st_size),
                    "mtime": datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
                    "mtime_ts": st.st_mtime,
                }
            )
        except OSError:
            continue
    rows.sort(key=lambda x: x["mtime_ts"], reverse=True)
    return rows[:limit]


def _resolve_backup_archive_path(rel):
    root = _admin_backup_dir().resolve()
    rel = (rel or "").replace("\\", "/").lstrip("/")
    if not rel or ".." in rel.split("/"):
        return None
    target = (root / rel).resolve()
    try:
        target.relative_to(root)
    except ValueError:
        return None
    if not target.is_file() or target.suffix.casefold() != ".db":
        return None
    return target


def _auto_backup_loop():
    """매일 지정 시각에 DB 스냅샷. 서버 기동 중이면 미실행분 보정."""
    while True:
        try:
            state = _load_backup_monitor_state()
            if state.get("enabled", True):
                now = datetime.now()
                due = (now.hour, now.minute) >= (
                    int(state.get("hour", 2)),
                    int(state.get("minute", 0)),
                )
                today_path = _daily_backup_path_for(now.date())
                if due and not today_path.exists():
                    try:
                        _run_db_backup(trigger="auto")
                    except Exception:
                        pass
        except Exception:
            pass
        time.sleep(30)


def _start_auto_backup_scheduler():
    global _AUTO_BACKUP_THREAD, _AUTO_BACKUP_STARTED
    if _AUTO_BACKUP_STARTED:
        return
    # Flask debug reloader: 감시 부모 프로세스가 아닌 실행 프로세스에서만 기동
    if app.debug and os.environ.get("WERKZEUG_RUN_MAIN") != "true":
        return
    _AUTO_BACKUP_STARTED = True
    _admin_backup_dir()
    _backup_daily_dir()
    _backup_manual_dir()
    _AUTO_BACKUP_THREAD = Thread(
        target=_auto_backup_loop,
        name="pns-daily-db-backup",
        daemon=True,
    )
    _AUTO_BACKUP_THREAD.start()


@app.route("/admin/backup/auto/settings", methods=["POST"])
@_admin_required
def admin_backup_auto_settings():
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_backup"))

    state = _load_backup_monitor_state()
    state["enabled"] = request.form.get("enabled") in ("1", "on", "true", "yes")
    try:
        state["hour"] = max(0, min(23, int(request.form.get("hour", state.get("hour", 2)))))
    except (TypeError, ValueError):
        flash("백업 시각(시)이 올바르지 않습니다.", "error")
        return redirect(url_for("admin_backup"))
    try:
        state["minute"] = max(0, min(59, int(request.form.get("minute", state.get("minute", 0)))))
    except (TypeError, ValueError):
        flash("백업 시각(분)이 올바르지 않습니다.", "error")
        return redirect(url_for("admin_backup"))
    try:
        state["keep_days"] = max(1, min(365, int(request.form.get("keep_days", state.get("keep_days", 30)))))
    except (TypeError, ValueError):
        flash("보관 일수가 올바르지 않습니다.", "error")
        return redirect(url_for("admin_backup"))

    _save_backup_monitor_state(state)
    _prune_old_backups(state["keep_days"])
    flash(
        f"자동 백업 설정을 저장했습니다. "
        f"({'사용' if state['enabled'] else '중지'} · "
        f"{state['hour']:02d}:{state['minute']:02d} · {state['keep_days']}일 보관)",
        "success",
    )
    return redirect(url_for("admin_backup"))


@app.route("/admin/backup/auto/run", methods=["POST"])
@_admin_required
def admin_backup_auto_run():
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_backup"))
    try:
        path = _run_db_backup(trigger="manual", force=True)
        flash(f"수동 DB 백업을 완료했습니다: {path.name}", "success")
    except Exception as exc:
        flash(f"수동 백업에 실패했습니다: {exc}", "error")
    return redirect(url_for("admin_backup"))


@app.route("/admin/backup/archive/<path:rel>")
@_admin_required
def admin_backup_archive_download(rel):
    target = _resolve_backup_archive_path(rel)
    if not target:
        flash("백업 파일을 찾을 수 없습니다.", "error")
        return redirect(url_for("admin_backup"))
    return send_file(
        target,
        mimetype="application/octet-stream",
        as_attachment=True,
        download_name=target.name,
    )


@app.route("/admin/backup/archive/delete", methods=["POST"])
@_admin_required
def admin_backup_archive_delete():
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_backup"))
    rel = (request.form.get("rel") or "").strip()
    target = _resolve_backup_archive_path(rel)
    if not target:
        flash("삭제할 백업 파일을 찾을 수 없습니다.", "error")
        return redirect(url_for("admin_backup"))
    try:
        target.unlink(missing_ok=True)
        flash(f"백업 파일을 삭제했습니다: {target.name}", "success")
    except OSError as exc:
        flash(f"삭제에 실패했습니다: {exc}", "error")
    return redirect(url_for("admin_backup"))


@app.route("/admin/backup/archive/restore", methods=["POST"])
@_admin_required
def admin_backup_archive_restore():
    """보관 중인 백업 파일로 운영 DB 복원."""
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_backup"))
    rel = (request.form.get("rel") or "").strip()
    target = _resolve_backup_archive_path(rel)
    if not target:
        flash("복원할 백업 파일을 찾을 수 없습니다.", "error")
        return redirect(url_for("admin_backup"))
    try:
        _validate_sqlite_backup_file(target)
        pre_path = _replace_users_db_from_file(target)
    except Exception as exc:
        flash(f"DB 복원에 실패했습니다: {exc}", "error")
        return redirect(url_for("admin_backup"))
    flash(
        f"DB 복원이 완료되었습니다. ({target.name} → 접수 {len(RECEIPTS)}건 · 이력 {len(REPORT_DETAILS)}건) "
        f"이전 DB는 {pre_path.name} 에 보관했습니다.",
        "success",
    )
    return redirect(url_for("admin_backup"))


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _parse_flag01(value):
    text = _cell_text(value).casefold()
    if text in ("1", "y", "yes", "true", "예", "o", "on"):
        return 1
    if text in ("0", "n", "no", "false", "아니오", "x", "off", ""):
        return 0
    try:
        return 1 if int(float(text)) else 0
    except (TypeError, ValueError):
        return 0


def _parse_parts_from_excel(text):
    """엑셀 부품내역(name|규격:..|수량:..|단가:..|소계:..) → parts 리스트."""
    parts = []
    for chunk in (_cell_text(text) or "").split(";"):
        chunk = chunk.strip()
        if not chunk:
            continue
        bits = [b.strip() for b in chunk.split("|") if b.strip()]
        if not bits:
            continue
        item = {"name": bits[0], "spec": "", "qty": 0, "unit_price": 0, "subtotal": 0}
        for bit in bits[1:]:
            if ":" not in bit:
                continue
            key, raw = bit.split(":", 1)
            key = key.strip()
            raw = raw.strip().replace(",", "")
            if key == "규격":
                item["spec"] = raw
            elif key == "수량":
                try:
                    item["qty"] = int(float(raw))
                except (TypeError, ValueError):
                    item["qty"] = 0
            elif key == "단가":
                try:
                    item["unit_price"] = int(float(raw))
                except (TypeError, ValueError):
                    item["unit_price"] = 0
            elif key == "소계":
                try:
                    item["subtotal"] = int(float(raw))
                except (TypeError, ValueError):
                    item["subtotal"] = 0
        if not item["subtotal"] and item["qty"] and item["unit_price"]:
            item["subtotal"] = int(item["qty"]) * int(item["unit_price"])
        parts.append(item)
    return parts


def _xlsx_read_sheet_dicts(wb, title):
    if title not in wb.sheetnames:
        return []
    ws = wb[title]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [_cell_text(h) for h in next(rows_iter)]
    except StopIteration:
        return []
    if not any(headers):
        return []
    out = []
    for raw in rows_iter:
        if raw is None or all(v is None or _cell_text(v) == "" for v in raw):
            continue
        row = {}
        for idx, key in enumerate(headers):
            if not key:
                continue
            row[key] = raw[idx] if idx < len(raw) else ""
        out.append(row)
    return out


def _label_to_code(options, label, default=""):
    text = _cell_text(label)
    if not text:
        return default
    for code, name in options:
        if text == code or text == name:
            return code
    return default or text


def _compose_ymd(year, month, day):
    y = _cell_text(year)
    m = _cell_text(month).zfill(2) if _cell_text(month) else ""
    d = _cell_text(day).zfill(2) if _cell_text(day) else ""
    if len(y) == 4 and m and d:
        return f"{y}-{m}-{d}"
    return _cell_text(day) if "-" in _cell_text(day) else ""


def _receipt_from_excel_row(row, fallback_id):
    receipt_no = _cell_text(row.get("접수번호"))
    if not receipt_no:
        return None
    try:
        rid = int(float(_cell_text(row.get("id")) or fallback_id))
    except (TypeError, ValueError):
        rid = fallback_id
    charge = _cell_text(row.get("유무상") or row.get("유무상코드")).casefold()
    if charge in ("무상", "free"):
        charge_type = "free"
    elif charge in ("유상", "paid"):
        charge_type = "paid"
    else:
        charge_type = _cell_text(row.get("유무상코드")) or "free"
        if charge_type not in ("free", "paid"):
            charge_type = "free"
    status = _cell_text(row.get("상태") or row.get("AS상태코드")) or "received"
    status_map = {
        "접수": "received",
        "임시저장": "draft",
        "진행중": "in_progress",
        "부품대기": "parts_waiting",
        "보류": "on_hold",
        "완료": "completed",
    }
    status = status_map.get(status, status)
    billing = _cell_text(row.get("청구비용"))
    try:
        billing_cost = int(float(billing.replace(",", ""))) if billing not in ("", None) else None
    except (TypeError, ValueError):
        billing_cost = billing or None
    return {
        "id": rid,
        "receipt_no": receipt_no,
        "receipt_date": _cell_text(row.get("접수일")),
        "team_id": _cell_text(row.get("팀코드")),
        "team": _cell_text(row.get("팀명")),
        "customer_name": _cell_text(row.get("고객명")),
        "customer_contact": _cell_text(row.get("고객담당자")),
        "phone": _cell_text(row.get("연락처")),
        "equipment_name": _cell_text(row.get("장비명")),
        "spec": _cell_text(row.get("규격")),
        "lot": _cell_text(row.get("로트") or row.get("로트번호")),
        "symptoms": _cell_text(row.get("증상")),
        "as_staff": _cell_text(row.get("AS담당코드")),
        "as_staff_name": _cell_text(row.get("AS담당명") or row.get("AS배정담당자")),
        "receipt_staff": _cell_text(row.get("접수담당코드")),
        "receipt_staff_name": _cell_text(row.get("접수담당명") or row.get("접수담당자")),
        "office": _cell_text(row.get("사무소") or row.get("접수사무소")),
        "status": status,
        "charge_type": charge_type,
        "billing_cost": billing_cost,
        "completed_date": _cell_text(row.get("완료일")) or None,
        "created_at": _cell_text(row.get("등록일시")) or datetime.now().isoformat(timespec="seconds"),
        "photos": [],
        "ship_date": _cell_text(row.get("출하일")),
        "support_quality": bool(_parse_flag01(row.get("품질지원") or row.get("품질지원요청"))),
        "support_tech": bool(_parse_flag01(row.get("기술지원") or row.get("기술지원요청"))),
        "created_by_user_id": None,
        "created_by_username": "",
        "severity": "medium",
        "intervention_status": "none",
        "compensation_status": "none",
        "intervention_note": "",
        "compensation_note": "",
        "intervention_by": "",
        "intervention_at": None,
        "escalation_level": "normal",
        "compensation_amount": "",
        "intervention_log": [],
        "billing_margin_rate": None,
        "billing_manual": 0,
        "billing_set_by": "",
        "product_line": plex.normalize_receipt_line(row.get("제품군") or "equipment"),
        "notice_sent": _parse_flag01(row.get("알림발송")),
    }


def _detail_from_history_excel_row(row):
    inspect_date = _compose_ymd(row.get("점검연도"), row.get("점검월"), row.get("점검일"))
    if not inspect_date:
        inspect_date = _cell_text(row.get("점검일"))
    re_repair = _csv_re_repair_label(row.get("재수리여부"))
    if re_repair == "Y":
        re_repair_val = "1회"
    elif re_repair == "N":
        re_repair_val = "해당없음"
    else:
        re_repair_val = _cell_text(row.get("재수리여부")) or "해당없음"
    esc = _label_to_code(v32.ESCALATION_OPTIONS, row.get("에스컬레이션"), "normal")
    return {
        "serial_no": _cell_text(row.get("시리얼번호")),
        "phone": _cell_text(row.get("연락처")),
        "ship_date": _cell_text(row.get("출하일")),
        "inspection_date": inspect_date,
        "inspection_result": _cell_text(row.get("점검결과")),
        "repair_content": _cell_text(row.get("수리내용")),
        "detailed_work": _cell_text(row.get("상세작업내용")),
        "remarks": _cell_text(row.get("비고")),
        "support_history": _cell_text(row.get("지원요청이력")),
        "support_quality": bool(_parse_flag01(row.get("품질지원요청"))),
        "support_tech": bool(_parse_flag01(row.get("기술지원요청"))),
        "support_tech_mgmt": bool(_parse_flag01(row.get("최고기술경영실보고"))),
        "as_technician": _cell_text(row.get("AS처리담당자")),
        "affiliation": _cell_text(row.get("소속")),
        "process_location": _cell_text(row.get("처리위치")),
        "ship_out_date": _cell_text(row.get("출고일")),
        "receipt_staff_name": _cell_text(row.get("접수담당자")),
        "report_date": _cell_text(row.get("보고서작성일")),
        "re_repair": re_repair_val,
        "parts": _parse_parts_from_excel(row.get("부품내역")),
        "photos": [],
        "escalation_level": esc,
        "history_updated_at": datetime.now().isoformat(timespec="seconds"),
    }


def _apply_excel_service_import(wb):
    """엑셀에서 접수·이력만 복원(계정·조직은 유지)."""
    global _next_id

    receipt_rows = _xlsx_read_sheet_dicts(wb, "접수목록")
    history_rows = _xlsx_read_sheet_dicts(wb, "AS이력상세")
    if not receipt_rows and not history_rows:
        raise ValueError("시트 '접수목록' 또는 'AS이력상세'가 없거나 비어 있습니다.")

    new_receipts = []
    seen = set()
    next_fallback = 1
    for row in receipt_rows:
        item = _receipt_from_excel_row(row, next_fallback)
        if not item:
            continue
        if item["receipt_no"] in seen:
            continue
        seen.add(item["receipt_no"])
        new_receipts.append(item)
        next_fallback = max(next_fallback, int(item["id"]) + 1)

    new_details = {}
    for row in history_rows:
        receipt_no = _cell_text(row.get("접수번호"))
        if not receipt_no:
            continue
        detail = _detail_from_history_excel_row(row)
        new_details[receipt_no] = detail

        # 접수목록에 없으면 이력 시트로 최소 접수 생성
        if receipt_no not in seen:
            item = _receipt_from_excel_row(row, next_fallback)
            if item:
                # 이력 시트 컬럼명 보정
                item["as_staff_name"] = item["as_staff_name"] or _cell_text(row.get("AS배정담당자"))
                item["receipt_staff_name"] = item["receipt_staff_name"] or _cell_text(row.get("접수담당자"))
                item["office"] = item["office"] or _cell_text(row.get("접수사무소"))
                item["status"] = _cell_text(row.get("AS상태코드")) or item["status"]
                status_map = {
                    "접수": "received",
                    "임시저장": "draft",
                    "진행중": "in_progress",
                    "부품대기": "parts_waiting",
                    "보류": "on_hold",
                    "완료": "completed",
                }
                item["status"] = status_map.get(item["status"], item["status"])
                charge_code = _cell_text(row.get("유무상코드"))
                if charge_code in ("free", "paid"):
                    item["charge_type"] = charge_code
                item["completed_date"] = _compose_ymd(
                    row.get("완료연도"), row.get("완료월"), row.get("완료일")
                ) or item.get("completed_date")
                item["receipt_date"] = _compose_ymd(
                    row.get("접수연도"), row.get("접수월"), row.get("접수일")
                ) or item.get("receipt_date")
                item["ship_date"] = detail.get("ship_date") or ""
                item["support_quality"] = detail.get("support_quality")
                item["support_tech"] = detail.get("support_tech")
                item["escalation_level"] = detail.get("escalation_level") or "normal"
                item["intervention_status"] = _label_to_code(
                    INTERVENTION_OPTIONS, row.get("개입상태"), "none"
                )
                item["compensation_status"] = _label_to_code(
                    COMPENSATION_OPTIONS, row.get("보상상태"), "none"
                )
                item["compensation_amount"] = _cell_text(row.get("보상금액"))
                margin = _cell_text(row.get("청구마진율"))
                try:
                    item["billing_margin_rate"] = float(margin) if margin else None
                except (TypeError, ValueError):
                    item["billing_margin_rate"] = None
                item["billing_manual"] = _parse_flag01(row.get("청구수동책정"))
                item["billing_set_by"] = _cell_text(row.get("청구책정자"))
                new_receipts.append(item)
                seen.add(receipt_no)
                next_fallback = max(next_fallback, int(item["id"]) + 1)
        else:
            # 기존 접수에 이력 필드 보강
            for r in new_receipts:
                if r["receipt_no"] != receipt_no:
                    continue
                if detail.get("ship_date") and not r.get("ship_date"):
                    r["ship_date"] = detail["ship_date"]
                if detail.get("phone") and not r.get("phone"):
                    r["phone"] = detail["phone"]
                esc = detail.get("escalation_level")
                if esc:
                    r["escalation_level"] = esc
                r["intervention_status"] = _label_to_code(
                    INTERVENTION_OPTIONS, row.get("개입상태"), r.get("intervention_status") or "none"
                )
                r["compensation_status"] = _label_to_code(
                    COMPENSATION_OPTIONS, row.get("보상상태"), r.get("compensation_status") or "none"
                )
                if _cell_text(row.get("보상금액")):
                    r["compensation_amount"] = _cell_text(row.get("보상금액"))
                margin = _cell_text(row.get("청구마진율"))
                if margin:
                    try:
                        r["billing_margin_rate"] = float(margin)
                    except (TypeError, ValueError):
                        pass
                if _cell_text(row.get("청구책정자")):
                    r["billing_set_by"] = _cell_text(row.get("청구책정자"))
                r["billing_manual"] = _parse_flag01(row.get("청구수동책정")) or r.get("billing_manual") or 0
                break

    if not new_receipts:
        raise ValueError("가져올 접수 데이터가 없습니다.")

    with SERVICE_DATA_LOCK:
        RECEIPTS.clear()
        RECEIPTS.extend(sorted(new_receipts, key=lambda x: int(x.get("id") or 0)))
        REPORT_DETAILS.clear()
        REPORT_DETAILS.update(new_details)
        _next_id = max((int(item.get("id", 0)) for item in RECEIPTS), default=0) + 1
        _save_service_data()

    return len(RECEIPTS), len(REPORT_DETAILS)


def _validate_sqlite_backup_file(path):
    header = path.read_bytes()[:16]
    if not header.startswith(b"SQLite format 3"):
        raise ValueError("SQLite DB 파일(.db)이 아닙니다.")
    uri = path.resolve().as_uri() + "?mode=ro"
    conn = sqlite3.connect(uri, uri=True)
    try:
        tables = {
            r[0]
            for r in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            ).fetchall()
        }
        required = {"users", "receipts", "report_details"}
        missing = sorted(required - tables)
        if missing:
            raise ValueError("필수 테이블이 없습니다: " + ", ".join(missing))
        user_count = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        if user_count < 1:
            raise ValueError("users 테이블이 비어 있습니다.")
    finally:
        conn.close()


def _replace_users_db_from_file(src_path):
    """업로드된 DB로 users.db 내용을 덮어쓰고 메모리 데이터를 다시 로드.

    Windows에서 파일이 사용 중이면 os.replace가 실패하므로
    SQLite backup API로 기존 파일에 내용을 복사한다.
    """
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = _admin_backup_dir()
    pre_path = backup_dir / f"users_pre_restore_{stamp}.db"

    with SERVICE_DATA_LOCK:
        # 현재 DB 안전 백업
        if USER_DB_PATH.exists():
            live = sqlite3.connect(str(USER_DB_PATH))
            try:
                snap = sqlite3.connect(str(pre_path))
                try:
                    live.backup(snap)
                finally:
                    snap.close()
                try:
                    live.execute("PRAGMA wal_checkpoint(TRUNCATE)")
                except Exception:
                    pass
            finally:
                live.close()

        # 업로드본 → 운영 DB (동일 파일 경로에 내용 복원)
        src = sqlite3.connect(str(src_path))
        try:
            dst = sqlite3.connect(str(USER_DB_PATH))
            try:
                src.backup(dst)
                try:
                    dst.execute("PRAGMA wal_checkpoint(TRUNCATE)")
                except Exception:
                    pass
            finally:
                dst.close()
        finally:
            src.close()

        _init_user_db()
        _load_service_data()

    return pre_path


@app.route("/admin/backup/db/upload", methods=["POST"])
@_admin_required
def admin_backup_db_upload():
    """DB 백업 파일 업로드·복원."""
    import tempfile

    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_backup"))

    file = request.files.get("db_file")
    if not file or not (file.filename or "").strip():
        flash("업로드할 DB 파일을 선택해 주세요.", "error")
        return redirect(url_for("admin_backup"))

    filename = secure_filename(file.filename) or "upload.db"
    lower = filename.casefold()
    if not (lower.endswith(".db") or lower.endswith(".sqlite") or lower.endswith(".sqlite3")):
        # secure_filename이 한글만 있는 이름을 비울 수 있어 원본 확장자로 재검사
        orig = (file.filename or "").casefold()
        if not (orig.endswith(".db") or orig.endswith(".sqlite") or orig.endswith(".sqlite3")):
            flash("확장자가 .db / .sqlite 인 파일만 업로드할 수 있습니다.", "error")
            return redirect(url_for("admin_backup"))
        filename = "upload.db"

    with tempfile.NamedTemporaryFile(delete=False, suffix=".db") as tf:
        tmp_path = Path(tf.name)
    try:
        file.save(tmp_path)
        size = tmp_path.stat().st_size
        if size <= 0:
            flash("빈 파일은 업로드할 수 없습니다.", "error")
            return redirect(url_for("admin_backup"))
        if size > 200 * 1024 * 1024:
            flash("DB 파일 크기가 너무 큽니다. (최대 200MB)", "error")
            return redirect(url_for("admin_backup"))
        try:
            _validate_sqlite_backup_file(tmp_path)
            pre_path = _replace_users_db_from_file(tmp_path)
        except Exception as exc:
            flash(f"DB 복원에 실패했습니다: {exc}", "error")
            return redirect(url_for("admin_backup"))
    finally:
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass

    flash(
        f"DB 복원이 완료되었습니다. (접수 {len(RECEIPTS)}건 · 이력 {len(REPORT_DETAILS)}건) "
        f"이전 DB는 {pre_path.name} 에 보관했습니다.",
        "success",
    )
    return redirect(url_for("admin_backup"))


@app.route("/admin/backup/excel/upload", methods=["POST"])
@_admin_required
def admin_backup_excel_upload():
    """엑셀에서 접수·이력 데이터 복원(계정·조직은 유지)."""
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_backup"))

    file = request.files.get("excel_file")
    if not file or not (file.filename or "").strip():
        flash("업로드할 엑셀 파일을 선택해 주세요.", "error")
        return redirect(url_for("admin_backup"))

    filename = secure_filename(file.filename) or "upload.xlsx"
    if not filename.casefold().endswith(".xlsx"):
        orig = (file.filename or "").casefold()
        if not orig.endswith(".xlsx"):
            flash("확장자가 .xlsx 인 파일만 업로드할 수 있습니다.", "error")
            return redirect(url_for("admin_backup"))
        filename = "upload.xlsx"

    try:
        from openpyxl import load_workbook
    except ImportError:
        flash("엑셀 업로드에 openpyxl이 필요합니다. (pip install openpyxl)", "error")
        return redirect(url_for("admin_backup"))

    payload = file.read()
    if not payload:
        flash("빈 파일은 업로드할 수 없습니다.", "error")
        return redirect(url_for("admin_backup"))
    if len(payload) > 80 * 1024 * 1024:
        flash("엑셀 파일 크기가 너무 큽니다. (최대 80MB)", "error")
        return redirect(url_for("admin_backup"))

    try:
        wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            receipt_count, detail_count = _apply_excel_service_import(wb)
        finally:
            wb.close()
    except Exception as exc:
        flash(f"엑셀 복원에 실패했습니다: {exc}", "error")
        return redirect(url_for("admin_backup"))

    flash(
        f"엑셀 업로드가 완료되었습니다. 접수 {receipt_count}건 · 이력상세 {detail_count}건으로 교체했습니다. "
        "(계정·조직은 변경되지 않습니다.)",
        "success",
    )
    return redirect(url_for("admin_backup"))


@app.route("/admin/org/department/create", methods=["POST"])
@_admin_required
def admin_org_department_create():
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    name = request.form.get("name", "").strip()
    note = request.form.get("note", "").strip()
    team_category = "domestic"
    if not name:
        flash("부서명을 입력해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    try:
        with _user_db() as conn:
            code = _next_master_code(conn, "teams", "dept")
            next_order = conn.execute(
                "SELECT COALESCE(MAX(sort_order), 0) + 10 FROM teams"
            ).fetchone()[0]
            conn.execute(
                """
                INSERT INTO teams (code, name, category, sort_order, note)
                VALUES (?, ?, ?, ?, ?)
                """,
                (code, name, team_category, next_order, note),
            )
            conn.execute(
                "DELETE FROM app_meta WHERE key = ?",
                (f"team_removed:{code}",),
            )
            conn.commit()
    except sqlite3.IntegrityError:
        flash("부서를 추가하지 못했습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    flash(f"부서 '{name}'을 추가했습니다.", "success")
    return redirect(url_for("admin_users", _anchor="orgMaster"))


@app.route("/admin/org/department/<int:item_id>/update", methods=["POST"])
@_admin_required
def admin_org_department_update(item_id):
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    name = request.form.get("name", "").strip()
    if not name:
        flash("부서명을 입력해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    with _user_db() as conn:
        item = conn.execute("SELECT id FROM teams WHERE id = ?", (item_id,)).fetchone()
        if not item:
            flash("수정할 부서를 찾을 수 없습니다.", "error")
            return redirect(url_for("admin_users", _anchor="orgMaster"))
        conn.execute(
            """
            UPDATE teams
            SET name = ?
            WHERE id = ?
            """,
            (name, item_id),
        )

    flash(f"부서 '{name}'을 수정했습니다.", "success")
    return redirect(url_for("admin_users", _anchor="orgMaster"))


@app.route("/admin/org/department/<int:item_id>/delete", methods=["POST"])
@_admin_required
def admin_org_department_delete(item_id):
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    with _user_db() as conn:
        item = conn.execute("SELECT * FROM teams WHERE id = ?", (item_id,)).fetchone()
        if not item:
            flash("삭제할 부서를 찾을 수 없습니다.", "error")
            return redirect(url_for("admin_users", _anchor="orgMaster"))

        team_code = item["code"]
        team_name = item["name"]
        people_count = conn.execute(
            "SELECT COUNT(*) FROM org_people WHERE team_code = ? COLLATE NOCASE",
            (team_code,),
        ).fetchone()[0]

        # 인원은 유지하고 소속만 해제(미배정)
        if people_count:
            conn.execute(
                "UPDATE org_people SET team_code = '' "
                "WHERE team_code = ? COLLATE NOCASE",
                (team_code,),
            )
        for table_name in ("receipt_staff", "as_staff"):
            try:
                conn.execute(
                    f"UPDATE {table_name} SET team_code = '' "
                    f"WHERE team_code = ? COLLATE NOCASE",
                    (team_code,),
                )
            except sqlite3.OperationalError:
                pass

        conn.execute("DELETE FROM teams WHERE id = ?", (item_id,))
        conn.execute(
            "INSERT OR REPLACE INTO app_meta (key, value) VALUES (?, ?)",
            (f"team_removed:{team_code}", "1"),
        )
        conn.commit()

    if people_count:
        flash(
            f"부서 '{team_name}'을 삭제했습니다. "
            f"소속 인원 {people_count}명은 미배정으로 유지됩니다.",
            "success",
        )
    else:
        flash(f"부서 '{team_name}'을 삭제했습니다.", "success")
    return redirect(url_for("admin_users", _anchor="orgMaster"))


@app.route("/admin/org/person/create", methods=["POST"])
@_admin_required
def admin_org_person_create():
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    name = request.form.get("name", "").strip()
    title = request.form.get("title", "").strip()
    role_note = request.form.get("role_note", "").strip()
    team_code = request.form.get("team_code", "").strip()
    email = _collect_person_emails_from_form(request.form)
    can_receipt = 1 if request.form.get("can_receipt") else 0
    can_as = 1 if request.form.get("can_as") else 0

    if not name or not team_code:
        flash("이름과 소속 부서를 입력해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    try:
        with _user_db() as conn:
            team = conn.execute(
                "SELECT code FROM teams WHERE code = ? COLLATE NOCASE", (team_code,)
            ).fetchone()
            if not team:
                flash("연결할 부서를 선택해 주세요.", "error")
                return redirect(url_for("admin_users", _anchor="orgMaster"))
            code = _next_master_code(conn, "org_people", "person")
            next_order = conn.execute(
                """
                SELECT COALESCE(MAX(sort_order), 0) + 10
                FROM org_people
                WHERE team_code = ? COLLATE NOCASE
                """,
                (team_code,),
            ).fetchone()[0]
            conn.execute(
                """
                INSERT INTO org_people
                    (code, name, title, role_note, team_code, can_receipt, can_as, sort_order, email)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (code, name, title, role_note, team_code, can_receipt, can_as, next_order, email),
            )
            _ensure_person_login(conn, name)
    except sqlite3.IntegrityError:
        flash("인원을 추가하지 못했습니다. 같은 이름(로그인 ID)이 이미 있을 수 있습니다.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    flash(f"인원 '{name}'을 추가했습니다. (로그인 ID: {name} / PW: 1234)", "success")
    return redirect(url_for("admin_users", _anchor="orgMaster"))


@app.route("/admin/org/person/<int:item_id>/update", methods=["POST"])
@_admin_required
def admin_org_person_update(item_id):
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    name = request.form.get("name", "").strip()
    title = request.form.get("title", "").strip()
    role_note = request.form.get("role_note", "").strip()
    team_code = request.form.get("team_code", "").strip()
    email = _collect_person_emails_from_form(request.form)
    password = _normalize_login_password(request.form.get("password", ""))
    can_receipt = 1 if request.form.get("can_receipt") else 0
    can_as = 1 if request.form.get("can_as") else 0
    account_active = 1 if request.form.get("account_active") else 0

    if not name or not team_code:
        flash("이름과 소속 부서를 입력해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))
    if password and len(password) < 4:
        flash("비밀번호는 4자 이상이어야 합니다.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    with _user_db() as conn:
        item = conn.execute(
            "SELECT id, name FROM org_people WHERE id = ?", (item_id,)
        ).fetchone()
        if not item:
            flash("수정할 인원을 찾을 수 없습니다.", "error")
            return redirect(url_for("admin_users", _anchor="orgMaster"))
        team = conn.execute(
            "SELECT code FROM teams WHERE code = ? COLLATE NOCASE", (team_code,)
        ).fetchone()
        if not team:
            flash("연결할 부서를 선택해 주세요.", "error")
            return redirect(url_for("admin_users", _anchor="orgMaster"))
        old_name = item["name"]
        conn.execute(
            """
            UPDATE org_people
            SET name = ?, title = ?, role_note = ?, team_code = ?,
                can_receipt = ?, can_as = ?, email = ?
            WHERE id = ?
            """,
            (name, title, role_note, team_code, can_receipt, can_as, email, item_id),
        )
        _ensure_person_login(
            conn, name, old_name=old_name, password=password or None
        )
        user = conn.execute(
            "SELECT id, role FROM users WHERE username = ? COLLATE NOCASE",
            (name,),
        ).fetchone()
        if user and user["role"] != "admin":
            conn.execute(
                """
                UPDATE users SET active = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    account_active,
                    datetime.now().isoformat(timespec="seconds"),
                    user["id"],
                ),
            )

    if password:
        flash(f"인원 '{name}'을 수정하고 비밀번호를 변경했습니다.", "success")
    else:
        flash(f"인원 '{name}'을 수정했습니다.", "success")
    return redirect(url_for("admin_users", _anchor="orgMaster"))


@app.route("/admin/org/person/<int:item_id>/delete", methods=["POST"])
@_admin_required
def admin_org_person_delete(item_id):
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    with _user_db() as conn:
        item = conn.execute("SELECT * FROM org_people WHERE id = ?", (item_id,)).fetchone()
        if not item:
            flash("삭제할 인원을 찾을 수 없습니다.", "error")
            return redirect(url_for("admin_users", _anchor="orgMaster"))
        conn.execute("DELETE FROM org_people WHERE id = ?", (item_id,))
        _delete_person_login(conn, item["name"])
        if (item["team_code"] or "") == plex.SALES_OFFICE_TEAM:
            office_code = (item["code"] or "").strip()
            if office_code:
                conn.execute(
                    "INSERT OR REPLACE INTO app_meta (key, value) VALUES (?, ?)",
                    (f"office_removed:{office_code}", "1"),
                )
        conn.commit()

    flash(f"인원 '{item['name']}'을 삭제했습니다.", "success")
    return redirect(url_for("admin_users", _anchor="orgMaster"))


@app.route("/admin/users/create", methods=["POST"])
@_admin_required
def admin_user_create():
    flash("로그인 계정은 「인원 추가」로 등록해 주세요. (부서 · 인원 목록)", "error")
    return redirect(url_for("admin_users", _anchor="orgMaster"))


@app.route("/admin/users/<int:user_id>/update", methods=["POST"])
@_admin_required
def admin_user_update(user_id):
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    username = request.form.get("username", "").strip()
    display_name = request.form.get("display_name", "").strip()
    password = request.form.get("password", "")
    role = request.form.get("role", "staff")
    active = 1 if request.form.get("active") else 0

    if (
        len(username) < 3
        or not all(ch.isalnum() or ch in "._-" for ch in username)
        or not display_name
        or role not in ("admin", "staff")
        or (password and len(password) < 4)
    ):
        flash("입력값을 확인해 주세요. 새 비밀번호는 4자 이상이어야 합니다.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    with _user_db() as conn:
        target = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if not target:
            flash("계정을 찾을 수 없습니다.", "error")
            return redirect(url_for("admin_users", _anchor="orgMaster"))

        if user_id == session.get("user_id"):
            active = 1
            role = "admin"

        if target["role"] == "admin" and (role != "admin" or not active):
            admin_count = conn.execute(
                "SELECT COUNT(*) FROM users WHERE role = 'admin' AND active = 1"
            ).fetchone()[0]
            if admin_count <= 1:
                flash("마지막 활성 관리자 계정은 권한을 변경하거나 비활성화할 수 없습니다.", "error")
                return redirect(url_for("admin_users", _anchor="orgMaster"))

        params = [username, display_name, role, active, datetime.now().isoformat(timespec="seconds")]
        sql = """
            UPDATE users
            SET username = ?, display_name = ?, role = ?, active = ?, updated_at = ?
        """
        if password:
            sql += ", password_hash = ?"
            params.append(generate_password_hash(password))
        sql += " WHERE id = ?"
        params.append(user_id)
        try:
            conn.execute(sql, params)
        except sqlite3.IntegrityError:
            flash("이미 사용 중인 아이디입니다.", "error")
            return redirect(url_for("admin_users", _anchor="orgMaster"))

    if user_id == session.get("user_id"):
        session["username"] = username
        session["user_name"] = display_name
        session["user_role"] = "admin"
    flash(f"계정 '{username}' 정보를 저장했습니다.", "success")
    return redirect(url_for("admin_users", _anchor="orgMaster"))


@app.route("/admin/users/<int:user_id>/delete", methods=["POST"])
@_admin_required
def admin_user_delete(user_id):
    if not _valid_admin_csrf():
        flash("요청이 만료되었습니다. 다시 시도해 주세요.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))
    if user_id == session.get("user_id"):
        flash("현재 로그인한 관리자 계정은 삭제할 수 없습니다.", "error")
        return redirect(url_for("admin_users", _anchor="orgMaster"))

    with _user_db() as conn:
        target = conn.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if not target:
            flash("계정을 찾을 수 없습니다.", "error")
            return redirect(url_for("admin_users", _anchor="orgMaster"))
        if target["role"] == "admin":
            admin_count = conn.execute(
                "SELECT COUNT(*) FROM users WHERE role = 'admin' AND active = 1"
            ).fetchone()[0]
            if target["active"] and admin_count <= 1:
                flash("마지막 활성 관리자 계정은 삭제할 수 없습니다.", "error")
                return redirect(url_for("admin_users", _anchor="orgMaster"))
        conn.execute("DELETE FROM users WHERE id = ?", (user_id,))

    flash(f"계정 '{target['username']}'을 삭제했습니다.", "success")
    return redirect(url_for("admin_users", _anchor="orgMaster"))





@app.route("/logout")

def logout():

    session.clear()

    return redirect(url_for("index"))


# v3.2 확장 설치 (권한·알림·청구·자동감지 API)
from v32_app_ext import install as _install_v32

_install_v32(__import__(__name__))

# 2026-09-05 확장 (이관·공동담당·지원피드백·장비통계·담당자실적)
from as_ext_0905 import install as _install_0905

_install_0905(__import__(__name__))

# 일일 자동 DB 백업 스케줄러 (데몬 스레드)
_start_auto_backup_scheduler()


@app.route("/rods", endpoint="rods_home")
def rods_home():
    _enter_product_line("rods")
    if session.get("user_id"):
        return render_template(
            "index.html",
            active_page="rods_home",
            dashboard=_build_home_dashboard(),
        )
    return render_template(
        "rods_home.html",
        active_page="rods_home",
    )


@app.route("/rods/sales", endpoint="rods_sales")
def rods_sales():
    _enter_product_line("rods")
    return redirect(url_for("sales", **request.args.to_dict(flat=True)))


@app.route("/rods/history", endpoint="rods_history")
def rods_history():
    _enter_product_line("rods")
    return redirect(url_for("history", **request.args.to_dict(flat=True)))


@app.route("/rods/report", endpoint="rods_report")
def rods_report():
    _enter_product_line("rods")
    return redirect(url_for("report", **request.args.to_dict(flat=True)))


if __name__ == "__main__":

    app.run(host="0.0.0.0", port=5004, debug=True)


