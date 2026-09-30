# -*- coding: utf-8 -*-
"""AS이력기본정보.xlsx 를 접수·이력(report_details) DB에 등록한다.

기본: 같은 접수번호가 있으면 건너뛴다.
재실행 시 엑셀 기준으로 덮어쓰려면: python import_as_history_xlsx.py --replace
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

import app as m

SEED_MARKER = "__xlsx_as_history__"
DEFAULT_XLSX = Path(r"c:\Users\AdministratorB\Desktop\AS이력기본정보.xlsx")

TEAM_NAME_TO_ID = {
    "내수영업팀": "domestic",
    "관계사영업팀": "affiliate",
    "수출영업팀": "export",
    "AS담당관": "rods_as",
    "AS담당관실": "rods_as",
}


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text.lower() in {"none", "nan"}:
        return ""
    return text


def _as_date(value):
    text = _cell(value)
    if not text:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = text.replace("/", "-").split(" ")[0]
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    return text


def _int_or_none(value):
    text = _cell(value).replace(",", "")
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _charge_type(raw):
    text = _cell(raw)
    if "유상" in text and "무상" not in text:
        return "paid"
    return "free"


def _re_repair(raw):
    text = _cell(raw).upper()
    if text in {"Y", "1", "예", "1회"}:
        return "Y"
    if text in {"N", "0", "아니오", "해당없음"}:
        return "N"
    return text or "N"


def _process_location(raw):
    text = _cell(raw)
    if text in {"공장", "factory"}:
        return "공장"
    if text:
        return text
    return "현장"


def _match_person(name):
    name = _cell(name)
    if not name:
        return None
    people = m._get_org_people()
    for person in people:
        if (person.get("name") or "").strip() == name:
            return person
    for person in people:
        pname = (person.get("name") or "").strip()
        if pname and (name.startswith(pname) or pname.startswith(name)):
            return person
    return None


def _person_id_name(raw):
    person = _match_person(raw)
    if person:
        return person.get("id") or "", person.get("name") or _cell(raw)
    return "", _cell(raw)


def _team_id_from_name(raw, fallback="affiliate"):
    name = _cell(raw)
    if name in TEAM_NAME_TO_ID:
        return TEAM_NAME_TO_ID[name]
    for code, label in TEAM_NAME_TO_ID.items():
        if code in name:
            return label
    return fallback


def _office_name(raw):
    text = _cell(raw)
    if not text:
        return ""
    for _code, name, _ord in m.plex.SALES_OFFICES:
        if text == name or text == _code:
            return name
    return text


def _row_map(headers, values):
    out = {}
    for i, header in enumerate(headers):
        key = _cell(header)
        if key:
            out[key] = values[i] if i < len(values) else None
    return out


def _read_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = None
    rows = []
    for i, raw in enumerate(ws.iter_rows(values_only=True), 1):
        values = list(raw)
        if i == 1:
            headers = [_cell(v) for v in values]
            continue
        data = _row_map(headers, values)
        if not _cell(data.get("접수번호")):
            continue
        rows.append(data)
    wb.close()
    return rows


def _build_receipt(row, next_id):
    receipt_no = _cell(row.get("접수번호"))
    team_name = _cell(row.get("담당팀")) or "관계사영업팀"
    team_id = _team_id_from_name(team_name)
    team_label = m._team_name(team_id) or team_name
    receipt_staff, receipt_staff_name = _person_id_name(row.get("접수담당자"))
    as_staff, as_staff_name = _person_id_name(row.get("AS처리담당자"))
    charge = _charge_type(row.get("유상무상"))
    billing = _int_or_none(row.get("최종청구비용"))
    margin = _cell(row.get("마진율(%)"))
    completed = _as_date(row.get("완료일")) or _as_date(row.get("출고일자"))
    receipt_date = _as_date(row.get("접수일자"))
    office = _office_name(row.get("영업사무소"))
    phone = m._format_mobile_phone(_cell(row.get("연락처"))) if hasattr(m, "_format_mobile_phone") else _cell(row.get("연락처"))

    receipt = {
        "id": next_id,
        "receipt_no": receipt_no,
        "receipt_date": receipt_date,
        "team_id": team_id,
        "team": team_label,
        "customer_name": _cell(row.get("고객명")) or "—",
        "customer_contact": _cell(row.get("고객담당자")),
        "equipment_name": _cell(row.get("장비명")) or "—",
        "symptoms": _cell(row.get("증상")) or "—",
        "as_staff": as_staff,
        "as_staff_name": as_staff_name,
        "receipt_staff": receipt_staff,
        "receipt_staff_name": receipt_staff_name,
        "office": office,
        "status": "completed" if completed else "in_progress",
        "charge_type": charge,
        "billing_cost": billing if charge == "paid" else None,
        "completed_date": completed or None,
        "created_at": f"{receipt_date}T09:00:00" if receipt_date else datetime.now().isoformat(),
        "photos": [],
        "phone": phone,
        "ship_date": _as_date(row.get("최초출하일")),
        "support_quality": False,
        "support_tech": False,
        "created_by_user_id": None,
        "created_by_username": receipt_staff_name or "엑셀이력등록",
        "severity": "medium",
        "escalation_level": "normal",
        "intervention_status": "none",
        "compensation_status": "none",
        "intervention_note": "",
        "compensation_note": "",
        "compensation_amount": "",
        "intervention_log": [],
        "intervention_by": "",
        "intervention_at": None,
        "billing_margin_rate": margin if charge == "paid" and margin else None,
        "billing_manual": 1 if charge == "paid" and billing is not None else 0,
        "billing_set_by": receipt_staff_name if charge == "paid" else "",
        "product_line": "equipment",
        "spec": "",
        "lot": "",
        "notice_sent": 1,
    }

    inspect = _as_date(row.get("점검날짜")) or receipt_date
    ship_out = _as_date(row.get("출고일자")) or completed
    affiliation = _cell(row.get("처리팀")) or team_label
    remarks = _cell(row.get("비고"))
    detail = {
        "report_date": completed or inspect or receipt_date,
        "phone": phone,
        "serial_no": _cell(row.get("시리얼넘버")),
        "re_repair": _re_repair(row.get("재수리여부")),
        "receipt_staff_name": receipt_staff_name or "—",
        "as_technician": as_staff_name or "",
        "processed_by_username": as_staff_name or "",
        "ship_date": receipt.get("ship_date") or "",
        "warranty": "",
        "inspection_date": inspect,
        "affiliation": affiliation,
        "process_location": _process_location(row.get("처리위치")),
        "ship_out_date": ship_out,
        "inspection_result": _cell(row.get("점검결과")),
        "repair_content": _cell(row.get("수리내용")),
        "detailed_work": "",
        "parts": [],
        "support_history": "지원 요청 이력 없음",
        "support_quality": False,
        "support_tech": False,
        "support_tech_mgmt": False,
        "photos": [],
        "remarks": remarks,
        "as_status": receipt["status"],
        "seed_tag": SEED_MARKER,
    }
    m._stamp_history_modified(detail, {"name": as_staff_name or "엑셀이력등록", "id": as_staff})
    return receipt, detail


def main():
    parser = argparse.ArgumentParser(description="AS이력기본정보.xlsx DB 등록")
    parser.add_argument("xlsx", nargs="?", default=str(DEFAULT_XLSX))
    parser.add_argument("--replace", action="store_true", help="같은 접수번호가 있으면 엑셀 값으로 덮어씀")
    args = parser.parse_args()
    path = Path(args.xlsx)
    if not path.exists():
        print(f"파일을 찾을 수 없습니다: {path}")
        return 1

    rows = _read_rows(path)
    if not rows:
        print("엑셀에 등록할 행이 없습니다.")
        return 1

    existing = {r.get("receipt_no") for r in m.RECEIPTS}
    by_no = {r.get("receipt_no"): r for r in m.RECEIPTS}
    next_id = max((int(r.get("id") or 0) for r in m.RECEIPTS), default=0) + 1

    added = 0
    replaced = 0
    skipped = 0
    errors = []

    for row in rows:
        receipt_no = _cell(row.get("접수번호"))
        try:
            receipt, detail = _build_receipt(row, next_id)
        except Exception as exc:
            errors.append(f"{receipt_no}: {exc}")
            continue
        if receipt_no in existing:
            if not args.replace:
                skipped += 1
                continue
            old = by_no[receipt_no]
            receipt["id"] = old.get("id") or next_id
            for i, item in enumerate(m.RECEIPTS):
                if item.get("receipt_no") == receipt_no:
                    m.RECEIPTS[i] = receipt
                    break
            m.REPORT_DETAILS[receipt_no] = detail
            replaced += 1
        else:
            receipt["id"] = next_id
            m.RECEIPTS.append(receipt)
            m.REPORT_DETAILS[receipt_no] = detail
            existing.add(receipt_no)
            by_no[receipt_no] = receipt
            next_id += 1
            added += 1

    m._next_id = next_id
    m._save_service_data()

    print(f"파일: {path}")
    print(f"엑셀 행: {len(rows)}건")
    print(f"신규 등록: {added}건")
    print(f"덮어씀: {replaced}건")
    print(f"기존 접수번호 건너뜀: {skipped}건")
    if errors:
        print("오류:")
        for line in errors:
            print(" -", line)
    print(f"현재 접수 총 {len(m.RECEIPTS)}건")
    return 0 if not errors else 2


if __name__ == "__main__":
    raise SystemExit(main())
